#!/usr/bin/env python3
"""
jmxgen - author ready-to-use JMeter .jmx test plans from a URL, an Excel sheet,
a HAR recording, or a short YAML/JSON spec - and verify the result is valid.

Platform-agnostic: the same input builds a plan that runs locally, on
HyperExecute / TestMu (LambdaTest), or any other JMeter runner. Only the
`platform:` block changes (result paths, driver paths).

Commands
  template  write an Excel/CSV template to fill in
  from-excel  Excel/CSV -> .jmx      (verified on the way out)
  from-url    crawl a live site -> .jmx
  from-har    DevTools HAR -> .jmx
  init/build  starter spec -> .jmx   (full control)
  verify      is this .jmx valid and loadable? (--deep asks JMeter itself)
  validate    performance lint (heap/CPU checklist)

PyYAML is used for specs, openpyxl for .xlsx; both optional (JSON/CSV work without).
"""

import argparse
import datetime
import json
import os
import re
import sys
from xml.sax.saxutils import escape as _xml_escape

JMETER_VERSION = "5.6.3"

# --------------------------------------------------------------------------
# tiny XML helpers
# --------------------------------------------------------------------------


def esc(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    return _xml_escape(str(v), {'"': "&quot;"})


def sp(name, val):
    return '<stringProp name="%s">%s</stringProp>' % (name, esc(val))


def bp(name, val):
    return '<boolProp name="%s">%s</boolProp>' % (name, "true" if val else "false")


def ip(name, val):
    return '<intProp name="%s">%d</intProp>' % (name, int(val))


def lp(name, val):
    return '<longProp name="%s">%d</longProp>' % (name, int(val))


def indent(lines, level):
    pad = "  " * level
    return [pad + l for l in lines]


def node(tag, attrs, children=None, subtree=None):
    """Render <tag attrs>children</tag> followed by its hashTree."""
    a = " ".join('%s="%s"' % (k, esc(v)) for k, v in attrs.items())
    out = []
    if children:
        out.append("<%s %s>" % (tag, a))
        out += indent(children, 1)
        out.append("</%s>" % tag)
    else:
        out.append("<%s %s/>" % (tag, a))
    if subtree:
        out.append("<hashTree>")
        out += indent(subtree, 1)
        out.append("</hashTree>")
    else:
        out.append("<hashTree/>")
    return out


def gui(cls, testclass, name, enabled=True):
    return {
        "guiclass": cls,
        "testclass": testclass,
        "testname": name,
        "enabled": "true" if enabled else "false",
    }


def arguments_collection(pairs, http=False, encode=True):
    """pairs: list of (name, value) -> <collectionProp> body lines.

    encode=False for a raw body: with postBodyRaw set, always_encode makes
    JMeter URL-encode the whole body, so a JSON payload arrives as
    %7B%22user%22... and the API rejects it.
    """
    out = ["<collectionProp name=\"Arguments.arguments\">"]
    for k, v in pairs:
        etype = "HTTPArgument" if http else "Argument"
        inner = []
        if http:
            inner.append(bp("HTTPArgument.always_encode", encode))
        inner.append(sp("Argument.value", v))
        inner.append(sp("Argument.metadata", "="))
        if k:
            inner.append(bp("HTTPArgument.use_equals", True) if http else "")
            inner.insert(0, sp("Argument.name", k))
        inner = [x for x in inner if x]
        out.append('  <elementProp name="%s" elementType="%s">' % (esc(k or ""), etype))
        out += indent(inner, 2)
        out.append("  </elementProp>")
    out.append("</collectionProp>")
    return out


# --------------------------------------------------------------------------
# config elements
# --------------------------------------------------------------------------


def build_user_vars(variables, name="User Defined Variables"):
    if not variables:
        return []
    body = arguments_collection(list(variables.items()))
    return node("Arguments", gui("ArgumentsPanel", "Arguments", name), body)


def _proxy_props(proxy):
    """JMeter reads these off HTTP Request Defaults / each sampler."""
    if not proxy:
        return []
    host = proxy.get("host", "")
    if not host and ":" in str(proxy.get("server", "")):
        host, _, _p = str(proxy["server"]).rpartition(":")
        proxy = dict(proxy, host=host.split("//")[-1], port=_p)
        host = proxy["host"]
    out = [sp("HTTPSampler.proxyHost", host),
           sp("HTTPSampler.proxyPort", proxy.get("port", "")),
           sp("HTTPSampler.proxyScheme", proxy.get("scheme", "http"))]
    if proxy.get("user"):
        out.append(sp("HTTPSampler.proxyUser", proxy["user"]))
    if proxy.get("password"):
        out.append(sp("HTTPSampler.proxyPass", proxy["password"]))
    return out


def build_http_defaults(d, proxy=None):
    if not d and not proxy:
        return []
    d = d or {}
    body = ['<elementProp name="HTTPsampler.Arguments" elementType="Arguments" '
            'guiclass="HTTPArgumentsPanel" testclass="Arguments" testname="User Defined Variables">',
            '  <collectionProp name="Arguments.arguments"/>',
            "</elementProp>"]
    body += [
        sp("HTTPSampler.domain", d.get("domain", "")),
        sp("HTTPSampler.port", d.get("port", "")),
        sp("HTTPSampler.protocol", d.get("protocol", "https")),
        sp("HTTPSampler.contentEncoding", d.get("encoding", "UTF-8")),
        sp("HTTPSampler.path", d.get("path", "")),
        sp("HTTPSampler.concurrentPool", "6"),
        sp("HTTPSampler.connect_timeout", d.get("connect_timeout", 5000)),
        sp("HTTPSampler.response_timeout", d.get("response_timeout", 30000)),
    ] + _proxy_props(proxy)
    return node(
        "ConfigTestElement",
        gui("HttpDefaultsGui", "ConfigTestElement", "HTTP Request Defaults"),
        body,
    )


def build_header_manager(headers, name="HTTP Header Manager"):
    if not headers:
        return []
    body = ['<collectionProp name="HeaderManager.headers">']
    for k, v in headers.items():
        body.append('  <elementProp name="" elementType="Header">')
        body.append("    " + sp("Header.name", k))
        body.append("    " + sp("Header.value", v))
        body.append("  </elementProp>")
    body.append("</collectionProp>")
    return node("HeaderManager", gui("HeaderPanel", "HeaderManager", name), body)


def build_cookie_manager(cfg):
    if not cfg:
        return []
    cfg = {} if cfg is True else cfg
    body = [
        '<collectionProp name="CookieManager.cookies"/>',
        bp("CookieManager.clearEachIteration", cfg.get("clear_each_iteration", True)),
        bp("CookieManager.controlledByThreadGroup", False),
    ]
    return node("CookieManager", gui("CookiePanel", "CookieManager", "HTTP Cookie Manager"), body)


def build_cache_manager(cfg):
    if not cfg:
        return []
    cfg = {} if cfg is True else cfg
    body = [
        bp("clearEachIteration", cfg.get("clear_each_iteration", True)),
        bp("useExpires", True),
        bp("CacheManager.controlledByThread", False),
    ]
    return node("CacheManager", gui("CacheManagerGui", "CacheManager", "HTTP Cache Manager"), body)


def build_keystore(cfg):
    """Keystore Configuration - client certificates for mutual TLS.

    JMeter reads the keystore itself from SYSTEM properties, not from the .jmx:
    this element only controls which alias a thread picks and whether the store
    is preloaded. So a plan needing mTLS is not self-contained - it travels with
    a system.properties, which `write_system_properties` emits alongside it.
    """
    cfg = {} if cfg is True else cfg
    body = [
        # a variable here lets each thread use a different client identity,
        # which is how you load-test a per-customer certificate estate
        sp("clientCertAliasVarName", cfg.get("alias_variable", "")),
        sp("startIndex", cfg.get("start_index", 0)),
        sp("endIndex", cfg.get("end_index", 0)),
        bp("preload", cfg.get("preload", True)),
    ]
    return node("KeystoreConfig", gui("TestBeanGUI", "KeystoreConfig",
                                      "Keystore Configuration"), body)


def system_properties_for(spec):
    """The system.properties a plan needs, or None.

    HyperExecute rejects -D and -J in the args array, so a property that must
    reach the JVM has to arrive as a file. This is that file.
    """
    tls = spec.get("tls")
    if not tls:
        return None
    tls = {} if tls is True else tls
    if not tls.get("keystore"):
        raise ValueError("tls: needs a 'keystore' path (a .p12 or .jks)")
    # system.properties is read by the JVM as plain java.util.Properties, so
    # JMeter functions are NOT expanded here - a ${__P(...)} password would be
    # sent literally and the handshake would fail with a confusing error
    for key in ("password", "truststore_password"):
        if "${" in str(tls.get(key) or ""):
            raise ValueError(
                "tls.%s cannot use ${...} - system.properties is read by the JVM, "
                "not by JMeter, so functions are never expanded there. Put the "
                "literal password here, or leave it blank and pass "
                "-Djavax.net.ssl.keyStorePassword=... on the command line." % key)

    lines = [
        "# Written by jmxgen. JMeter reads the client certificate from these",
        "# system properties - the .jmx cannot carry them.",
        "#",
        "#   jmeter -n -t plan.jmx -S system.properties ...",
        "#",
        "# -S, not -p: -p loads JMeter properties, and the JVM never sees these.",
        "#",
        "# On HyperExecute, upload this file with the plan; -D flags are rejected",
        "# in the args array.",
        "",
        "javax.net.ssl.keyStore=%s" % tls["keystore"],
        "javax.net.ssl.keyStorePassword=%s" % tls.get("password", ""),
    ]
    if tls.get("type"):
        lines.append("javax.net.ssl.keyStoreType=%s" % tls["type"])
    if tls.get("truststore"):
        lines += [
            "",
            "javax.net.ssl.trustStore=%s" % tls["truststore"],
            "javax.net.ssl.trustStorePassword=%s" % tls.get("truststore_password", ""),
        ]
    if tls.get("protocols"):
        lines += ["", "https.socket.protocols=%s" % tls["protocols"]]
    # JMeter reuses one SSL context per thread by default, which means the first
    # certificate wins for the life of that thread
    lines += ["", "https.use.cached.ssl.context=%s"
              % str(bool(tls.get("cache_ssl_context", False))).lower()]
    return "\n".join(lines) + "\n"


def write_system_properties(spec, jmx_path):
    """Emit system.properties next to the plan when the spec needs it."""
    text = system_properties_for(spec)
    if not text:
        return None
    out = os.path.join(os.path.dirname(os.path.abspath(jmx_path)) or ".",
                       "system.properties")
    open(out, "w", encoding="utf-8").write(text)
    return out


def build_auth_manager(entries):
    """HTTP Authorization Manager - BASIC / DIGEST / KERBEROS / NTLM."""
    body = ['<collectionProp name="AuthManager.auth_list">']
    for a in entries:
        body.append('  <elementProp name="" elementType="Authorization">')
        for k, v in (("Authorization.url", a.get("url", "")),
                     ("Authorization.username", a.get("username", "")),
                     ("Authorization.password", a.get("password", "")),
                     ("Authorization.domain", a.get("domain", "")),
                     ("Authorization.realm", a.get("realm", ""))):
            body.append("    " + sp(k, v))
        body.append("    " + sp("Authorization.mechanism",
                                (a.get("mechanism") or "BASIC").upper()))
        body.append("  </elementProp>")
    body.append("</collectionProp>")
    body.append(bp("AuthManager.clearEachIteration", False))
    body.append(bp("AuthManager.controlledByThreadGroup", False))
    return node("AuthManager", gui("AuthPanel", "AuthManager", "HTTP Authorization Manager"),
                body)


def build_jdbc_pool(cfg):
    """JDBC Connection Configuration - the pool JDBC samplers bind to by name."""
    body = [
        bp("autocommit", cfg.get("autocommit", True)),
        sp("checkQuery", cfg.get("check_query", "")),
        sp("connectionAge", cfg.get("connection_age", 5000)),
        sp("dataSource", cfg["name"]),
        sp("dbUrl", cfg["url"]),
        sp("driver", cfg["driver"]),
        sp("initQuery", cfg.get("init_query", "")),
        bp("keepAlive", True),
        sp("password", cfg.get("password", "")),
        sp("poolMax", cfg.get("pool_max", 10)),
        bp("preinit", cfg.get("preinit", False)),
        sp("timeout", cfg.get("timeout", 10000)),
        sp("transactionIsolation", "DEFAULT"),
        sp("trimInterval", 60000),
        sp("username", cfg.get("username", "")),
    ]
    return node("JDBCDataSource",
                gui("TestBeanGUI", "JDBCDataSource", "JDBC Connection - %s" % cfg["name"]),
                body)


def build_backend_listener(cfg):
    """Streams live metrics out during the run (InfluxDB / Graphite / custom)."""
    impl = cfg.get("classname",
                   "org.apache.jmeter.visualizers.backend.influxdb.InfluxdbBackendListenerClient")
    args = dict(cfg.get("arguments") or {})
    args.setdefault("influxdbUrl", cfg.get("url", ""))
    args.setdefault("application", cfg.get("application", "jmxgen"))
    args.setdefault("measurement", "jmeter")
    args.setdefault("summaryOnly", "false")
    args.setdefault("samplersRegex", ".*")
    args.setdefault("percentiles", "90;95;99")
    args.setdefault("testTitle", cfg.get("title", "Test"))
    body = ['<elementProp name="arguments" elementType="Arguments" '
            'guiclass="ArgumentsPanel" testclass="Arguments" testname="User Defined Variables">']
    body += indent(arguments_collection(list(args.items())), 1)
    body.append("</elementProp>")
    body.append(sp("classname", impl))
    body.append(sp("QUEUE_SIZE", cfg.get("queue_size", 5000)))
    return node("BackendListener",
                gui("BackendListenerGui", "BackendListener", "Backend Listener"), body)


def build_csv(cfg):
    name = cfg.get("name") or "CSV Data Set - %s" % os.path.basename(str(cfg.get("file", "")))
    body = [
        sp("filename", cfg["file"]),
        sp("fileEncoding", cfg.get("encoding", "UTF-8")),
        sp("variableNames", ",".join(cfg["variables"]) if isinstance(cfg.get("variables"), list) else cfg.get("variables", "")),
        bp("ignoreFirstLine", cfg.get("ignore_first_line", True)),
        sp("delimiter", cfg.get("delimiter", ",")),
        bp("quotedData", cfg.get("quoted", False)),
        bp("recycle", cfg.get("recycle", True)),
        bp("stopThread", cfg.get("stop_thread", False)),
        sp("shareMode", cfg.get("share_mode", "shareMode.all")),
    ]
    return node("CSVDataSet", gui("TestBeanGUI", "CSVDataSet", name), body)


def build_jdbc_or_raw(raw):
    """Escape hatch: paste any raw JMX element XML into the plan."""
    return raw.strip().splitlines() + ["<hashTree/>"] if "hashTree" not in raw else raw.strip().splitlines()


# --------------------------------------------------------------------------
# WebDriver (jp@gc plugin) config
# --------------------------------------------------------------------------

CHROME_CFG = "com.googlecode.jmeter.plugins.webdriver.config.ChromeDriverConfig"


def build_chrome_config(cfg, proxy=None):
    if not cfg:
        return []
    cfg = dict(cfg)
    if proxy and not cfg.get("proxy_type"):
        cfg["proxy_type"] = "MANUAL"
        cfg["proxy_host"] = proxy.get("host", "")
        cfg["proxy_port"] = int(proxy.get("port") or 8080)
    body = [
        bp("WebDriverConfig.acceptinsecurecerts", cfg.get("accept_insecure_certs", False)),
        bp("WebDriverConfig.reset_per_iteration", cfg.get("reset_per_iteration", True)),
        sp("WebDriverConfig.driver_path", cfg.get("driver_path", "")),
        bp("WebDriverConfig.dev_mode", False),
        bp("WebDriverConfig.headless", cfg.get("headless", True)),
        bp("WebDriverConfig.maximize_browser", cfg.get("maximize", False)),
        sp("ChromeDriverConfig.additional_args", cfg.get("args", "--no-sandbox --disable-dev-shm-usage --disable-notifications")),
        sp("ChromeDriverConfig.binary_path", cfg.get("binary_path", "")),
        sp("WebDriverConfig.proxy_type", cfg.get("proxy_type", "DIRECT")),
        sp("WebDriverConfig.proxy_pac_url", cfg.get("proxy_pac_url", "")),
        sp("WebDriverConfig.http_host", cfg.get("proxy_host", "")),
        ip("WebDriverConfig.http_port", cfg.get("proxy_port", 8080)),
        bp("WebDriverConfig.use_http_for_all_protocols", True),
        sp("WebDriverConfig.https_host", cfg.get("proxy_host", "")),
        ip("WebDriverConfig.https_port", cfg.get("proxy_port", 8080)),
        sp("WebDriverConfig.ftp_host", ""),
        ip("WebDriverConfig.ftp_port", 8080),
        sp("WebDriverConfig.socks_host", ""),
        ip("WebDriverConfig.socks_port", 8080),
        sp("WebDriverConfig.no_proxy", "localhost"),
        sp("WebDriverConfig.custom_capabilites", cfg.get("capabilities", "")),
        bp("WebDriverConfig.use_ftp_proxy", True),
        bp("WebDriverConfig.use_socks_proxy", True),
    ]
    return node(
        CHROME_CFG,
        gui(CHROME_CFG + ".gui.ChromeDriverConfigGui", CHROME_CFG, "jp@gc - Chrome Driver Config"),
        body,
    )


# --------------------------------------------------------------------------
# post-processors / assertions / timers
# --------------------------------------------------------------------------


def build_extractor(ex):
    kind = (ex.get("type") or "json").lower()
    var = ex["var"]
    if kind == "json":
        body = [
            sp("JSONPostProcessor.referenceNames", var),
            sp("JSONPostProcessor.jsonPathExprs", ex["query"]),
            sp("JSONPostProcessor.match_numbers", ex.get("match", "1")),
            sp("JSONPostProcessor.defaultValues", ex.get("default", "NOT_FOUND")),
        ]
        return node("JSONPostProcessor",
                    gui("JSONPostProcessorGui", "JSONPostProcessor", "Extract %s" % var), body)
    if kind == "regex":
        body = [
            sp("RegexExtractor.useHeaders", ex.get("scope_field", "false")),
            sp("RegexExtractor.refname", var),
            sp("RegexExtractor.regex", ex["query"]),
            sp("RegexExtractor.template", ex.get("template", "$1$")),
            sp("RegexExtractor.default", ex.get("default", "NOT_FOUND")),
            sp("RegexExtractor.match_number", ex.get("match", "1")),
        ]
        return node("RegexExtractor",
                    gui("RegexExtractorGui", "RegexExtractor", "Extract %s" % var), body)
    if kind == "boundary":
        body = [
            sp("BoundaryExtractor.useHeaders", "false"),
            sp("BoundaryExtractor.refname", var),
            sp("BoundaryExtractor.lboundary", ex["left"]),
            sp("BoundaryExtractor.rboundary", ex["right"]),
            sp("BoundaryExtractor.default", ex.get("default", "NOT_FOUND")),
            sp("BoundaryExtractor.match_number", ex.get("match", "1")),
        ]
        return node("BoundaryExtractor",
                    gui("BoundaryExtractorGui", "BoundaryExtractor", "Extract %s" % var), body)
    if kind in ("xpath", "xpath2"):
        body = [
            sp("XPathExtractor2.refname", var),
            sp("XPathExtractor2.xpathQuery", ex["query"]),
            sp("XPathExtractor2.namespaces", ex.get("namespaces", "")),
            sp("XPathExtractor2.default", ex.get("default", "NOT_FOUND")),
            sp("XPathExtractor2.matchNumber", ex.get("match", "1")),
            bp("XPathExtractor2.fragment", False),
        ]
        return node("XPath2Extractor",
                    gui("XPath2ExtractorGui", "XPath2Extractor", "Extract %s" % var), body)
    if kind in ("css", "css_selector", "jquery"):
        body = [
            sp("HtmlExtractor.refname", var),
            sp("HtmlExtractor.expr", ex["query"]),
            sp("HtmlExtractor.attribute", ex.get("attribute", "")),
            sp("HtmlExtractor.default", ex.get("default", "NOT_FOUND")),
            sp("HtmlExtractor.match_number", ex.get("match", "1")),
            sp("HtmlExtractor.extractor_impl", ex.get("impl", "JSOUP")),
        ]
        return node("HtmlExtractor",
                    gui("HtmlExtractorGui", "HtmlExtractor", "Extract %s" % var), body)
    raise ValueError("unknown extractor type: %s" % kind)


# Assertion.test_type bit flags
_ASSERT_TYPE = {"contains": 2, "matches": 1, "equals": 8, "substring": 16}
_ASSERT_FIELD = {
    "body": "Assertion.response_data",
    "code": "Assertion.response_code",
    "message": "Assertion.response_message",
    "headers": "Assertion.response_headers",
}


def build_json_assertion(a):
    """JSON Assertion - what API testers reach for first."""
    body = [
        sp("JSON_PATH", a["query"]),
        sp("EXPECTED_VALUE", a.get("pattern", "")),
        bp("JSONVALIDATION", "pattern" in a or "value" in a),
        bp("EXPECT_NULL", a.get("expect_null", False)),
        bp("INVERT", a.get("invert", False)),
        bp("ISREGEX", a.get("regex", False)),
    ]
    name = a.get("name") or "Assert %s" % a["query"]
    return node("JSONPathAssertion", gui("JSONPathAssertionGui", "JSONPathAssertion", name), body)


def build_duration_assertion(a):
    ms = a.get("max_ms", a.get("duration", 2000))
    return node("DurationAssertion",
                gui("DurationAssertionGui", "DurationAssertion",
                    a.get("name") or "Under %sms" % ms),
                [sp("DurationAssertion.duration", int(ms))])


_SIZE_OPS = {">": 3, ">=": 4, "<": 5, "<=": 6, "=": 1, "==": 1, "!=": 2}


def build_size_assertion(a):
    op = _SIZE_OPS.get(str(a.get("op", ">")), 3)
    size = int(a.get("bytes", a.get("size", 0)))
    return node("SizeAssertion",
                gui("SizeAssertionGui", "SizeAssertion",
                    a.get("name") or "Size %s %s bytes" % (a.get("op", ">"), size)),
                [sp("Assertion.test_field", "SizeAssertion.response_network_size"),
                 sp("SizeAssertion.size", size),
                 ip("SizeAssertion.operator", op)])


def build_assertion(a):
    kind = (a.get("type") or "").lower()
    if kind == "json" or (a.get("field") or "").lower() == "json":
        return build_json_assertion(a)
    if kind == "duration" or "max_ms" in a:
        return build_duration_assertion(a)
    if kind == "size" or "bytes" in a:
        return build_size_assertion(a)
    field = _ASSERT_FIELD[(a.get("field") or "body").lower()]
    ttype = _ASSERT_TYPE[(a.get("match") or "contains").lower()]
    patterns = a.get("patterns") or [a["pattern"]]
    body = ['<collectionProp name="Asserion.test_strings">']
    for p in patterns:
        body.append("  " + sp(str(abs(hash(p)) % 10**9), p))
    body.append("</collectionProp>")
    body += [
        sp("Assertion.custom_message", a.get("message", "")),
        sp("Assertion.test_field", field),
        bp("Assertion.assume_success", False),
        ip("Assertion.test_type", ttype),
    ]
    name = a.get("name") or "Assert %s %s" % (a.get("field", "body"), patterns[0][:30])
    return node("ResponseAssertion", gui("AssertionGui", "ResponseAssertion", name), body)


def build_timer(t):
    """t: int ms (constant), {min,max}, or {type: throughput|gaussian|poisson|sync}."""
    if isinstance(t, dict) and t.get("type"):
        kind = t["type"].lower()
        if kind == "throughput":                 # requests per minute, whole test
            per_min = float(t.get("per_minute", t.get("rpm", 60)))
            mode = int(t.get("calc_mode", 2))    # 2 = all active threads
            return node("ConstantThroughputTimer",
                        gui("TestBeanGUI", "ConstantThroughputTimer",
                            t.get("name") or "Throughput %s/min" % per_min),
                        [ip("calcMode", mode),
                         '<doubleProp><name>throughput</name><value>%s</value>'
                         '<savedValue>0.0</savedValue></doubleProp>' % per_min])
        if kind in ("gaussian", "poisson"):
            cls = "GaussianRandomTimer" if kind == "gaussian" else "PoissonRandomTimer"
            return node(cls,
                        gui(cls + "Gui", cls, t.get("name") or "%s timer" % kind.title()),
                        [sp("ConstantTimer.delay", int(t.get("constant", 300))),
                         sp("RandomTimer.range", float(t.get("range", 100)))])
        if kind == "sync":                       # rendezvous: release N threads together
            return node("SyncTimer",
                        gui("TestBeanGUI", "SyncTimer",
                            t.get("name") or "Rendezvous %s users" % t.get("users", 10)),
                        [ip("groupSize", int(t.get("users", 10))),
                         lp("timeoutInMs", int(t.get("timeout_ms", 0)))])
        raise ValueError("unknown timer type: %s" % kind)
    if isinstance(t, (int, float, str)):
        return node("ConstantTimer",
                    gui("ConstantTimerGui", "ConstantTimer", "Think Time %sms" % t),
                    [sp("ConstantTimer.delay", int(t))])
    base = int(t.get("constant", t.get("min", 0)))
    rng = int(t.get("random", int(t.get("max", base)) - base))
    return node("UniformRandomTimer",
                gui("UniformRandomTimerGui", "UniformRandomTimer",
                    "Think Time %s-%sms" % (base, base + rng)),
                [sp("ConstantTimer.delay", base),
                 sp("RandomTimer.range", rng)])


# --------------------------------------------------------------------------
# samplers
# --------------------------------------------------------------------------


def build_http_sampler(s):
    name = s.get("name") or "%s %s" % (s.get("method", "GET"), s.get("path", "/"))
    body_data = s.get("body")
    params = s.get("params") or {}
    raw = body_data is not None

    if raw:
        if not isinstance(body_data, str):
            body_data = json.dumps(body_data, indent=2)
        args = arguments_collection([("", body_data)], http=True, encode=False)
    else:
        args = arguments_collection(list(params.items()), http=True)

    inner = ['<elementProp name="HTTPsampler.Arguments" elementType="Arguments" '
             'guiclass="HTTPArgumentsPanel" testclass="Arguments" testname="User Defined Variables">']
    inner += indent(args, 1)
    inner.append("</elementProp>")
    if raw:
        inner.append(bp("HTTPSampler.postBodyRaw", True))
    inner += [
        sp("HTTPSampler.domain", s.get("domain", "")),
        sp("HTTPSampler.port", s.get("port", "")),
        sp("HTTPSampler.protocol", s.get("protocol", "")),
        sp("HTTPSampler.contentEncoding", s.get("encoding", "")),
        sp("HTTPSampler.path", s.get("path", "/")),
        sp("HTTPSampler.method", s.get("method", "GET")),
        bp("HTTPSampler.follow_redirects", s.get("follow_redirects", True)),
        bp("HTTPSampler.auto_redirects", False),
        bp("HTTPSampler.use_keepalive", True),
        bp("HTTPSampler.DO_MULTIPART_POST", s.get("multipart", False)),
        # embedded resources OFF unless explicitly asked: big memory cost
        bp("HTTPSampler.image_parser", s.get("embedded_resources", False)),
        sp("HTTPSampler.connect_timeout", s.get("connect_timeout", "")),
        sp("HTTPSampler.response_timeout", s.get("response_timeout", "")),
    ]

    children = []
    if s.get("headers"):
        children += build_header_manager(s["headers"], "Headers - %s" % name)
    for ex in s.get("extract", []):
        children += build_extractor(ex)
    for a in s.get("assert", []):
        children += build_assertion(a)
    if s.get("post_script"):
        # must be a CHILD of the sampler: post-processors at an outer scope can run
        # before the sampler's own extractors, and would read the variable too early
        children += build_jsr223_post(s.get("post_script_name", "Post-process"),
                                     s["post_script"])

    return node("HTTPSamplerProxy",
                gui("HttpTestSampleGui", "HTTPSamplerProxy", name, s.get("enabled", True)),
                inner, children or None)


WD_SAMPLER = "com.googlecode.jmeter.plugins.webdriver.sampler.WebDriverSampler"


def build_webdriver_sampler(s):
    name = s.get("name") or "WebDriver Step"
    script = s.get("script", "")
    if not script and s.get("actions"):
        script = groovy_from_actions(s["actions"])
    inner = [
        sp("WebDriverSampler.script", script),
        sp("WebDriverSampler.parameters", s.get("parameters", "")),
        sp("WebDriverSampler.language", "groovy"),
    ]
    return node(WD_SAMPLER,
                gui(WD_SAMPLER.replace(".sampler.", ".sampler.gui.") + "Gui", WD_SAMPLER, name,
                    s.get("enabled", True)),
                inner)


def _by_expr(action):
    """A Selenium `By` for an action, from either action shape.

    Hand-written specs carry a single `xpath`. The recorder emits a ranked
    `locators` list instead, so the best surviving locator is chosen here -
    Selenium has no fallback chain of its own, which is one more reason the
    Playwright emitter is the better target for recorded journeys.
    """
    if action.get("xpath"):
        return "By.xpath('%s')" % action["xpath"]
    for loc in action.get("locators") or []:
        kind, value = loc.get("type"), (loc.get("value") or "").replace("'", "\\'")
        if kind == "testid":
            return "By.cssSelector('[%s=\"%s\"]')" % (loc.get("attr", "data-testid"), value)
        if kind == "id":
            return "By.id('%s')" % value
        if kind == "name":
            return "By.name('%s')" % value
        if kind == "css":
            return "By.cssSelector('%s')" % value
        if kind == "text":
            return "By.xpath(\"//*[normalize-space(text())='%s']\")" % value
        if kind == "xpath":
            return "By.xpath('%s')" % value
    return None


def groovy_from_actions(actions):
    """Turn a short action list into a WebDriver Sampler groovy script."""
    head = [
        "import org.openqa.selenium.By",
        "import org.openqa.selenium.support.ui.WebDriverWait",
        "import org.openqa.selenium.support.ui.ExpectedConditions",
        "import java.time.Duration",
        "",
        "def wait = new WebDriverWait(WDS.browser, Duration.ofSeconds("
        "(vars.get('TIMEOUT') ?: '30') as int))",
        "WDS.sampleResult.sampleStart()",
        "try {",
    ]
    body = []
    for a in actions:
        kind = (a.get("do") or "").lower()
        by = _by_expr(a)
        if kind in ("open", "navigate"):
            body.append("  WDS.browser.get('%s')" % a["url"])
        elif kind in ("click", "check", "uncheck"):
            if not by:
                continue
            body.append("  wait.until(ExpectedConditions.elementToBeClickable("
                        "%s)).click()" % by)
        elif kind == "type":
            if not by:
                continue
            body.append("  def el = wait.until(ExpectedConditions.presenceOfElementLocated("
                        "%s))" % by)
            body.append("  el.clear(); el.sendKeys('%s' as String)" % a.get("text", ""))
        elif kind == "select":
            if not by:
                continue
            body.append("  new org.openqa.selenium.support.ui.Select("
                        "wait.until(ExpectedConditions.presenceOfElementLocated(%s)))"
                        ".selectByVisibleText('%s')" % (by, a.get("text", "")))
        elif kind == "wait_for":
            if not by:
                continue
            body.append("  wait.until(ExpectedConditions.presenceOfElementLocated("
                        "%s))" % by)
        elif kind == "assert_text":
            body.append("  assert WDS.browser.getPageSource().contains('%s') : "
                        "'missing: %s'" % (a["text"], a["text"]))
        elif kind == "sleep":
            body.append("  Thread.sleep(%d)" % int(a["ms"]))
        elif kind == "script":
            body.append("  " + a["code"])
    tail = [
        "  WDS.sampleResult.setSuccessful(true)",
        "} catch (Throwable t) {",
        "  WDS.sampleResult.setSuccessful(false)",
        "  WDS.sampleResult.setResponseMessage(t.toString())",
        "  WDS.log.error(t.toString())",
        "} finally {",
        "  WDS.sampleResult.sampleEnd()",
        "}",
    ]
    return "\n".join(head + body + tail)


def build_jsr223(s):
    inner = [
        sp("cacheKey", "true"),   # compile-cache on: keeps CPU sane under load
        sp("filename", s.get("file", "")),
        sp("parameters", s.get("parameters", "")),
        sp("script", s.get("script", "")),
        sp("scriptLanguage", s.get("language", "groovy")),
    ]
    return node("JSR223Sampler",
                gui("TestBeanGUI", "JSR223Sampler", s.get("name", "JSR223 Sampler")), inner)


def build_jdbc_sampler(s):
    body = [
        sp("dataSource", s.get("pool", s.get("dataSource", "db"))),
        sp("query", s.get("query", "")),
        sp("queryArguments", s.get("arguments", "")),
        sp("queryArgumentsTypes", s.get("argument_types", "")),
        sp("queryTimeout", s.get("timeout", "")),
        sp("queryType", s.get("query_type", "Select Statement")),
        sp("resultSetHandler", "Store as String"),
        sp("resultSetMaxRows", s.get("max_rows", "")),
        sp("resultVariable", s.get("result_var", "")),
        sp("variableNames", ",".join(s["variables"]) if isinstance(s.get("variables"), list)
           else s.get("variables", "")),
    ]
    children = []
    for a in s.get("assert", []):
        children += build_assertion(a)
    return node("JDBCSampler",
                gui("TestBeanGUI", "JDBCSampler", s.get("name", "JDBC Request")),
                body, children or None)


def build_graphql_sampler(s):
    """A GraphQL call is an HTTP POST of {query, variables, operationName}."""
    payload = {"query": s.get("query", "")}
    if s.get("variables") is not None:
        payload["variables"] = s["variables"]
    if s.get("operation"):
        payload["operationName"] = s["operation"]
    step = dict(s)
    step.pop("query", None)
    step.pop("variables", None)
    step.pop("operation", None)
    step["method"] = "POST"
    step["body"] = payload
    step.setdefault("path", s.get("path", "/graphql"))
    step.setdefault("name", s.get("name", "GraphQL %s" % (s.get("operation") or "query")))
    step.setdefault("headers", {}).setdefault("Content-Type", "application/json")
    return build_http_sampler(step)


def build_sampler(s):
    kind = (s.get("type") or "http").lower()
    if kind == "http":
        return build_http_sampler(s)
    if kind == "jdbc":
        return build_jdbc_sampler(s)
    if kind == "graphql":
        return build_graphql_sampler(s)
    if kind in ("webdriver", "wd", "browser"):
        return build_webdriver_sampler(s)
    if kind in ("jsr223", "groovy"):
        return build_jsr223(s)
    if kind == "raw":
        return build_jdbc_or_raw(s["xml"])
    raise ValueError("unknown sampler type: %s" % kind)


# --------------------------------------------------------------------------
# controllers / thread groups
# --------------------------------------------------------------------------


PARALLEL_CTRL = "com.blazemeter.jmeter.controller.ParallelSampler"


def build_parallel(name, children, max_threads=6):
    """bzm - Parallel Controller. The cap is only honoured when
    LIMIT_MAX_THREAD_NUMBER is true, so it is always set here."""
    inner = [
        bp("PARENT_SAMPLE", True),
        bp("LIMIT_MAX_THREAD_NUMBER", True),
        ip("MAX_THREAD_NUMBER", max_threads),
        ip("ParallelSampler.maxThreadNumber", max_threads),
        ip("ParallelSampler.timeout", 0),
    ]
    return node(PARALLEL_CTRL,
                gui(PARALLEL_CTRL.replace(".ParallelSampler", ".ParallelControllerGui"),
                    PARALLEL_CTRL, name),
                inner, children)


def build_pause(ms, name=None):
    """Flow Control Action 'pause' - waits exactly here, unlike a scoped timer."""
    return node("TestAction",
                gui("TestActionGui", "TestAction", name or "Pause %sms" % ms),
                [ip("ActionProcessor.action", 1),      # 1 = pause
                 ip("ActionProcessor.target", 0),      # 0 = current thread
                 sp("ActionProcessor.duration", int(ms))])


def build_steps(steps):
    out = []
    for s in steps:
        if "pause" in s:
            out += build_pause(s["pause"], s.get("name"))
            continue
        if "parallel" in s:
            out += build_parallel(s["parallel"], build_steps(s["steps"]),
                                  s.get("max_threads", 6))
            continue
        if "transaction" in s:
            children = build_steps(s["steps"])
            out += node("TransactionController",
                        gui("TransactionControllerGui", "TransactionController", s["transaction"]),
                        [bp("TransactionController.includeTimers", False),
                         bp("TransactionController.parent", True)],
                        children)
            continue
        if "if" in s:
            children = build_steps(s["steps"])
            out += node("IfController", gui("IfControllerPanel", "IfController",
                                            s.get("name", "If %s" % s["if"])),
                        [sp("IfController.condition", s["if"]),
                         bp("IfController.evaluateAll", False),
                         bp("IfController.useExpression", True)],
                        children)
            continue
        if "loop" in s:
            children = build_steps(s["steps"])
            out += node("LoopController", gui("LoopControlPanel", "LoopController",
                                              s.get("name", "Loop x%s" % s["loop"])),
                        [bp("LoopController.continue_forever", False),
                         sp("LoopController.loops", s["loop"])],
                        children)
            continue
        if "while" in s:
            out += node("WhileController",
                        gui("WhileControllerGui", "WhileController",
                            s.get("name", "While %s" % s["while"])),
                        [sp("WhileController.condition", s["while"])],
                        build_steps(s["steps"]))
            continue
        if "switch" in s:
            out += node("SwitchController",
                        gui("SwitchControllerGui", "SwitchController",
                            s.get("name", "Switch")),
                        [sp("SwitchController.value", s["switch"])],
                        build_steps(s["steps"]))
            continue
        if "interleave" in s:
            out += node("InterleaveControl",
                        gui("InterleaveControlGui", "InterleaveControl",
                            s.get("name", "Interleave")),
                        [ip("InterleaveControl.style", int(s.get("style", 1)))],
                        build_steps(s["steps"]))
            continue
        if "random" in s:
            out += node("RandomController",
                        gui("RandomControlGui", "RandomController",
                            s.get("name", "Random")),
                        [ip("InterleaveControl.style", int(s.get("style", 1)))],
                        build_steps(s["steps"]))
            continue
        if "throughput" in s:
            pct = float(s["throughput"])
            out += node("ThroughputController",
                        gui("ThroughputControllerGui", "ThroughputController",
                            s.get("name", "Only %s%% of users" % pct)),
                        [ip("ThroughputController.style", 1),      # 1 = percent
                         bp("ThroughputController.perThread", False),
                         ip("ThroughputController.maxThroughput", 1),
                         '<FloatProperty><name>ThroughputController.percentThroughput</name>'
                         '<value>%s</value><savedValue>0.0</savedValue></FloatProperty>' % pct],
                        build_steps(s["steps"]))
            continue
        if "runtime" in s:
            out += node("RunTime",
                        gui("RunTimeGui", "RunTime", s.get("name", "Run %ss" % s["runtime"])),
                        [sp("RunTime.seconds", s["runtime"])],
                        build_steps(s["steps"]))
            continue
        if "once" in s:
            children = build_steps(s["steps"])
            out += node("OnceOnlyController",
                        gui("OnceOnlyControllerGui", "OnceOnlyController", s.get("name", "Once Only")),
                        [], children)
            continue
        out += build_sampler(s)
        if s.get("think_time") is not None:
            out += build_timer(s["think_time"])
    return out


def build_jsr223_post(name, script, language="groovy"):
    return node("JSR223PostProcessor",
                gui("TestBeanGUI", "JSR223PostProcessor", name),
                [sp("cacheKey", "true"),
                 sp("filename", ""),
                 sp("parameters", ""),
                 sp("script", script),
                 sp("scriptLanguage", language)])


def build_auth_setup(auth):
    """A setUp Thread Group that logs in once and publishes the token.

    Variables are per-thread, so the token is promoted to a JMeter *property*;
    the request threads then read it with ${__P(NAME)}. This is the standard way
    to avoid every virtual user re-authenticating."""
    var = auth.get("var", "AUTH_TOKEN")
    login = dict(auth.get("login") or {})
    if not login.get("path"):
        raise ValueError("auth.login needs at least a 'path'")

    ex = dict(login.pop("extract", None) or {})
    ex.setdefault("type", "json")
    ex.setdefault("query", "$.access_token")
    ex["var"] = var

    login.setdefault("name", "Authenticate")
    login.setdefault("method", "POST")
    login["extract"] = [ex]
    login.setdefault("assert", [{"field": "code", "match": "equals", "pattern": "200"}])

    login["post_script_name"] = "Publish %s to a property" % var
    login["post_script"] = 'props.put("%s", vars.get("%s") ?: "")' % (var, var)
    children = build_sampler(login)

    inner = [
        sp("ThreadGroup.on_sample_error", "stopthread"),
        '<elementProp name="ThreadGroup.main_controller" elementType="LoopController" '
        'guiclass="LoopControlPanel" testclass="LoopController" testname="Loop Controller">',
        "  " + bp("LoopController.continue_forever", False),
        "  " + sp("LoopController.loops", 1),
        "</elementProp>",
        sp("ThreadGroup.num_threads", 1),
        sp("ThreadGroup.ramp_time", 1),
        bp("ThreadGroup.scheduler", False),
        sp("ThreadGroup.duration", ""),
        sp("ThreadGroup.delay", ""),
    ]
    return node("SetupThreadGroup",
                gui("SetupThreadGroupGui", "SetupThreadGroup", "setUp - authenticate"),
                inner, children)


def auth_header(auth):
    var = auth.get("var", "AUTH_TOKEN")
    if auth.get("header"):
        name, _, value = str(auth["header"]).partition(":")
        return {name.strip(): value.strip()}
    scheme = auth.get("scheme", "Bearer")
    return {"Authorization": "%s ${__P(%s)}" % (scheme, var) if scheme
            else "${__P(%s)}" % var}


def build_thread_group(tg, defaults):
    """Emit the thread group the spec asked for.

    `model:` picks the workload shape:
      closed      (default) N users looping - the plain JMeter Thread Group
      concurrency hold N concurrent users, JMeter manages the threads
      arrivals    N iterations started per second/minute, regardless of latency

    Arrival rate is the one worth reaching for when the target is written as
    throughput: under a closed model a slow response reduces the offered load,
    which quietly hides the very problem you are testing for.
    """
    model = str(tg.get("model", "closed")).lower()
    if model in ("arrivals", "arrival", "rate", "rps", "open"):
        return _casutg_group(tg, "arrivals")
    if model in ("concurrency", "concurrent"):
        return _casutg_group(tg, "concurrency")
    if model not in ("closed", "threads", ""):
        raise ValueError("unknown thread group model %r - use closed, concurrency or arrivals"
                        % tg.get("model"))

    name = tg.get("name", "Thread Group")
    threads = tg.get("threads", 1)
    ramp = tg.get("ramp_up", 1)
    duration = tg.get("duration")
    loops = tg.get("loops", 1)
    scheduler = bool(duration)

    loop_ctrl = [
        '<elementProp name="ThreadGroup.main_controller" elementType="LoopController" '
        'guiclass="LoopControlPanel" testclass="LoopController" testname="Loop Controller">',
        "  " + bp("LoopController.continue_forever", False),
        # -1 is only safe because the scheduler time-boxes the run
        "  " + sp("LoopController.loops", -1 if scheduler else loops),
        "</elementProp>",
    ]
    inner = [
        sp("ThreadGroup.on_sample_error", tg.get("on_error", "continue")),
    ] + loop_ctrl + [
        # authored values become the DEFAULTS of JMeter properties, so the same
        # file runs at any load without editing:  jmeter -Jthreads=50 -Jduration=900
        sp("ThreadGroup.num_threads", "${__P(threads,%s)}" % threads),
        sp("ThreadGroup.ramp_time", "${__P(ramp,%s)}" % ramp),
        bp("ThreadGroup.same_user_on_next_iteration", tg.get("same_user", True)),
        bp("ThreadGroup.scheduler", scheduler),
        sp("ThreadGroup.duration",
           "${__P(duration,%s)}" % duration if duration else ""),
        sp("ThreadGroup.delay", tg.get("startup_delay", "")),
    ]

    children = []
    if tg.get("variables"):
        children += build_user_vars(tg["variables"], "Vars - %s" % name)
    for csv in tg.get("csv", []):
        children += build_csv(csv)
    if tg.get("headers"):
        children += build_header_manager(tg["headers"], "Headers - %s" % name)
    children += build_steps(tg.get("steps", []))

    return node("ThreadGroup", gui("ThreadGroupGui", "ThreadGroup", name, tg.get("enabled", True)),
                inner, children)


# jpgc Custom Thread Groups. These express an OPEN workload - you state the rate
# or the concurrency you want and JMeter works out the threads - which is how
# load targets are normally written ("500 orders per second"), and what a plain
# Thread Group cannot say at all.
CASUTG = "com.blazemeter.jmeter.threads"
VU_CONTROLLER = ('<elementProp name="ThreadGroup.main_controller" '
                 'elementType="com.blazemeter.jmeter.control.VirtualUserController"/>')


def _casutg_group(tg, kind):
    """Shared body for the arrivals and concurrency thread groups.

    `rate` (arrivals) and `threads` (concurrency) both land in TargetLevel -
    the element is the same shape, only the meaning of the number changes.
    """
    name = tg.get("name", "Thread Group")
    if kind == "arrivals":
        cls = "%s.arrivals.ArrivalsThreadGroup" % CASUTG
        guicls = "%s.arrivals.ArrivalsThreadGroupGui" % CASUTG
        label = "bzm - Arrivals Thread Group"
        target = tg.get("rate", tg.get("arrival_rate", 1))
        prop = "rate"
    else:
        cls = "%s.concurrency.ConcurrencyThreadGroup" % CASUTG
        guicls = "%s.concurrency.ConcurrencyThreadGroupGui" % CASUTG
        label = "bzm - Concurrency Thread Group"
        target = tg.get("threads", 1)
        prop = "threads"

    # per second unless the spec says per minute - "600 per minute" reads better
    # for slow business flows and is what capacity docs usually quote
    unit = str(tg.get("unit", "S")).upper()[:1]
    unit = "M" if unit == "M" else "S"

    inner = [
        VU_CONTROLLER,
        sp("ThreadGroup.on_sample_error", tg.get("on_error", "continue")),
        # same __P() treatment as the plain group, so one file runs at any load
        sp("TargetLevel", "${__P(%s,%s)}" % (prop, target)),
        sp("RampUp", "${__P(ramp,%s)}" % tg.get("ramp_up", 0)),
        sp("Steps", tg.get("steps_count", 1)),
        sp("Hold", "${__P(duration,%s)}" % (tg.get("duration") or 0)),
        sp("LogFilename", tg.get("log_file", "")),
        sp("Unit", unit),
    ]
    if kind == "arrivals":
        # a runaway target must not open unbounded threads; blank means no cap
        inner.append(sp("ConcurrencyLimit", tg.get("max_concurrency", "")))
        inner.append(sp("Iterations", tg.get("loops", "")))

    children = []
    if tg.get("variables"):
        children += build_user_vars(tg["variables"], "Vars - %s" % name)
    for csv in tg.get("csv", []):
        children += build_csv(csv)
    if tg.get("headers"):
        children += build_header_manager(tg["headers"], "Headers - %s" % name)
    children += build_steps(tg.get("steps", []))

    return node(cls, gui(guicls, cls, "%s (%s)" % (name, label), tg.get("enabled", True)),
                inner, children)


# --------------------------------------------------------------------------
# listeners
# --------------------------------------------------------------------------

SAVE_CONFIG = [
    '<objProp>',
    '  <name>saveConfig</name>',
    '  <value class="SampleSaveConfiguration">',
    '    <time>true</time><latency>true</latency><timestamp>true</timestamp>',
    '    <success>true</success><label>true</label><code>true</code>',
    '    <message>true</message><threadName>true</threadName>',
    '    <dataType>false</dataType><encoding>false</encoding><assertions>true</assertions>',
    '    <subresults>false</subresults><responseData>false</responseData>',
    '    <samplerData>false</samplerData><xml>false</xml><fieldNames>true</fieldNames>',
    '    <responseHeaders>false</responseHeaders><requestHeaders>false</requestHeaders>',
    '    <responseDataOnError>false</responseDataOnError>',
    '    <saveAssertionResultsFailureMessage>true</saveAssertionResultsFailureMessage>',
    '    <assertionsResultsToSave>0</assertionsResultsToSave>',
    '    <bytes>true</bytes><sentBytes>true</sentBytes><url>true</url>',
    '    <threadCounts>true</threadCounts><idleTime>true</idleTime><connectTime>true</connectTime>',
    '  </value>',
    '</objProp>',
]


def build_data_writer(path):
    body = list(SAVE_CONFIG) + [sp("filename", path)]
    return node("ResultCollector", gui("SimpleDataWriter", "ResultCollector", "Simple Data Writer"), body)


# --------------------------------------------------------------------------
# platform profiles
# --------------------------------------------------------------------------

PLATFORMS = {
    "local": {
        "results": "results.csv",
        "driver_path": "",
        "binary_path": "",
        "headless": False,
    },
    "hyperexecute": {
        # HyperExecute/TestMu Linux runner: driver + chrome are pre-installed here
        "results": "${__P(results.file,results.csv)}",
        "driver_path": "/home/ltuser/lrc/drivers/chrome/141.0/chromedriver",
        "binary_path": "/home/ltuser/lrc/chrome/Google-Chrome-141.0/opt/google/chrome/chrome",
        "headless": True,
    },
}
PLATFORMS["testmu"] = PLATFORMS["hyperexecute"]


def resolve_platform(spec):
    pname = (spec.get("platform") or "local")
    if isinstance(pname, dict):
        base = dict(PLATFORMS.get(pname.get("name", "local"), PLATFORMS["local"]))
        base.update({k: v for k, v in pname.items() if k != "name"})
        return base
    return dict(PLATFORMS.get(pname, PLATFORMS["local"]))


# --------------------------------------------------------------------------
# top level
# --------------------------------------------------------------------------


def parameterize(spec, bindings, csv_files=None):
    """Replace hard-coded literals with ${CSV columns} everywhere in the plan.

    bindings: {"bob@example.com": "email", "hunter2": "password"} - the literal a
    recording captured, and the CSV column that should drive it instead."""
    if not bindings:
        return spec, 0
    hits = [0]

    def walk(node_):
        if isinstance(node_, dict):
            return {k: (v if k in ("extract", "assert") else walk(v))
                    for k, v in node_.items()}
        if isinstance(node_, list):
            return [walk(v) for v in node_]
        if isinstance(node_, str):
            for literal, col in bindings.items():
                if literal and literal in node_:
                    node_ = node_.replace(literal, "${%s}" % col)
                    hits[0] += 1
            return node_
        return node_

    spec = walk(spec)
    if csv_files:
        spec["csv"] = (spec.get("csv") or []) + csv_files
    return spec, hits[0]


SUGGEST_PATTERNS = [
    ("email", re.compile(r"[A-Za-z0-9._%%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")),
    ("phone", re.compile(r"\+?\d[\d\-. ]{8,15}\d")),
    ("username", re.compile(r'(?:user|username|login|userid|customer)"?\s*[:=]\s*"?'
                            r'([A-Za-z0-9._-]{3,40})')),
    ("password", re.compile(r'(?:password|passwd|pwd)"?\s*[:=]\s*"?([^"&,}\s]{4,60})')),
]


def suggest_parameters(spec):
    """Values that look like per-user data a load test should vary.

    Scans the serialized spec directly - the interesting values usually sit
    inside an escaped JSON body, not as tidy top-level fields."""
    text = json.dumps(spec)
    found = {}
    for label, pattern in SUGGEST_PATTERNS:
        for m in pattern.finditer(text):
            value = m.group(1) if pattern.groups else m.group(0)
            value = value.strip().strip('"')
            if not value or "${" in value or value.startswith(("http", "/", "application")):
                continue
            found.setdefault(value, label)
    return found


def build_plan(spec):
    # Validate the TLS block here rather than at the point the properties file
    # is written: otherwise a bad password expression fails AFTER the .jmx has
    # been written, leaving a broken plan on disk that looks like a build.
    if spec.get("tls"):
        system_properties_for(spec)

    plat = resolve_platform(spec)
    name = spec.get("name", "Test Plan")

    tp_inner = [
        bp("TestPlan.functional_mode", False),          # keeps response data out of heap
        bp("TestPlan.tearDown_on_shutdown", True),
        bp("TestPlan.serialize_threadgroups", spec.get("serialize_thread_groups", False)),
        sp("TestPlan.comments", spec.get("comments", "Generated by jmxgen")),
        sp("TestPlan.user_define_classpath", spec.get("classpath", "")),
    ]
    variables = dict(spec.get("variables") or {})
    if variables:
        vbody = arguments_collection(list(variables.items()))
        tp_inner.append('<elementProp name="TestPlan.user_defined_variables" elementType="Arguments" '
                        'guiclass="ArgumentsPanel" testclass="Arguments" testname="User Defined Variables">')
        tp_inner += indent(vbody, 1)
        tp_inner.append("</elementProp>")
    else:
        tp_inner.append('<elementProp name="TestPlan.user_defined_variables" elementType="Arguments" '
                        'guiclass="ArgumentsPanel" testclass="Arguments" testname="User Defined Variables">'
                        '<collectionProp name="Arguments.arguments"/></elementProp>')

    proxy = spec.get("proxy")
    auth = spec.get("auth")
    plan_headers = dict(spec.get("headers") or {})
    if auth:
        plan_headers.update(auth_header(auth))

    children = []
    children += build_http_defaults(spec.get("defaults"), proxy)
    children += build_header_manager(plan_headers)
    children += build_cookie_manager(spec.get("cookies", True))
    children += build_cache_manager(spec.get("cache", False))
    if spec.get("http_auth"):
        children += build_auth_manager(spec["http_auth"])
    if spec.get("tls"):
        children += build_keystore(spec["tls"])
    for pool in spec.get("jdbc", []):
        children += build_jdbc_pool(pool)
    for csv in spec.get("csv", []):
        children += build_csv(csv)

    wd = spec.get("webdriver")
    if wd:
        wd = dict(wd)
        wd.setdefault("driver_path", plat["driver_path"])
        wd.setdefault("binary_path", plat["binary_path"])
        wd.setdefault("headless", plat["headless"])
        children += build_chrome_config(wd, proxy)

    if auth:
        children += build_auth_setup(auth)

    for tg in spec.get("thread_groups", []):
        children += build_thread_group(tg, spec.get("defaults"))

    if spec.get("backend_listener"):
        children += build_backend_listener(spec["backend_listener"])
    children += build_data_writer(spec.get("results_file", plat["results"]))

    plan = node("TestPlan", gui("TestPlanGui", "TestPlan", name), tp_inner, children)

    out = ['<?xml version="1.0" encoding="UTF-8"?>',
           '<jmeterTestPlan version="1.2" properties="5.0" jmeter="%s">' % JMETER_VERSION,
           "  <hashTree>"]
    out += indent(plan, 2)
    out += ["  </hashTree>", "</jmeterTestPlan>", ""]
    return "\n".join(out)


# --------------------------------------------------------------------------
# spec loading
# --------------------------------------------------------------------------


def _yaml_help(path, exc):
    """A YAML parse error the author can act on, not a stack trace.

    Nearly every one of these is a ${...} or a [0] inside a {flow mapping},
    where the braces and brackets are YAML syntax - so say so.
    """
    mark = getattr(exc, "problem_mark", None)
    where = " line %d, column %d" % (mark.line + 1, mark.column + 1) if mark else ""
    lines = ["%s: %s%s" % (path, getattr(exc, "problem", None) or "could not parse", where)]
    if mark:
        try:
            src = open(path, "r", encoding="utf-8").read().splitlines()
            if mark.line < len(src):
                lines.append("    %s" % src[mark.line].rstrip())
                lines.append("    %s^" % (" " * mark.column))
        except OSError:
            pass
    lines.append("")
    lines.append("  A value containing ${...}, [0] or {...} must be quoted inside a")
    lines.append("  {flow mapping} - the braces and brackets are YAML syntax:")
    lines.append("      body: {sku: \"SKU-${ID}\"}")
    lines.append("      extract: [{type: json, var: ID, query: \"$.items[0].id\"}]")
    return "\n".join(lines)


def load_spec(path):
    text = open(path, "r", encoding="utf-8").read()
    if path.endswith((".yaml", ".yml")):
        try:
            import yaml  # type: ignore
        except ImportError:
            sys.exit("PyYAML not installed. Run: pip install pyyaml   (or write the spec as .json)")
        try:
            return yaml.safe_load(text)
        except yaml.YAMLError as exc:
            sys.exit(_yaml_help(path, exc))
    try:
        return json.loads(text)
    except ValueError as exc:
        sys.exit("%s is not valid JSON: %s" % (path, exc))


def _have_yaml():
    try:
        import yaml  # noqa: F401
        return True
    except ImportError:
        return False


def dump_spec(spec, path):
    if path.endswith((".yaml", ".yml")) and _have_yaml():
        import yaml
        return yaml.safe_dump(spec, sort_keys=False, default_flow_style=False)
    return json.dumps(spec, indent=2)


# --------------------------------------------------------------------------
# display helpers - shared by the console and the in-browser engine
# --------------------------------------------------------------------------


def _flatten(spec, depth=0, out=None, group=""):
    """A flat, display-friendly view of the plan's steps."""
    if out is None:
        out = []
    for tg in spec.get("thread_groups", []):
        _walk_display(tg.get("steps", []), out, tg.get("name", "Thread Group"))
    return out


def _walk_display(steps, out, group):
    for st in steps:
        if "transaction" in st or "parallel" in st:
            label = st.get("transaction") or st.get("parallel")
            _walk_display(st.get("steps", []), out, label)
            continue
        if "steps" in st:
            _walk_display(st["steps"], out, group)
            continue
        if "pause" in st:
            out.append({"group": group, "name": "pause %sms" % st["pause"],
                        "method": "", "path": "", "kind": "pause"})
            continue
        out.append({
            "group": group,
            "name": st.get("name", ""),
            "method": st.get("method", st.get("type", "")),
            "path": str(st.get("path", ""))[:120],
            "asserts": len(st.get("assert", []) or []),
            "extracts": [e.get("var") for e in (st.get("extract", []) or [])],
            "think_time": st.get("think_time"),
            "kind": st.get("type", "http"),
        })


def _plan_load(spec):
    """The first thread group's load, for prefilling the run form."""
    tgs = spec.get("thread_groups") or []
    tg = tgs[0] if tgs else {}
    return {"threads": tg.get("threads") or 1,
            "ramp_up": tg.get("ramp_up") or 1,
            "duration": tg.get("duration") or ""}


# --------------------------------------------------------------------------
# recorded browser actions -> spec steps -> Playwright
# --------------------------------------------------------------------------


def actions_to_steps(actions):
    """Recorded GUI actions -> webdriver steps, grouped by transaction.

    The recorder emits a flat list because that is what actually happened; the
    spec wants them grouped the way the rest of a plan is grouped, so a browser
    test reads with the same shape as the protocol one.
    """
    groups = []
    for act in actions or []:
        tx = act.get("transaction") or "Flow"
        if not groups or groups[-1]["transaction"] != tx:
            groups.append({"transaction": tx, "steps": []})
        groups[-1]["steps"].append(act)

    out = []
    for grp in groups:
        inner = []
        for act in grp["steps"]:
            inner.append({"type": "webdriver",
                          "name": _action_name(act),
                          "actions": [_action_spec(act)]})
        out.append({"transaction": grp["transaction"], "steps": inner})
    return out


def _action_name(act):
    verb = act.get("do", "do")
    if verb == "navigate":
        return "Open %s" % (act.get("url") or "")
    label = act.get("label") or "element"
    if verb == "type":
        return "Type into %s" % label
    return "%s %s" % (verb.capitalize(), label)


def _action_spec(act):
    """One recorded action -> one spec action, keeping the whole locator list.

    The locators are carried through rather than collapsed to one: whichever
    runner consumes this can try them in order, which is the entire point of
    verifying several at capture time.
    """
    verb = act.get("do")
    out = {"do": verb}
    if verb == "navigate":
        out["url"] = act.get("url")
        return out
    out["locators"] = act.get("locators") or []
    if act.get("text") is not None:
        out["text"] = act["text"]
    if act.get("weak"):
        out["weak"] = True
    return out


def _pw_locator(loc):
    """One recorded locator -> the Playwright call that finds it."""
    kind, value = loc.get("type"), loc.get("value")
    if kind == "testid":
        return "page.get_by_test_id(%r)" % value
    if kind == "id":
        return "page.locator(%r)" % ("#" + value)
    if kind == "name":
        return "page.locator(%r)" % ('[name="%s"]' % value)
    if kind == "text":
        return "page.get_by_text(%r, exact=True)" % value
    if kind == "xpath":
        return "page.locator(%r)" % ("xpath=" + value)
    return "page.locator(%r)" % value


def _pw_action(act, indent="    "):
    """One spec action -> Playwright lines, with locator fallback."""
    verb = act.get("do")
    if verb == "navigate":
        return ["%spage.goto(%r)" % (indent, act.get("url") or "/")]

    locs = act.get("locators") or []
    if not locs:
        return ["%s# skipped %s - no locator was recorded" % (indent, verb)]

    if verb == "assert_text":
        return ["%sexpect(%s).to_contain_text(%r)"
                % (indent, _pw_locator(locs[0]), act.get("text") or "")]

    call = {"click": ".click()",
            "type": ".fill(%r)" % (act.get("text") or ""),
            "select": ".select_option(label=%r)" % (act.get("text") or ""),
            "check": ".check()",
            "uncheck": ".uncheck()"}.get(verb, ".click()")

    args = ", ".join(_pw_locator(l) for l in locs)
    return ["%sfirst_of(page, [%s])%s" % (indent, args, call)]


PW_PRELUDE = '''"""Browser test generated by jmxgen.

Run it:
    pip install playwright && python3 -m playwright install chromium
    python3 %(file)s

This is the FUNCTIONAL half of the recording - it drives one real browser and
proves the journey still works. It does not carry load: the .jmx from the same
session is what scales to thousands of users. A browser per virtual user is not
a load test, it is a way to measure your own hardware.
"""

from playwright.sync_api import sync_playwright, expect


def first_of(page, locators, timeout=5000):
    """The first locator that is actually present.

    Every element was recorded with several locators, ranked most-stable first
    and each verified against the live DOM at capture time. Trying them in
    order is what lets the script survive a page that has moved on.
    """
    last = None
    for loc in locators:
        try:
            loc.wait_for(state="attached", timeout=max(timeout // len(locators), 500))
            return loc
        except Exception as exc:          # noqa: BLE001 - try the next one
            last = exc
    raise AssertionError("no recorded locator matched: %%s" %% last)
'''


def spec_to_playwright(spec, filename="browser_test.py"):
    """Render the webdriver steps of a spec as a runnable Playwright script."""
    body = [PW_PRELUDE % {"file": filename}, "", "def run(page):"]
    wrote = False
    for tg in spec.get("thread_groups") or []:
        for node in tg.get("steps") or []:
            if node.get("transaction"):
                inner = [x for x in node.get("steps") or []
                         if x.get("type") == "webdriver"]
                if not inner:
                    continue
                body.append("    # --- %s ---" % node["transaction"])
                for step in inner:
                    body.append("    # %s" % step.get("name", ""))
                    for act in step.get("actions") or []:
                        body += _pw_action(act)
                    think = step.get("think_time")
                    if isinstance(think, dict):
                        think = think.get("min")
                    if think:
                        body.append("    page.wait_for_timeout(%d)" % int(think))
                    wrote = True
            elif node.get("type") == "webdriver":
                for act in node.get("actions") or []:
                    body += _pw_action(act)
                wrote = True
    if not wrote:
        body.append("    raise SystemExit('this plan has no browser steps')")

    headless = bool((spec.get("webdriver") or {}).get("headless", True))
    body += [
        "",
        "",
        "if __name__ == '__main__':",
        "    with sync_playwright() as pw:",
        "        browser = pw.chromium.launch(headless=%s)" % headless,
        "        page = browser.new_page()",
        "        try:",
        "            run(page)",
        "            print('browser test passed')",
        "        finally:",
        "            browser.close()",
        "",
    ]
    return "\n".join(body)


def spec_has_browser_steps(spec):
    """True when the plan carries anything Playwright could run."""
    for tg in spec.get("thread_groups") or []:
        for step in _walk_steps(tg.get("steps") or []):
            if step.get("type") == "webdriver":
                return True
    return False


# --------------------------------------------------------------------------
# spec -> Taurus YAML
# --------------------------------------------------------------------------


def _taurus_time(value, unit="s"):
    """Taurus wants a duration with a unit; the spec keeps plain seconds."""
    if value in (None, "", 0):
        return None
    return "%s%s" % (value, unit)


def _taurus_assert(rules):
    """spec assertions -> Taurus `assert` blocks.

    Taurus splits by subject, and its http-code subject only takes strings, so
    a numeric code from the spec has to be stringified or the run fails at
    validation rather than at assertion time.
    """
    out = []
    for rule in rules or []:
        field = (rule.get("field") or "body").lower()
        subject = {"code": "http-code", "headers": "http-headers",
                   "body": "body", "message": "body"}.get(field, "body")
        pattern = rule.get("pattern")
        if pattern is None:
            continue
        entry = {"contains": [str(pattern)], "subject": subject}
        match = (rule.get("match") or "contains").lower()
        if match in ("matches", "regex", "regexp"):
            entry["regexp"] = True
        elif match == "equals":
            # Taurus has no equals; an anchored regexp is the honest equivalent
            entry["contains"] = ["^%s$" % re.escape(str(pattern))]
            entry["regexp"] = True
        if rule.get("negate") or match.startswith("not"):
            entry["not"] = True
        out.append(entry)
    return out


def _taurus_extract(extracts, req):
    """spec extractors -> Taurus extract-* maps, keyed by variable name."""
    for ex in extracts or []:
        var = ex.get("var")
        if not var:
            continue
        kind = (ex.get("type") or "regex").lower()
        query = ex.get("query") or ex.get("pattern") or ""
        if kind in ("json", "jsonpath"):
            req.setdefault("extract-jsonpath", {})[var] = {
                "jsonpath": query, "default": ex.get("default", "NOT_FOUND")}
        elif kind in ("xpath",):
            req.setdefault("extract-xpath", {})[var] = {
                "xpath": query, "default": ex.get("default", "NOT_FOUND")}
        elif kind in ("css", "css-jquery"):
            req.setdefault("extract-css-jquery", {})[var] = {
                "expression": query, "default": ex.get("default", "NOT_FOUND")}
        else:
            req.setdefault("extract-regexp", {})[var] = {
                "regexp": query,
                "match-no": ex.get("match_no", 1),
                "template": ex.get("template", "$1$"),
                "default": ex.get("default", "NOT_FOUND")}


def _taurus_request(step):
    """One spec step -> one Taurus request."""
    req = {"label": step.get("name") or step.get("path") or "request",
           "url": step.get("path") or step.get("url") or "/"}
    method = (step.get("method") or "GET").upper()
    if method != "GET":
        req["method"] = method
    if step.get("headers"):
        req["headers"] = dict(step["headers"])
    body = step.get("body")
    if body is not None:
        # a dict becomes form/JSON params, a string goes through as a raw body
        req["body"] = body if isinstance(body, (dict, list)) else str(body)
    asserts = _taurus_assert(step.get("assert"))
    if asserts:
        req["assert"] = asserts
    _taurus_extract(step.get("extract"), req)
    think = step.get("think_time")
    if isinstance(think, dict):
        # Taurus has no range, so the midpoint is the honest single value
        lo, hi = think.get("min", 0), think.get("max", think.get("min", 0))
        think = int((lo + hi) / 2)
    if think:
        req["think-time"] = _taurus_time(think, "ms")
    return req


def _taurus_steps(steps):
    """Walk the spec's steps, preserving transaction nesting."""
    out = []
    for step in steps or []:
        if step.get("transaction"):
            out.append({"transaction": step["transaction"],
                        "do": _taurus_steps(step.get("steps"))})
        elif step.get("kind") in (None, "http"):
            out.append(_taurus_request(step))
    return out


def spec_to_taurus(spec):
    """Render the spec as a Taurus YAML config.

    Taurus is BlazeMeter's own runner format, so this is what makes a plan
    portable: the same file runs under `bzt` locally, on BlazeMeter, or as the
    JMX we already emit. Only the HTTP parts convert - a spec carrying
    webdriver steps keeps those in the .jmx, which is noted in the output.
    """
    defaults = spec.get("defaults") or {}
    domain = defaults.get("domain") or ""
    protocol = defaults.get("protocol") or "https"
    address = ""
    if domain:
        address = domain if "://" in domain else "%s://%s" % (protocol, domain)

    scenario = {}
    if address:
        scenario["default-address"] = address
    if spec.get("variables"):
        scenario["variables"] = dict(spec["variables"])
    headers = dict(spec.get("headers") or {})
    if spec.get("auth"):
        headers.update(auth_header(spec["auth"]))
    if headers:
        scenario["headers"] = headers
    if defaults.get("connect_timeout"):
        scenario["timeout"] = _taurus_time(int(defaults["connect_timeout"]) // 1000)
    scenario["store-cache"] = bool(spec.get("cache", False))
    scenario["keepalive"] = True
    # `store-cookie` is the boolean; Taurus's `cookies` key takes a list of
    # cookie objects and blows up on a bool
    scenario["store-cookie"] = bool(spec.get("cookies", True))

    sources = []
    for csv in spec.get("csv") or []:
        entry = {"path": csv.get("file"), "loop": True}
        if csv.get("variables"):
            entry["variable-names"] = ",".join(csv["variables"])
        sources.append(entry)
    if sources:
        scenario["data-sources"] = sources

    groups = spec.get("thread_groups") or []
    requests = []
    for tg in groups:
        requests += _taurus_steps(tg.get("steps"))
    scenario["requests"] = requests

    name = re.sub(r"[^A-Za-z0-9_-]+", "-", spec.get("name") or "jmxgen").strip("-").lower()
    name = name or "jmxgen"

    execution = []
    for tg in groups or [{}]:
        item = {"scenario": name, "concurrency": _num_prop(tg.get("threads")) or 1}
        ramp = _num_prop(tg.get("ramp_up"))
        hold = _num_prop(tg.get("duration"))
        if ramp:
            item["ramp-up"] = _taurus_time(ramp)
        if hold:
            item["hold-for"] = _taurus_time(hold)
        elif tg.get("loops"):
            item["iterations"] = tg["loops"]
        execution.append(item)

    config = {
        "execution": execution,
        "scenarios": {name: scenario},
        "reporting": [{"module": "console"}, {"module": "final-stats"}],
    }
    return config


def dump_taurus(spec):
    """Taurus YAML as text, ready to write or hand to `bzt`."""
    config = spec_to_taurus(spec)
    header = ("# Generated by jmxgen - run with:  bzt this-file.yml\n"
              "# Or upload it to BlazeMeter / run the .jmx on HyperExecute.\n")
    wd = any(s.get("kind") == "webdriver"
             for tg in spec.get("thread_groups") or []
             for s in _walk_steps(tg.get("steps") or []))
    if wd:
        header += ("# NOTE: this plan has browser steps, which Taurus JMeter\n"
                   "#       scenarios cannot express - they stay in the .jmx.\n")
    if _have_yaml():
        import yaml
        return header + yaml.safe_dump(config, sort_keys=False, default_flow_style=False)
    return header + json.dumps(config, indent=2)


def _walk_steps(steps):
    """Every leaf step, transactions flattened."""
    for step in steps or []:
        if step.get("transaction"):
            for inner in _walk_steps(step.get("steps")):
                yield inner
        else:
            yield step


# --------------------------------------------------------------------------
# HAR -> spec
# --------------------------------------------------------------------------


SKIP_TYPES = re.compile(
    r"image/|font/|text/css|javascript|octet-stream", re.I)


# --------------------------------------------------------------------------
# HAR -> spec  (page-aware, filtered, with automatic correlation)
# --------------------------------------------------------------------------

_CORRELATE_SKIP = re.compile(
    r"^(mozilla|text/|application/|https?://|gzip|deflate|utf-8|no-cache|keep-alive)", re.I)


# --------------------------------------------------------------------------
# correlation rule packs
#
# A rule says: "a request field named like THIS carries a server-issued value,
# and here is exactly how to pull it back out of the response." Rules beat pure
# heuristics because they are deterministic and, when one fires, explainable.
# --------------------------------------------------------------------------

CORRELATION_RULES = [
    {"name": "asp.net-viewstate", "fields": [r"^__VIEWSTATE$"],
     "extract": {"type": "boundary", "left": 'id="__VIEWSTATE" value="', "right": '"'},
     "confidence": "high"},
    {"name": "asp.net-viewstategenerator", "fields": [r"^__VIEWSTATEGENERATOR$"],
     "extract": {"type": "boundary", "left": 'id="__VIEWSTATEGENERATOR" value="', "right": '"'},
     "confidence": "high"},
    {"name": "asp.net-eventvalidation", "fields": [r"^__EVENTVALIDATION$"],
     "extract": {"type": "boundary", "left": 'id="__EVENTVALIDATION" value="', "right": '"'},
     "confidence": "high"},
    {"name": "jsf-viewstate", "fields": [r"^javax\.faces\.ViewState$", r"^jakarta\.faces\.ViewState$"],
     "extract": {"type": "regex",
                 "query": r'name="javax\.faces\.ViewState"[^>]*value="([^"]+)"'},
     "confidence": "high"},
    {"name": "rails-csrf", "fields": [r"^authenticity_token$"],
     "extract": {"type": "regex",
                 "query": r'name="authenticity_token"[^>]*value="([^"]+)"'},
     "confidence": "high"},
    {"name": "django-csrf", "fields": [r"^csrfmiddlewaretoken$"],
     "extract": {"type": "regex",
                 "query": r'name="csrfmiddlewaretoken"[^>]*value="([^"]+)"'},
     "confidence": "high"},
    {"name": "laravel-csrf", "fields": [r"^_token$"],
     "extract": {"type": "regex", "query": r'name="_token"[^>]*value="([^"]+)"'},
     "confidence": "high"},
    {"name": "spring-csrf", "fields": [r"^_csrf$", r"^X-CSRF-TOKEN$"],
     "extract": {"type": "regex", "query": r'name="_csrf"[^>]*value="([^"]+)"'},
     "confidence": "high"},
    {"name": "saml-response", "fields": [r"^SAMLResponse$", r"^SAMLRequest$"],
     "extract": {"type": "regex", "query": r'name="SAMLResponse" value="([^"]+)"'},
     "confidence": "high"},
    {"name": "saml-relaystate", "fields": [r"^RelayState$"],
     "extract": {"type": "regex", "query": r'name="RelayState" value="([^"]+)"'},
     "confidence": "medium"},
    {"name": "oauth-code", "fields": [r"^code$"], "header_only": True,
     "confidence": "high"},
    {"name": "oauth-state", "fields": [r"^state$", r"^nonce$"], "confidence": "medium"},
    {"name": "bearer-token", "fields": [r"^Authorization$", r"access_token", r"id_token",
                                        r"^token$", r"^jwt$"],
     "confidence": "high"},
    {"name": "session-id", "fields": [r"session", r"^sid$", r"^JSESSIONID$",
                                      r"^PHPSESSID$", r"^ASP\.NET_SessionId$"],
     "confidence": "medium"},
]


def load_rules(path=None):
    rules = list(CORRELATION_RULES)
    if path:
        extra = _load_json_or_yaml(path)
        if isinstance(extra, dict):
            extra = extra.get("rules") or []
        rules = list(extra) + rules          # user rules win
    return rules


def _match_rule(label, rules):
    for rule in rules:
        for pattern in rule.get("fields", []):
            if re.search(pattern, label or "", re.I):
                return rule
    return None


def _response_headers(entry):
    out = []
    for h in (entry.get("response") or {}).get("headers") or []:
        out.append(("%s: %s" % (h.get("name", ""), h.get("value", ""))))
    return "\n".join(out)


# path segments that are structure, never an identifier
_UUIDISH = re.compile(r"[0-9a-f]{8,}(?:-[0-9a-f]{4,}){0,4}", re.I)

_PATH_NOISE = {"api", "rest", "public", "internal", "graphql", "oauth", "auth",
               "login", "logout", "token", "session", "v1", "v2", "v3", "v4"}


def _path_segments(url):
    """(label, value) for path segments that could be a server-issued id.

    The label is taken from the segment before it, so /orders/ord_123 is named
    after what it identifies - order_id - rather than something anonymous.
    """
    try:
        path = urllib.parse.urlsplit(url).path
    except Exception:
        return []
    parts = [p for p in path.split("/") if p]
    out = []
    for i, seg in enumerate(parts):
        low = seg.lower()
        if low in _PATH_NOISE or re.fullmatch(r"v\d+", low):
            continue
        parent = parts[i - 1].lower() if i else ""
        if not parent or parent in _PATH_NOISE:
            continue                           # /something - no collection to name it after
        # a numeric id says nothing on its own, so it only qualifies when the
        # segment before it reads like a collection: /products/101, /orders/42
        if not re.fullmatch(r"[a-z][a-z_-]{2,}s", parent):
            if len(seg) < 6:
                continue
        elif not re.fullmatch(r"[\w-]{1,64}", seg):
            continue
        # /carts/add and /products/search are routes, not ids. A real identifier
        # carries a digit or is a uuid; a static sub-route almost never does.
        if not any(c.isdigit() for c in seg) and not _UUIDISH.fullmatch(seg):
            continue
        noun = parent[:-1] if parent.endswith("s") and len(parent) > 3 else parent
        out.append(("%s_id" % re.sub(r"\W+", "_", noun), seg))
    return out


def _candidate_values(req):
    """(source_label, value) pairs from a HAR request that might be dynamic."""
    out = []
    for q in req.get("queryString") or []:
        out.append((q.get("name") or "q", q.get("value") or ""))
    pd = req.get("postData") or {}
    for p in pd.get("params") or []:
        out.append((p.get("name") or "p", p.get("value") or ""))
    text = pd.get("text") or ""
    if text:
        try:
            data = json.loads(text)

            def walk(node, prefix=""):
                if isinstance(node, dict):
                    for k, v in node.items():
                        walk(v, k)
                elif isinstance(node, list):
                    for v in node[:5]:
                        walk(v, prefix)
                elif isinstance(node, str):
                    out.append((prefix or "body", node))
            walk(data)
        except Exception:
            for m in re.finditer(r"([\w.\-]+)=([^&\s]{8,})", text):
                out.append((m.group(1), m.group(2)))
    # REST APIs carry the id in the path - /orders/ord_4b7e2a19 - which is the
    # single most common chaining pattern there is. Without this the id stays
    # hardcoded and every virtual user hammers the same one record.
    for label, seg in _path_segments(req.get("url") or ""):
        out.append((label, seg))
    for h in req.get("headers") or []:
        hname = (h.get("name") or "").lower()
        if hname in ("authorization", "x-csrf-token", "x-xsrf-token", "x-auth-token",
                     "x-api-key", "x-session-id"):
            value = h.get("value") or ""
            # "Bearer eyJ..." - correlate the credential, not the scheme word
            m = re.match(r"^(Bearer|Basic|Token|JWT)\s+(\S+)$", value, re.I)
            out.append((h["name"], m.group(2) if m else value))
    return out


SECRETISH = re.compile(
    r"token|auth|session|sess|csrf|xsrf|sid|nonce|jwt|ticket|secret|key|"
    r"state|code|signature|hash|guid|uuid", re.I)


def _is_dynamic(value, label=""):
    """Does this look like a server-issued token rather than a constant?

    Kept deliberately strict: a false correlation turns a stable literal (a
    version, a page size, a slug) into a variable that can break on replay."""
    v = (value or "").strip()
    if not (8 <= len(v) <= 512):
        return False
    if _CORRELATE_SKIP.match(v):
        return False
    if re.fullmatch(r"[\d.,_-]+", v):          # versions, ids, numbers: "2024.11.0"
        return False
    if not re.fullmatch(r"[\w\-.:+/=%~]{8,512}", v):
        return False
    has_digit = any(c.isdigit() for c in v)
    has_alpha = any(c.isalpha() for c in v)
    named_secret = bool(SECRETISH.search(label or ""))
    if not (has_digit and has_alpha) and not named_secret:
        return False
    if named_secret:                            # the field name is the strong signal
        return True
    # otherwise a short value must look like hex/base64 noise to qualify
    return len(v) >= 12 or bool(re.fullmatch(r"[0-9a-f]{8,}", v, re.I))


def _extractor_from_rule(rule, var, body, value):
    """Build the extractor a rule prescribes, checking it would actually match."""
    tmpl = dict(rule.get("extract") or {})
    if not tmpl:
        return None
    ex = dict(tmpl)
    ex["var"] = var
    if ex.get("type") == "boundary":
        left, right = ex.get("left", ""), ex.get("right", "")
        if left and left in body:
            return ex
        return None
    if ex.get("type") == "regex":
        try:
            m = re.search(ex["query"], body, re.S)
        except re.error:
            return None
        if m and (m.group(1) == value or value in m.group(0)):
            return ex
    return None


def _extractor_from_body(value, body, label):
    """No rule matched - infer from where the value sits in the response."""
    var = re.sub(r"\W+", "_", label).strip("_").upper()[:40] or "CORR"
    m = re.search(r'"([\w.\-]{1,60})"\s*:\s*"%s"' % re.escape(value), body)
    if m:
        var = re.sub(r"\W+", "_", m.group(1)).strip("_").upper()[:40] or var
        return {"type": "json", "var": var, "query": "$..%s" % m.group(1)}, var
    idx = body.find(value)
    if idx < 0:
        return None, None
    left = body[max(0, idx - 40):idx][-25:].split("\n")[-1]
    right = body[idx + len(value):idx + len(value) + 40][:25].split("\n")[0]
    if not left or not right:
        return None, None
    return {"type": "boundary", "var": var, "left": left, "right": right}, var


def _extractor_from_headers(value, headers_text, label):
    """The value came back in a response header - Location, Set-Cookie, custom."""
    var = re.sub(r"\W+", "_", label).strip("_").upper()[:40] or "CORR"
    idx = headers_text.find(value)
    if idx < 0:
        return None, None
    left = headers_text[max(0, idx - 40):idx][-25:].split("\n")[-1]
    if not left:
        return None, None
    ex = {"type": "regex", "var": var, "scope_field": "true",
          "query": "%s([^&;\\s\"]+)" % re.escape(left), "template": "$1$"}
    return ex, var


# How far back a value may have come from, counted in kept requests.
#
# Two reasons for a horizon, and the second matters more than the first.
#
# Cost: without one, every candidate value is compared against every response
# before it, so an hour-long recording is quadratic. Measured in the extension,
# 3,000 requests correlate in 3.7s and 6,000 in 36.8s - the same session twice
# the length takes ten times as long, and 12,000 is minutes.
#
# Correctness: a token that first appears forty minutes and two thousand
# requests earlier is almost certainly a coincidence, not a source. Wiring it
# up produces a plan that looks correlated and is wrong, which is worse than
# one that admits it found nothing. Real sources are near: a login, then the
# calls that use it.
CORRELATION_HORIZON = 300


def _correlate(entries, kept, steps_by_entry, limit=60, rules=None,
               horizon=CORRELATION_HORIZON):
    """Wire values a later request sends back to the response that produced them.

    Returns provenance for every correlation so the result is reviewable rather
    than a silent rewrite. `horizon` bounds how far back a source may be, in
    kept requests; 0 or None searches the whole session."""
    rules = rules if rules is not None else CORRELATION_RULES
    sources = []
    for idx in kept:
        entry = entries[idx]
        content = (entry.get("response") or {}).get("content") or {}
        sources.append((idx, content.get("text") or "", _response_headers(entry)))

    # position within `sources`, so the search can start at the nearest
    # preceding response rather than walking from the beginning of the session
    pos_of = {src[0]: p for p, src in enumerate(sources)}

    found, made = [], {}
    for idx in kept:
        step = steps_by_entry.get(idx)
        if step is None:
            continue
        for label, value in _candidate_values(entries[idx].get("request") or {}):
            if value in made or len(found) >= limit:
                continue
            rule = _match_rule(label, rules)
            from_path = label.endswith("_id") and value in (entries[idx].get("request")
                                                            or {}).get("url", "")
            if not rule and not _is_dynamic(value, label) and not from_path:
                continue
            if rule and not _is_dynamic(value, label) and len(value) < 4:
                continue

            src_idx = src_body = src_headers = None
            where = None
            # backwards from the request that used it: the nearest response
            # carrying the value is the one that produced it
            here = pos_of.get(idx, len(sources))
            stop = 0 if not horizon else max(0, here - horizon)
            for p in range(here - 1, stop - 1, -1):
                j, body, headers = sources[p]
                if body and value in body:
                    src_idx, src_body, src_headers, where = j, body, headers, "body"
                    break
                if headers and value in headers:
                    src_idx, src_body, src_headers, where = j, body, headers, "headers"
                    break
            if src_idx is None:
                continue

            src_step = steps_by_entry.get(src_idx)
            if src_step is None:
                continue

            var = re.sub(r"\W+", "_", label).strip("_").upper()[:40] or "CORR"
            ex = None
            if rule and where == "body":
                ex = _extractor_from_rule(rule, var, src_body, value)
            if ex is None:
                if where == "headers":
                    ex, var = _extractor_from_headers(value, src_headers, label)
                else:
                    ex, var = _extractor_from_body(value, src_body, label)
            if not ex:
                continue

            if any(e.get("var") == ex["var"] for e in src_step.get("extract", [])):
                ex["var"] = ex["var"] + "_%d" % len(found)
            var = ex["var"]

            src_step.setdefault("extract", []).append(ex)
            made[value] = var
            found.append({
                "var": var,
                "field": label,
                "value": value,
                "value_preview": value[:24],
                "rule": (rule or {}).get("name", "heuristic"),
                "confidence": (rule or {}).get("confidence", "low"),
                "found_in": where,
                "source_step": src_step.get("name", "?"),
                "used_in": step.get("name", "?"),
            })

    if made:
        def substitute(node_):
            if isinstance(node_, dict):
                return {k: substitute(v) for k, v in node_.items()}
            if isinstance(node_, list):
                return [substitute(v) for v in node_]
            if isinstance(node_, str):
                for value, var in made.items():
                    if value in node_:
                        node_ = node_.replace(value, "${%s}" % var)
                return node_
            return node_
        for idx in kept:
            step = steps_by_entry.get(idx)
            if step is None:
                continue
            keep_extract = step.pop("extract", None)
            keep_name = step.pop("name", None)
            step.update(substitute(step))
            if keep_extract is not None:
                step["extract"] = keep_extract
            if keep_name is not None:
                # JMeter resolves ${VARS} inside sampler labels at RUNTIME, so a
                # label carrying one fragments the report into a row per id.
                # Name it after the shape instead: "GET /objects/{id}".
                step["name"] = _stable_label(keep_name, made)
    return found


def _stable_label(name, made):
    """Replace correlated literals in a sampler label with a shape placeholder."""
    for value, var in made.items():
        if value and value in name:
            name = name.replace(value, "{%s}" % var.lower())
    return name


API_TYPES = {"xhr", "fetch", "websocket", "eventsource"}
ASSET_TYPES = {"script", "stylesheet", "image", "font", "media", "manifest",
               "texttrack", "other", "preflight", "ping", "csp_violation_report"}
API_MIME = re.compile(r"json|xml|x-protobuf|graphql|x-www-form-urlencoded", re.I)
WRITE_METHODS = {"POST", "PUT", "PATCH", "DELETE"}


def _is_api_call(entry, url, method, mime, rtype):
    """Is this a service call rather than a page resource?

    Chrome tells us the resource type outright; when it is missing (a converted
    or hand-made HAR) fall back to method, content type and path shape."""
    if rtype:
        if rtype in API_TYPES:
            return True
        if rtype in ASSET_TYPES:
            return False
        if rtype == "document":
            return bool(API_MIME.search(mime or ""))
    if method in WRITE_METHODS:
        return True
    if API_MIME.search(mime or ""):
        return not STATIC_EXT.search(urllib.parse.urlsplit(url).path or "")
    return bool(re.search(r"/(api|rest|graphql|v\d+)(/|$)", url, re.I))


def _parse_iso(ts):
    if not ts:
        return None
    try:
        return datetime.datetime.fromisoformat(ts.replace("Z", "+00:00"))
    except (ValueError, AttributeError):
        return None


def _apply_real_think_times(groups, order, cap_ms):
    """Turn recorded start-time gaps into think time on the preceding step.

    Only gaps big enough to be a human pause are kept; the sub-second bursts a
    page fires on load are not think time, they are parallel resource loading."""
    prev_step, prev_at = None, None
    for gname in order:
        for step in groups[gname]:
            at = _parse_iso(step.pop("_started", None))
            if prev_step is not None and prev_at is not None and at is not None:
                gap = int((at - prev_at).total_seconds() * 1000)
                if 300 <= gap <= cap_ms:
                    prev_step["think_time"] = gap
            if at is not None:
                prev_step, prev_at = step, at
    if prev_step is not None:
        prev_step.pop("_started", None)


def uncorrelate(spec, correlations, reject_vars):
    """Undo specific correlations - put the recorded literal back, drop the extractor.

    A reviewer rejecting a correlation must get the plan they would have had if it
    never fired, not a plan with a dangling variable."""
    reject = {v for v in reject_vars}
    by_var = {c["var"]: c for c in correlations if c["var"] in reject}
    if not by_var:
        return spec, 0
    reverted = [0]

    def walk(node_):
        if isinstance(node_, dict):
            out = {}
            for k, v in node_.items():
                if k == "extract" and isinstance(v, list):
                    kept = [e for e in v if e.get("var") not in by_var]
                    if kept:
                        out[k] = kept
                    continue
                out[k] = walk(v)
            return out
        if isinstance(node_, list):
            return [walk(v) for v in node_]
        if isinstance(node_, str):
            for var, c in by_var.items():
                token = "${%s}" % var
                if token in node_:
                    node_ = node_.replace(token, c.get("value", token))
                    reverted[0] += 1
            return node_
        return node_

    return walk(spec), reverted[0]


def har_to_spec(har_path, include=None, exclude=None, keep_static=False, name=None,
                pages=True, drop_third_party=True, keep_trackers=False,
                correlate=True, think_time=None, mode="auto", methods=None,
                real_think_time=False, max_think_ms=30000, rules=None):
    har = json.load(open(har_path, "r", encoding="utf-8-sig"))
    log = har.get("log", {})
    entries = log.get("entries", [])
    page_titles = {p.get("id"): (p.get("title") or p.get("id"))
                   for p in (log.get("pages") or [])}

    hosts, origins = {}, {}
    for e in entries:
        u_ = urllib.parse.urlsplit(e.get("request", {}).get("url", ""))
        if u_.hostname:
            hosts[u_.hostname] = hosts.get(u_.hostname, 0) + 1
            key = (u_.hostname, u_.scheme or "https", str(u_.port or ""))
            origins[key] = origins.get(key, 0) + 1
    # Pick the system under test by REGISTRABLE DOMAIN, not by hostname: a site's
    # CDN subdomain usually out-counts its www host, and a recording that starts
    # mid-session may contain no page document at all - only third-party pixels,
    # some of which serve html and would otherwise win.
    roots = {}
    for e in entries:
        h_ = urllib.parse.urlsplit(e.get("request", {}).get("url", "")).hostname
        if not h_ or _looks_like_tracker(h_):
            continue
        r_ = _reg_domain(h_)
        roots[r_] = roots.get(r_, 0) + 1
    if roots:
        keep_root = max(roots, key=roots.get)
    else:
        keep_root = _reg_domain(max(hosts, key=hosts.get)) if hosts else ""

    in_root = {h: n for h, n in hosts.items()
               if h == keep_root or (keep_root and h.endswith("." + keep_root))}
    doc_host = ""
    for e in entries:
        mime = ((e.get("response") or {}).get("content") or {}).get("mimeType", "")
        h_ = urllib.parse.urlsplit(e.get("request", {}).get("url", "")).hostname
        if "html" in (mime or "").lower() and h_ in in_root:
            doc_host = h_
            break
    www = next((h for h in in_root if h.startswith("www.")), "")
    top_host = (doc_host or www
                or (max(in_root, key=in_root.get) if in_root else "")
                or (max(hosts, key=hosts.get) if hosts else ""))
    # the recorded scheme and port for that host - not every site is https on 443
    top_origin = max([o for o in origins if o[0] == top_host] or [(top_host, "https", "")],
                     key=lambda o: origins.get(o, 0))
    top_scheme, top_port = top_origin[1], top_origin[2]

    kept, steps_by_entry, groups, order = [], {}, {}, []
    user_agents = {}
    want_methods = {m.strip().upper() for m in (methods or "").split(",") if m.strip()}

    # A recording contains every hop of a redirect chain, but JMeter samplers follow
    # redirects themselves - keeping the targets too would request each one twice.
    # Keep the head of the chain, drop what following it would fetch anyway.
    redirect_targets = set()
    for e in entries:
        resp = e.get("response") or {}
        status = int(resp.get("status") or 0)
        if 300 <= status < 400:
            loc = resp.get("redirectURL") or ""
            if not loc:
                for h in resp.get("headers") or []:
                    if (h.get("name") or "").lower() == "location":
                        loc = h.get("value") or ""
                        break
            if loc:
                redirect_targets.add(
                    urllib.parse.urljoin(e.get("request", {}).get("url", ""), loc))
    skipped = {"static": 0, "third_party": 0, "filtered": 0, "preflight": 0}

    for idx, e in enumerate(entries):
        req = e.get("request") or {}
        url = req.get("url") or ""
        if not url.startswith("http"):
            continue
        u = urllib.parse.urlsplit(url)
        host = u.hostname or ""
        if include and not re.search(include, url):
            skipped["filtered"] += 1
            continue
        if exclude and re.search(exclude, url):
            skipped["filtered"] += 1
            continue
        mime = ((e.get("response") or {}).get("content") or {}).get("mimeType", "")
        rtype = (e.get("_resourceType")
                 or ((e.get("_jmxgen") or {}).get("type") or "")).lower()
        method_ = (req.get("method") or "GET").upper()

        if want_methods and method_ not in want_methods:
            skipped["filtered"] += 1
            continue

        # A CORS preflight belongs to the browser, not to the test. The browser
        # sends OPTIONS before a cross-origin call because its security model
        # requires it; JMeter has no such model and will never send one, so a
        # preflight sampler measures a request that would not happen and
        # doubles the apparent request count of every cross-origin API.
        if method_ == "OPTIONS" and _is_cors_preflight(req):
            skipped["preflight"] += 1
            continue

        if mode == "api":
            # service calls only - the plan a backend team wants
            if not _is_api_call(e, url, method_, mime, rtype):
                skipped["static"] += 1
                continue
        elif mode == "web":
            pass                      # everything the browser fetched
        elif not keep_static and (SKIP_TYPES.search(mime or "")
                                  or STATIC_EXT.search(u.path or "")):
            skipped["static"] += 1
            continue
        if drop_third_party and host and not (host == keep_root or host.endswith("." + keep_root)):
            if not keep_trackers or _looks_like_tracker(host):
                skipped["third_party"] += 1
                continue
        if not keep_trackers and NOISE_PATHS.search(u.path or ""):
            skipped["third_party"] += 1      # first-party telemetry, still not the app
            continue

        step = {"name": "%s %s" % (req.get("method", "GET"), (u.path or "/")[:70]),
                "method": (req.get("method") or "GET").upper(),
                "path": (u.path or "/") + (("?" + u.query) if u.query else "")}
        if host and host != top_host:
            step["domain"], step["protocol"] = host, u.scheme
            if u.port:
                step["port"] = str(u.port)
        elif (u.scheme or "https") != top_scheme or str(u.port or "") != top_port:
            step["protocol"] = u.scheme
            step["port"] = str(u.port or "")

        hdrs = {}
        for h in req.get("headers") or []:
            n = (h.get("name") or "")
            if n.lower() == "user-agent":
                # one UA for the whole plan, not one per sampler
                user_agents[h.get("value", "")] = user_agents.get(h.get("value", ""), 0) + 1
                continue
            if n.startswith(":") or n.lower() in (
                    "host", "content-length", "cookie", "connection", "accept-encoding",
                    "sec-fetch-dest", "sec-fetch-mode", "sec-fetch-site", "sec-fetch-user",
                    "sec-ch-ua", "sec-ch-ua-mobile", "sec-ch-ua-platform", "upgrade-insecure-requests"):
                continue
            hdrs[n] = h.get("value", "")
        if hdrs:
            step["headers"] = hdrs

        pd = req.get("postData") or {}
        if pd.get("params"):
            step["params"] = {p.get("name"): p.get("value", "") for p in pd["params"] if p.get("name")}
        elif pd.get("text"):
            step["body"] = pd["text"]
        if url in redirect_targets:
            skipped["filtered"] += 1
            continue
        code = str((e.get("response") or {}).get("status") or "")
        # a 3xx sampler reports the FINAL code once JMeter follows it, so asserting
        # the recorded 302 would fail every time
        if code.startswith("2"):
            step["assert"] = [{"field": "code", "match": "equals", "pattern": code}]
        if think_time is not None:
            step["think_time"] = think_time
        elif real_think_time:
            # the gap the real user left between this request and the previous one,
            # which is what makes replayed load look like traffic instead of a flood
            step["_started"] = e.get("startedDateTime") or ""

        # steps authored by hand in the jmxgen Chrome recorder ride along in the
        # HAR as `_jmxgen` fields; the spec allows custom underscore keys
        ann = e.get("_jmxgen") or {}
        if ann.get("skip"):
            skipped["filtered"] += 1
            continue
        if ann.get("name"):
            step["name"] = str(ann["name"])[:110]
        if ann.get("assert"):
            step["assert"] = list(ann["assert"]) + step.get("assert", [])
        if ann.get("extract"):
            step["extract"] = list(ann["extract"]) + step.get("extract", [])
        if ann.get("transaction"):
            e["pageref"] = ann["transaction"]
        if ann.get("manual"):
            step["name"] = "[manual] " + step["name"]

        if not pages:
            gname = "Recorded"
        else:
            gname = (ann.get("transaction")
                     or page_titles.get(e.get("pageref"))
                     or e.get("pageref") or "Recorded")
        gname = re.sub(r"\s+", " ", str(gname))[:70]
        if gname not in groups:
            groups[gname] = []
            order.append(gname)
        groups[gname].append(step)
        kept.append(idx)
        steps_by_entry[idx] = step
        if ann.get("pause_after_ms"):
            groups[gname].append({"pause": int(ann["pause_after_ms"]),
                                  "name": "Pause after %s" % step["name"][:40]})

    if real_think_time:
        _apply_real_think_times(groups, order, max_think_ms)

    correlated = (_correlate(entries, kept, steps_by_entry, rules=rules)
                  if correlate else [])
    plan_ann = (log.get("_jmxgen") or {}) if isinstance(log.get("_jmxgen"), dict) else {}

    variables = {}
    plan_headers = {"Accept": "*/*"}
    if user_agents:
        ua = max(user_agents, key=user_agents.get)
        # a recorded headless UA would advertise the load generator as a bot
        ua = ua.replace("HeadlessChrome", "Chrome").replace("Headless", "")
        plan_headers["User-Agent"] = ua
    spec = {
        "name": name or ("%s load test" % (top_host or os.path.basename(har_path))),
        "platform": "local",
        "comments": "Authored by jmxgen from %s: %d of %d recorded requests kept"
                    % (os.path.basename(har_path), len(kept), len(entries)),
        "variables": variables,
        "defaults": {"protocol": top_scheme, "domain": top_host, "port": top_port,
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": plan_headers,
        "cookies": True,
        "cache": True,
        "thread_groups": [{
            "name": "Recorded Journey",
            "threads": int(plan_ann.get("threads", 10)),
            "ramp_up": int(plan_ann.get("ramp_up", 30)),
            "duration": int(plan_ann.get("duration", 300)),
            "steps": [{"transaction": g, "steps": groups[g]} for g in order],
        }],
    }
    if plan_ann.get("name") and not name:
        spec["name"] = plan_ann["name"]

    # The recorder rides the browser steps along in the HAR's own extension
    # field. They become a second thread group: one user, running the journey in
    # a real browser, next to the protocol group that carries the load.
    recorded = ((log.get("_jmxgen") or {}).get("actions")) or []
    if recorded:
        spec["thread_groups"].append({
            "name": "Browser journey",
            "threads": 1, "ramp_up": 1, "loops": 1,
            "steps": actions_to_steps(recorded),
        })
        spec.setdefault("webdriver", {"headless": True})

    return spec, {"kept": len(kept), "total": len(entries), "skipped": skipped,
                  "correlated": correlated, "pages": len(order),
                  "actions": len(recorded)}


# --------------------------------------------------------------------------
# validate
# --------------------------------------------------------------------------


def validate(path):
    text = open(path, "r", encoding="utf-8", errors="replace").read()
    issues, notes = [], []

    def count(pat):
        return len(re.findall(pat, text))

    n_tree = count(r"ViewResultsFullVisualizer|ViewResultsTree")
    if n_tree:
        issues.append("%d result-storing listener(s) (View Results Tree/Full Visualizer) "
                      "- top OOM cause, remove them" % n_tree)
    n_agg = count(r'guiclass="StatVisualizer"|guiclass="StatGraphVisualizer"')
    if n_agg:
        issues.append("%d Aggregate Report/Graph listener(s) - hold per-sample data in heap" % n_agg)

    scheduler_on = "ThreadGroup.scheduler\">true" in text
    if re.search(r'LoopController\.loops">-1<', text) and not scheduler_on:
        issues.append("loops=-1 with no scheduler+duration -> unbounded run")

    has_http = "<HTTPSamplerProxy " in text
    if has_http and not re.search(r'HTTPSampler\.connect_timeout">\d', text):
        issues.append("no connect_timeout set anywhere - hung requests hold threads/buffers")
    if has_http and not re.search(r'HTTPSampler\.response_timeout">\d', text):
        issues.append("no response_timeout set anywhere")

    n_disabled = count(r'enabled="false"')
    if n_disabled:
        issues.append("%d disabled element(s) - still parsed into heap, delete them" % n_disabled)

    n_bsh = count(r"BeanShell")
    if n_bsh:
        issues.append("%d BeanShell element(s) - interpreted, high CPU; use JSR223+Groovy" % n_bsh)

    if re.search(r'HTTPSampler\.image_parser">true', text):
        notes.append("'Retrieve All Embedded Resources' enabled on some samplers")
    if re.search(r'TestPlan\.functional_mode">true', text):
        issues.append("Functional Test Mode ON - forces response data into heap")
    if re.search(r"<responseData>true</responseData>", text):
        issues.append("a listener is saving response data - large heap/disk cost")
    if re.search(r'limitMaxThreadNumber">false', text):
        issues.append("Parallel Controller without 'Limit max thread number' - unbounded fan-out")

    size_mb = os.path.getsize(path) / 1024.0 / 1024.0
    n_samplers = count(r"<HTTPSamplerProxy ")
    if size_mb > 5:
        issues.append("plan is %.1f MB (%d HTTP samplers) - whole tree loads into heap; split it"
                      % (size_mb, n_samplers))
    else:
        notes.append("%.2f MB, %d HTTP samplers" % (size_mb, n_samplers))

    print("== %s ==" % path)
    for n in notes:
        print("  .  %s" % n)
    for i in issues:
        print("  !  %s" % i)
    if not issues:
        print("  ok - no checklist violations found")
    return 1 if issues else 0


# --------------------------------------------------------------------------
# starter specs
# --------------------------------------------------------------------------

HTTP_SPEC = {
    "name": "My API Load Test",
    "platform": "hyperexecute",
    "variables": {"BASE_HOST": "api.example.com"},
    "defaults": {"protocol": "https", "domain": "${BASE_HOST}",
                 "connect_timeout": 5000, "response_timeout": 30000},
    "headers": {"Content-Type": "application/json", "Accept": "application/json"},
    "cookies": True,
    "csv": [{"file": "users.csv", "variables": ["username", "password"]}],
    "thread_groups": [{
        "name": "Load",
        "threads": 50,
        "ramp_up": 60,
        "duration": 600,
        "steps": [{
            "transaction": "Login and fetch profile",
            "steps": [
                {"name": "POST /login", "method": "POST", "path": "/login",
                 "body": {"user": "${username}", "pass": "${password}"},
                 "extract": [{"type": "json", "var": "TOKEN", "query": "$.token"}],
                 "assert": [{"field": "code", "match": "equals", "pattern": "200"}],
                 "think_time": {"min": 500, "max": 1500}},
                {"name": "GET /profile", "method": "GET", "path": "/profile",
                 "headers": {"Authorization": "Bearer ${TOKEN}"},
                 "assert": [{"field": "body", "match": "contains", "pattern": "email"}],
                 "think_time": 1000},
            ],
        }],
    }],
}

WD_SPEC = {
    "name": "Browser E2E Flow",
    "platform": "hyperexecute",
    "variables": {"BASE_URL": "https://example.com/login",
                  "USERNAME": "user", "PASSWORD": "pass", "TIMEOUT": "30"},
    "webdriver": {"headless": True},
    "thread_groups": [{
        "name": "UI Flow",
        "threads": 1, "ramp_up": 1, "loops": 1,
        "steps": [{
            "transaction": "Login",
            "steps": [
                {"type": "webdriver", "name": "Open login page",
                 "actions": [{"do": "open", "url": "${BASE_URL}"},
                             {"do": "wait_for", "xpath": "//input[@name='username']"}],
                 "think_time": 500},
                {"type": "webdriver", "name": "Enter credentials and submit",
                 "actions": [{"do": "type", "xpath": "//input[@name='username']", "text": "${USERNAME}"},
                             {"do": "type", "xpath": "//input[@name='password']", "text": "${PASSWORD}"},
                             {"do": "click", "xpath": "//button[@type='submit']"},
                             {"do": "assert_text", "text": "Dashboard"}],
                 "think_time": 1000},
            ],
        }],
    }],
}


# --------------------------------------------------------------------------
# URL -> spec  (crawl a live site and author a plan from what's there)
# --------------------------------------------------------------------------

import urllib.request
import urllib.parse
from html.parser import HTMLParser


class _PageParser(HTMLParser):
    """Pull title, forms (+inputs) and same-page links out of an HTML page."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = ""
        self._in_title = False
        self.forms = []
        self.links = []
        self._form = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "title":
            self._in_title = True
        elif tag == "form":
            self._form = {"action": a.get("action", ""),
                          "method": (a.get("method") or "GET").upper(),
                          "enctype": a.get("enctype", ""),
                          "inputs": []}
        elif tag in ("input", "select", "textarea") and self._form is not None:
            name = a.get("name")
            if name and (a.get("type") or "text").lower() not in ("submit", "button", "reset", "image"):
                self._form["inputs"].append(
                    {"name": name, "value": a.get("value", ""),
                     "type": (a.get("type") or "text").lower()})
        elif tag == "a":
            href = a.get("href")
            if href:
                self.links.append(href)

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        elif tag == "form" and self._form is not None:
            self.forms.append(self._form)
            self._form = None

    def handle_data(self, data):
        if self._in_title:
            self.title += data.strip()


FETCH_PROXY = {}          # {"http": "http://host:port", "https": ...} when --proxy is used


def _opener():
    """A urlopen-compatible callable, routed through --proxy when one is set."""
    if not FETCH_PROXY:
        return urllib.request.urlopen
    return urllib.request.build_opener(urllib.request.ProxyHandler(FETCH_PROXY)).open


def _fetch(url, timeout=20, user_agent="Mozilla/5.0 (jmxgen)"):
    req = urllib.request.Request(url, headers={"User-Agent": user_agent,
                                               "Accept": "text/html,*/*"})
    with _opener()(req, timeout=timeout) as r:
        ctype = r.headers.get("Content-Type", "")
        raw = r.read(3_000_000)
    charset = "utf-8"
    if "charset=" in ctype:
        charset = ctype.split("charset=")[-1].split(";")[0].strip() or "utf-8"
    return raw.decode(charset, errors="replace"), ctype


def _placeholder(inp):
    """Sensible ${VAR} / literal for a form field, so the plan is fillable."""
    n, t = inp["name"], inp["type"]
    if inp["value"]:
        return inp["value"]          # hidden tokens, defaults - keep as recorded
    low = n.lower()
    if t == "password" or "pass" in low:
        return "${PASSWORD}"
    if "user" in low or "email" in low or "login" in low:
        return "${USERNAME}"
    if t == "email":
        return "${EMAIL}"
    if t in ("checkbox", "radio"):
        return "on"
    return "${%s}" % re.sub(r"\W+", "_", n).upper()


def url_to_spec(start_url, depth=1, max_pages=8, name=None, follow_links=True,
                timeout=20, think_time=1000):
    """Crawl from start_url and author a plan: page GETs + form submissions."""
    parts = urllib.parse.urlsplit(start_url)
    if not parts.scheme:
        start_url = "https://" + start_url
        parts = urllib.parse.urlsplit(start_url)
    base_domain = parts.hostname
    protocol = parts.scheme
    port = str(parts.port) if parts.port else ""

    seen, queue, steps = set(), [(start_url, 0)], []
    variables = {}
    pages_done = 0
    errors = []

    while queue and pages_done < max_pages:
        url, d = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        try:
            html, ctype = _fetch(url, timeout)
        except Exception as exc:                       # unreachable page - record, keep going
            errors.append("%s -> %s" % (url, exc))
            continue
        pages_done += 1
        if "html" not in ctype.lower():
            continue

        p = _PageParser()
        try:
            p.feed(html)
        except Exception:
            pass

        page_path = urllib.parse.urlsplit(url).path or "/"
        query = urllib.parse.urlsplit(url).query
        label = (p.title or page_path)[:60].strip() or page_path
        page_steps = [{
            "name": "GET %s" % page_path,
            "method": "GET",
            "path": page_path + (("?" + query) if query else ""),
            "assert": [{"field": "code", "match": "equals", "pattern": "200"}],
            "think_time": think_time,
        }]

        # hidden fields (csrf/viewstate) become boundary extractors so the POST works
        for form in p.forms:
            for inp in form["inputs"]:
                if inp["type"] == "hidden":
                    var = re.sub(r"\W+", "_", inp["name"]).upper()
                    page_steps[0].setdefault("extract", []).append({
                        "type": "boundary", "var": var,
                        "left": 'name="%s" value="' % inp["name"], "right": '"',
                        "default": inp["value"] or "NOT_FOUND",
                    })

        for i, form in enumerate(p.forms, 1):
            action = urllib.parse.urljoin(url, form["action"] or url)
            ap = urllib.parse.urlsplit(action)
            if ap.hostname and ap.hostname != base_domain:
                continue
            params = {}
            for inp in form["inputs"]:
                if inp["type"] == "hidden":
                    params[inp["name"]] = "${%s}" % re.sub(r"\W+", "_", inp["name"]).upper()
                else:
                    val = _placeholder(inp)
                    params[inp["name"]] = val
                    m = re.fullmatch(r"\$\{(\w+)\}", val)
                    if m:
                        variables.setdefault(m.group(1), "CHANGE_ME")
            step = {
                "name": "%s %s (form %d)" % (form["method"], ap.path or "/", i),
                "method": form["method"],
                "path": (ap.path or "/") + (("?" + ap.query) if ap.query else ""),
                "params": params,
                "assert": [{"field": "code", "match": "equals", "pattern": "200"}],
                "think_time": think_time,
            }
            if "multipart" in (form["enctype"] or ""):
                step["multipart"] = True
            page_steps.append(step)

        steps.append({"transaction": label, "steps": page_steps})

        if follow_links and d < depth:
            for href in p.links:
                nxt = urllib.parse.urljoin(url, href)
                np_ = urllib.parse.urlsplit(nxt)
                if np_.scheme not in ("http", "https") or np_.hostname != base_domain:
                    continue
                nxt = urllib.parse.urlunsplit((np_.scheme, np_.netloc, np_.path, np_.query, ""))
                if nxt not in seen:
                    queue.append((nxt, d + 1))

    spec = {
        "name": name or ("%s load test" % base_domain),
        "platform": "local",
        "comments": "Authored by jmxgen from %s (%d page(s) crawled)" % (start_url, pages_done),
        "variables": variables,
        "defaults": {"protocol": protocol, "domain": base_domain, "port": port,
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": {"User-Agent": "Mozilla/5.0 (JMeter)",
                    "Accept": "text/html,application/xhtml+xml,application/json;q=0.9,*/*;q=0.8"},
        "cookies": True,
        "thread_groups": [{
            "name": "User Journey",
            "threads": 10, "ramp_up": 30, "duration": 300,
            "steps": steps,
        }],
    }
    return spec, errors


# --------------------------------------------------------------------------
# Excel / CSV -> spec
# --------------------------------------------------------------------------

STEP_COLUMNS = ["transaction", "name", "method", "url", "path", "headers", "params",
                "body", "extract", "assert", "think_time", "enabled", "thread_group"]

CONFIG_KEYS_HELP = """name, platform, protocol, domain, port, base_url, threads,
ramp_up, duration, loops, results_file, connect_timeout, response_timeout,
header.<Name>, var.<NAME>, csv.file, csv.variables"""


def _split_pairs(text, pair_sep=";", kv_sep=":"):
    out = {}
    for chunk in str(text).replace("\n", pair_sep).split(pair_sep):
        chunk = chunk.strip()
        if not chunk:
            continue
        if kv_sep not in chunk:
            continue
        k, v = chunk.split(kv_sep, 1)
        out[k.strip()] = v.strip()
    return out


def _parse_extract(text):
    """'TOKEN=json:$.token; SID=regex:sid=([^;]+); C=boundary:left|right'"""
    out = []
    for chunk in re.split(r";|\n", str(text)):
        chunk = chunk.strip()
        if not chunk or "=" not in chunk:
            continue
        var, rest = chunk.split("=", 1)
        var, rest = var.strip(), rest.strip()
        kind, _, expr = rest.partition(":")
        kind = kind.strip().lower() or "json"
        expr = expr.strip()
        if kind == "boundary":
            left, _, right = expr.partition("|")
            out.append({"type": "boundary", "var": var, "left": left, "right": right})
        elif kind in ("regex", "regexp"):
            out.append({"type": "regex", "var": var, "query": expr})
        else:
            out.append({"type": "json", "var": var, "query": expr or rest})
    return out


def _parse_assert(text):
    """'code=200; body contains email; body matches ^\\{.*\\}$'"""
    out = []
    for chunk in re.split(r";|\n", str(text)):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"^(code|body|message|headers)\s*(=|==|contains|matches|equals|substring)\s*(.+)$",
                     chunk, re.I)
        if m:
            field, op, pat = m.group(1).lower(), m.group(2).lower(), m.group(3).strip()
            match = {"=": "equals", "==": "equals"}.get(op, op)
            out.append({"field": field, "match": match, "pattern": pat})
        else:
            out.append({"field": "body", "match": "contains", "pattern": chunk})
    return out


def _parse_think(text):
    t = str(text).strip()
    m = re.match(r"^(\d+)\s*[-to]+\s*(\d+)$", t)
    if m:
        return {"min": int(m.group(1)), "max": int(m.group(2))}
    return int(float(t))


def _rows_from_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl not installed. Run: pip install openpyxl   (or export the sheet as .csv)")
    wb = openpyxl.load_workbook(path, data_only=True)
    config, steps = {}, []
    names = {n.lower(): n for n in wb.sheetnames}

    if "config" in names:
        for row in wb[names["config"]].iter_rows(values_only=True):
            if not row or row[0] in (None, ""):
                continue
            key = str(row[0]).strip()
            if key.lower() in ("key", "setting"):
                continue
            val = row[1] if len(row) > 1 else None
            if val not in (None, ""):
                config[key] = val

    sheet = wb[names.get("steps", wb.sheetnames[0])]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return config, steps
    header = [str(c).strip().lower().replace(" ", "_") if c else "" for c in rows[0]]
    for row in rows[1:]:
        if not row or all(c in (None, "") for c in row):
            continue
        steps.append({header[i]: row[i] for i in range(min(len(header), len(row)))
                      if header[i] and row[i] not in (None, "")})
    return config, steps


def _rows_from_csv(path):
    import csv as _csv
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rdr = _csv.DictReader(fh)
        rows = [{(k or "").strip().lower().replace(" ", "_"): v
                 for k, v in r.items() if v not in (None, "")} for r in rdr]
    return {}, rows


def sheet_to_spec(path, name=None):
    config, rows = (_rows_from_csv(path) if path.lower().endswith(".csv")
                    else _rows_from_xlsx(path))

    def cfg(key, default=None):
        for k, v in config.items():
            if k.strip().lower() == key:
                return v
        return default

    base_url = cfg("base_url")
    protocol, domain, port = cfg("protocol", "https"), cfg("domain", ""), cfg("port", "")
    if base_url:
        u = urllib.parse.urlsplit(str(base_url) if "//" in str(base_url) else "https://" + str(base_url))
        protocol, domain = u.scheme or protocol, u.hostname or domain
        port = str(u.port) if u.port else port

    headers = {k.split(".", 1)[1]: v for k, v in config.items() if k.lower().startswith("header.")}
    variables = {k.split(".", 1)[1]: v for k, v in config.items() if k.lower().startswith("var.")}

    groups = {}
    order = []
    for row in rows:
        if str(row.get("enabled", "yes")).strip().lower() in ("no", "false", "0"):
            continue
        gname = str(row.get("thread_group") or cfg("thread_group_name") or "Load").strip()
        if gname not in groups:
            groups[gname] = []
            order.append(gname)

        url = row.get("url") or row.get("path") or "/"
        url = str(url).strip()
        step = {"method": str(row.get("method", "GET")).strip().upper()}
        if "//" in url:
            u = urllib.parse.urlsplit(url)
            step["path"] = (u.path or "/") + (("?" + u.query) if u.query else "")
            if u.hostname and u.hostname != domain:
                step["domain"] = u.hostname
                step["protocol"] = u.scheme
                if u.port:
                    step["port"] = str(u.port)
            if not domain:
                domain, protocol = u.hostname, u.scheme
        else:
            step["path"] = url if url.startswith("/") or url.startswith("$") else "/" + url

        step["name"] = str(row.get("name") or "%s %s" % (step["method"], step["path"]))[:120]
        if row.get("headers"):
            step["headers"] = _split_pairs(row["headers"])
        if row.get("params"):
            step["params"] = _split_pairs(row["params"], kv_sep="=")
        if row.get("body"):
            step["body"] = str(row["body"])
        if row.get("extract"):
            step["extract"] = _parse_extract(row["extract"])
        if row.get("assert"):
            step["assert"] = _parse_assert(row["assert"])
        if row.get("think_time") not in (None, ""):
            step["think_time"] = _parse_think(row["think_time"])

        txn = str(row.get("transaction") or "").strip()
        bucket = groups[gname]
        if txn:
            if bucket and bucket[-1].get("transaction") == txn:
                bucket[-1]["steps"].append(step)
            else:
                bucket.append({"transaction": txn, "steps": [step]})
        else:
            bucket.append(step)

    def num(key, default):
        v = cfg(key, default)
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return default

    duration = cfg("duration")
    tgs = []
    for gname in order:
        tg = {"name": gname,
              "threads": num("threads", 1),
              "ramp_up": num("ramp_up", 1)}
        if duration not in (None, ""):
            tg["duration"] = int(float(duration))
        else:
            tg["loops"] = num("loops", 1)
        tg["steps"] = groups[gname]
        tgs.append(tg)

    csv_file = cfg("csv.file")
    spec = {
        "name": str(name or cfg("name") or os.path.basename(path).rsplit(".", 1)[0]),
        "platform": str(cfg("platform", "local")),
        "comments": "Authored by jmxgen from %s" % os.path.basename(path),
        "variables": variables,
        "defaults": {"protocol": protocol, "domain": domain or "", "port": str(port or ""),
                     "connect_timeout": num("connect_timeout", 5000),
                     "response_timeout": num("response_timeout", 30000)},
        "headers": headers or {"Accept": "*/*"},
        "cookies": True,
        "thread_groups": tgs,
    }
    if cfg("results_file"):
        spec["results_file"] = str(cfg("results_file"))
    if csv_file:
        variables_col = cfg("csv.variables", "")
        spec["csv"] = [{"file": str(csv_file),
                        "variables": [v.strip() for v in str(variables_col).split(",") if v.strip()]}]
    return spec


def write_template(path):
    """Emit a filled-in Excel (or CSV) template the user can edit and feed back in."""
    example = [
        ["Login", "POST login", "POST", "/api/login", "Content-Type: application/json", "",
         '{"user":"${USERNAME}","pass":"${PASSWORD}"}', "TOKEN=json:$.token", "code=200",
         "500-1500", "yes", "Load"],
        ["Login", "GET profile", "GET", "/api/profile", "Authorization: Bearer ${TOKEN}", "", "",
         "", "code=200; body contains email", "1000", "yes", "Load"],
        ["Search", "GET search", "GET", "/api/search", "", "q=shoes; page=1", "", "", "code=200",
         "1000", "yes", "Load"],
    ]
    config_rows = [
        ["key", "value"],
        ["name", "My Load Test"],
        ["platform", "hyperexecute"],
        ["base_url", "https://api.example.com"],
        ["threads", 50],
        ["ramp_up", 60],
        ["duration", 600],
        ["connect_timeout", 5000],
        ["response_timeout", 30000],
        ["header.Accept", "application/json"],
        ["var.USERNAME", "testuser"],
        ["var.PASSWORD", "changeme"],
        ["csv.file", "users.csv"],
        ["csv.variables", "username,password"],
    ]
    if path.lower().endswith(".csv"):
        import csv as _csv
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = _csv.writer(fh)
            w.writerow(STEP_COLUMNS[:1] + ["name", "method", "url", "headers", "params", "body",
                                           "extract", "assert", "think_time", "enabled", "thread_group"])
            for r in example:
                w.writerow(r)
        return path
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        sys.exit("openpyxl not installed. Run: pip install openpyxl   (or use a .csv template)")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "steps"
    head = ["transaction", "name", "method", "url", "headers", "params", "body",
            "extract", "assert", "think_time", "enabled", "thread_group"]
    ws.append(head)
    for r in example:
        ws.append(r)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    widths = [14, 22, 8, 34, 34, 22, 38, 26, 30, 12, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    cs = wb.create_sheet("config")
    for r in config_rows:
        cs.append(r)
    for c in cs[1]:
        c.font, c.fill = bold, fill
    cs.column_dimensions["A"].width = 24
    cs.column_dimensions["B"].width = 40

    hs = wb.create_sheet("help")
    for line in [
        ["steps sheet columns"],
        ["transaction", "groups consecutive rows into one Transaction Controller (optional)"],
        ["name", "sampler label shown in reports"],
        ["method", "GET / POST / PUT / DELETE / PATCH"],
        ["url", "full URL or a path like /api/login (paths use the config base_url)"],
        ["headers", "Name: value; Name2: value2"],
        ["params", "a=1; b=2   (form/query parameters)"],
        ["body", "raw request body - JSON, XML, anything; use ${VAR} freely"],
        ["extract", "VAR=json:$.token; VAR2=regex:sid=([^;]+); VAR3=boundary:left|right"],
        ["assert", "code=200; body contains email; body matches ^\\{.*\\}$"],
        ["think_time", "1000  or  500-1500 (random range), in ms"],
        ["enabled", "no / false / 0 skips the row"],
        ["thread_group", "rows with the same value land in the same Thread Group"],
        [""],
        ["config sheet keys"],
        [CONFIG_KEYS_HELP.replace("\n", " ")],
    ]:
        hs.append(line)
    hs.column_dimensions["A"].width = 16
    hs.column_dimensions["B"].width = 90
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# verify - is the generated XML actually a valid, loadable JMeter plan?
# --------------------------------------------------------------------------


def _walk_hashtree(ht, errors, trail):
    kids = list(ht)
    i = 0
    while i < len(kids):
        el = kids[i]
        if el.tag == "hashTree":
            errors.append("%s: stray <hashTree> with no element before it" % trail)
            i += 1
            continue
        label = el.get("testname") or el.tag
        nxt = kids[i + 1] if i + 1 < len(kids) else None
        if nxt is None or nxt.tag != "hashTree":
            errors.append("%s > <%s> \"%s\" is not followed by a <hashTree> "
                          "(JMeter will not load this)" % (trail, el.tag, label))
            i += 1
            continue
        _walk_hashtree(nxt, errors, "%s > %s" % (trail, label))
        i += 2


def verify(path, deep=False, quiet=False):
    """Structural + semantic validity. Returns (errors, warnings)."""
    import xml.etree.ElementTree as ET

    errors, warnings = [], []
    try:
        tree = ET.parse(path)
    except ET.ParseError as exc:
        errs = ["not well-formed XML: %s" % exc]
        hint = None
        if "invalid character" in str(exc):
            hint = ("a control character is embedded in an element value (often a binary "
                    "or gzipped payload pasted into a request). JMeter's XStream reader is "
                    "lenient and may still load it, but no standard XML tool will - "
                    "base64 or externalise that payload")
        if not quiet:
            print("== verify %s ==" % path)
            print("  X  %s" % errs[0])
            if hint:
                print("     hint: %s" % hint)
        loads = False
        if deep:
            deep_errs = _deep_load_check(path)
            loads = not deep_errs
            errs += deep_errs
        if not quiet:
            print("  %s" % ("LOADS IN JMETER, but not standard-XML valid - "
                            "no other XML tool will read it"
                            if loads else "INVALID - fix the X lines above"))
        return errs, []

    root = tree.getroot()
    if root.tag != "jmeterTestPlan":
        errors.append("root element is <%s>, expected <jmeterTestPlan>" % root.tag)
    tops = list(root)
    if len(tops) != 1 or tops[0].tag != "hashTree":
        errors.append("<jmeterTestPlan> must contain exactly one <hashTree>")
    else:
        _walk_hashtree(tops[0], errors, "TestPlan")

    text = open(path, "r", encoding="utf-8", errors="replace").read()

    if root.find(".//TestPlan") is None:
        errors.append("no <TestPlan> element")
    tgs = (root.findall(".//ThreadGroup") + root.findall(".//SetupThreadGroup")
           + root.findall(".//PostThreadGroup")
           + root.findall(".//kg.apc.jmeter.threads.UltimateThreadGroup")
           + root.findall(".//%s.arrivals.ArrivalsThreadGroup" % CASUTG)
           + root.findall(".//%s.concurrency.ConcurrencyThreadGroup" % CASUTG)
           + root.findall(".//%s.arrivals.FreeFormArrivalsThreadGroup" % CASUTG))
    if not tgs:
        errors.append("no Thread Group - nothing would run")
    samplers = [e for e in root.iter() if e.tag.endswith("Sampler")
                or e.tag in ("HTTPSamplerProxy", "JSR223Sampler", "DebugSampler")]
    if not samplers:
        errors.append("no samplers - the plan makes no requests")

    def _prop(el, name, deep_search=False):
        """A JMeter property may be written as string/int/long/boolProp."""
        for kind in ("stringProp", "intProp", "longProp", "boolProp"):
            q = ".//%s[@name='%s']" % (kind, name) if deep_search \
                else "%s[@name='%s']" % (kind, name)
            v = el.findtext(q)
            if v is not None:
                return v
        return None

    for tg in tgs:
        n = _prop(tg, "ThreadGroup.num_threads")
        if n is not None and str(n).strip() and not re.match(r"^\s*(\d+|\$\{.*\})\s*$", str(n)):
            errors.append("Thread Group '%s' has non-numeric thread count: %r"
                          % (tg.get("testname"), n))
        sched = (_prop(tg, "ThreadGroup.scheduler") or "false") == "true"
        loops = _prop(tg, "LoopController.loops", deep_search=True)
        dur = _prop(tg, "ThreadGroup.duration") or ""
        if sched and not dur.strip():
            errors.append("Thread Group '%s' has the scheduler on but no duration"
                          % tg.get("testname"))
        if str(loops).strip() == "-1" and not sched:
            errors.append("Thread Group '%s' loops forever (-1) with no scheduler"
                          % tg.get("testname"))

    # plugin elements that must be installed on the runner
    plugins = sorted({e.tag for e in root.iter()
                      if e.tag.startswith(("com.googlecode.jmeter.plugins",
                                           "kg.apc.jmeter", "com.blazemeter"))})
    # naming the jar matters: "needs a plugin" sends people hunting, and on
    # HyperExecute the fix is simply to upload the jar with the plan
    PLUGIN_JARS = {
        "com.blazemeter.jmeter.threads": "jmeter-plugins-casutg (Custom Thread Groups)",
        "com.googlecode.jmeter.plugins.webdriver": "jmeter-plugins-webdriver",
        "kg.apc.jmeter.threads": "jmeter-plugins-casutg (Custom Thread Groups)",
        "kg.apc.jmeter": "jmeter-plugins-standard",
    }
    for p in plugins:
        jar = next((v for k, v in PLUGIN_JARS.items() if p.startswith(k)), None)
        if jar:
            warnings.append("needs %s on the runner - install it in JMeter's "
                            "lib/ext, or upload the jar alongside the plan (%s)"
                            % (jar, p.rsplit(".", 1)[-1]))
        else:
            warnings.append("needs plugin element on the runner: %s" % p)

    # referenced files
    base = os.path.dirname(os.path.abspath(path))
    for prop in root.iter("stringProp"):
        if prop.get("name") in ("filename", "TestPlan.user_define_classpath") and (prop.text or "").strip():
            f = prop.text.strip()
            if "${" in f or f.startswith("/dev/"):
                continue
            target = f if os.path.isabs(f) else os.path.join(base, f)
            parent_is_writer = False
            if prop.get("name") == "filename":
                # result files are written, not read - only data-set inputs must exist
                parent_is_writer = not any(
                    f == (c.text or "").strip()
                    for cds in root.iter("CSVDataSet") for c in cds.iter("stringProp")
                    if c.get("name") == "filename")
            if not parent_is_writer and not os.path.exists(target):
                warnings.append("referenced file not found: %s" % f)

    # variables used but never defined
    defined = set(re.findall(r'<stringProp name="Argument\.name">([^<]+)<', text))
    defined |= set(re.findall(r'<stringProp name="\w+PostProcessor\.referenceNames">([^<]+)<', text))
    defined |= set(re.findall(r'<stringProp name="\w*Extractor\.refname">([^<]+)<', text))
    for names in re.findall(r'<stringProp name="variableNames">([^<]*)<', text):
        defined |= {n.strip() for n in names.split(",") if n.strip()}
    defined |= set(re.findall(r"vars\.put\(\s*['\"](\w+)['\"]", text))
    scriptless = re.sub(r'<stringProp name="(?:script|WebDriverSampler\.script)">.*?</stringProp>',
                        "", text, flags=re.S)
    used = set(re.findall(r"\$\{([A-Za-z_]\w*)\}", scriptless))
    unknown = sorted(used - defined - {"__jexl3", "__P", "__V"})
    for u in unknown:
        warnings.append("variable ${%s} is used but never defined (typo, or set at runtime by -J/-D)" % u)

    if deep:
        errors += _deep_load_check(path)

    if not quiet:
        print("== verify %s ==" % path)
        print("  .  well-formed XML, %d thread group(s), %d sampler(s)" % (len(tgs), len(samplers)))
        for w in warnings:
            print("  ?  %s" % w)
        for e in errors:
            print("  X  %s" % e)
        print("  %s" % ("VALID - ready to run" if not errors else "INVALID - fix the X lines above"))
    return errors, warnings


REPLAY_COLLECTOR = """<ResultCollector guiclass="SimpleDataWriter" testclass="ResultCollector" testname="jmxgen replay" enabled="true">
  <boolProp name="ResultCollector.error_logging">false</boolProp>
  <objProp><name>saveConfig</name><value class="SampleSaveConfiguration">
    <time>true</time><latency>true</latency><timestamp>true</timestamp><success>true</success>
    <label>true</label><code>true</code><message>true</message><threadName>true</threadName>
    <dataType>false</dataType><encoding>false</encoding><assertions>true</assertions>
    <subresults>true</subresults><responseData>false</responseData><samplerData>true</samplerData>
    <xml>true</xml><fieldNames>true</fieldNames><responseHeaders>false</responseHeaders>
    <requestHeaders>true</requestHeaders><responseDataOnError>true</responseDataOnError>
    <saveAssertionResultsFailureMessage>true</saveAssertionResultsFailureMessage>
    <assertionsResultsToSave>0</assertionsResultsToSave><bytes>true</bytes><url>true</url>
  </value></objProp>
  <stringProp name="filename">%s</stringProp>
</ResultCollector>"""


def _prepare_replay_plan(text, jtl_path, ET):
    """One thread, one loop, no scheduler, plus a collector inside the Test Plan.

    Done on the parsed tree: splicing the collector in as text lands it outside
    the Test Plan's subtree, where it is out of scope and silently records nothing.
    """
    root = ET.fromstring(text)

    for tg in list(root.iter()):
        if not tg.tag.endswith("ThreadGroup"):
            continue
        def setprop(el, name, value, kinds=("stringProp", "intProp", "longProp")):
            for kind in kinds:
                node_ = el.find("%s[@name='%s']" % (kind, name))
                if node_ is not None:
                    node_.text = value
                    return
            new_ = ET.SubElement(el, "stringProp")
            new_.set("name", name)
            new_.text = value
        setprop(tg, "ThreadGroup.num_threads", "1")
        setprop(tg, "ThreadGroup.ramp_time", "1")
        setprop(tg, "ThreadGroup.duration", "")
        sched = tg.find("boolProp[@name='ThreadGroup.scheduler']")
        if sched is not None:
            sched.text = "false"
        for loops in tg.iter():
            if loops.get("name") == "LoopController.loops":
                loops.text = "1"
            if loops.get("name") == "LoopController.continue_forever":
                loops.text = "false"

    top = root.find("hashTree")
    plan_ht = top.find("hashTree") if top is not None else None
    if plan_ht is None:
        raise ValueError("plan has no Test Plan subtree")
    plan_ht.append(ET.fromstring(REPLAY_COLLECTOR % jtl_path))
    plan_ht.append(ET.Element("hashTree"))
    return '<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding="unicode")


UNRESOLVED = re.compile(r"NOT_FOUND|\$\{[A-Za-z_]\w*\}")


def replay(path, out_dir=None, timeout=300, quiet=False):
    """Run the plan once with a single user and report what a real run would hit.

    A generated plan can be perfectly valid XML and still fail on the first
    request because a correlation did not match. This is the check that catches
    it - before anyone spends cloud minutes."""
    import shutil
    import subprocess
    import tempfile
    import xml.etree.ElementTree as ET

    jmeter = shutil.which("jmeter")
    if not jmeter:
        raise ValueError("jmeter is not on PATH - cannot replay")

    work = out_dir or tempfile.mkdtemp(prefix="jmxgen-replay-")
    os.makedirs(work, exist_ok=True)
    jtl = os.path.join(work, "replay.jtl")

    text = _prepare_replay_plan(
        open(path, "r", encoding="utf-8", errors="replace").read(), jtl, ET)
    plan = os.path.join(work, "replay.jmx")
    open(plan, "w", encoding="utf-8").write(text)

    # The plan references its data files by relative name, so JMeter has to run
    # with the plan's own directory as the working directory - otherwise every
    # ${column} silently resolves to nothing and the whole replay fails at login.
    cwd = os.path.dirname(os.path.abspath(path)) or os.getcwd()

    log_path = os.path.join(work, "replay.log")
    try:
        proc = subprocess.run([jmeter, "-n", "-t", plan, "-j", log_path],
                              capture_output=True, text=True, timeout=timeout, cwd=cwd)
    except subprocess.TimeoutExpired:
        raise ValueError("replay exceeded %ds - is the plan time-boxed?" % timeout)

    if not os.path.exists(jtl):
        raise ValueError("replay produced no results (see %s)" % work)

    # which ${VARS} each sampler depends on - so a failure can name its likely cause.
    # The Authorization header usually lives in a HeaderManager inside the sampler's
    # hashTree, so the sampler element alone is not enough.
    plan_vars = {}
    try:
        plan_root = ET.parse(path).getroot()

        def scan(ht):
            kids = list(ht)
            i = 0
            while i < len(kids):
                el = kids[i]
                sub = kids[i + 1] if i + 1 < len(kids) and kids[i + 1].tag == "hashTree" else None
                if el.tag.endswith("Sampler") or el.tag == "HTTPSamplerProxy":
                    blob = ET.tostring(el, encoding="unicode")
                    if sub is not None:
                        blob += ET.tostring(sub, encoding="unicode")
                    refs = sorted({r for r in re.findall(r"\$\{([A-Za-z_]\w*)\}", blob)
                                   if not r.startswith("__")})
                    props = re.findall(r"\$\{__P\(([A-Za-z_]\w*)", blob)
                    refs = sorted(set(refs) | set(props))
                    if refs:
                        plan_vars[el.get("testname", "?")] = refs
                if sub is not None:
                    scan(sub)
                i += 2 if sub is not None else 1

        top = plan_root.find("hashTree")
        if top is not None:
            scan(top)
    except Exception:
        pass

    samples = []
    try:
        root = ET.parse(jtl).getroot()
    except ET.ParseError:
        # JMeter leaves the file unclosed if it was interrupted
        raw = open(jtl, "r", encoding="utf-8", errors="replace").read()
        root = ET.fromstring(raw + "</testResults>")

    def visit(node_):
        for el in node_:
            if el.tag in ("httpSample", "sample"):
                req = " ".join(filter(None, [
                    (el.findtext("samplerData") or ""),
                    (el.findtext("queryString") or ""),
                    # JMeter writes this as <requestHeader>, singular
                    (el.findtext("requestHeader") or el.findtext("requestHeaders") or ""),
                    (el.findtext("java.net.URL") or ""),
                ]))
                label = el.get("lb", "?")
                assertion = re.sub(
                    r"\s+", " ",
                    (el.findtext("assertionResult/failureMessage") or "")).strip()
                samples.append({
                    "label": label,
                    "code": el.get("rc", ""),
                    "success": el.get("s") == "true",
                    "message": re.sub(r"\s+", " ", el.get("rm", "")).strip(),
                    "assertion": assertion,
                    "unresolved": sorted(set(UNRESOLVED.findall(req))),
                    "depends_on": plan_vars.get(label, []),
                })
                visit(el)
    visit(root)

    failures = [x for x in samples if not x["success"]]
    unresolved = [x for x in samples if x["unresolved"]]

    # A run that produced nothing is not a pass, and the caller cannot tell the
    # difference from the sample list alone - both are empty. JMeter has
    # usually said why on stdout or in its log ("File users.csv must exist and
    # be readable" is the common one), so the reason is carried out with it
    # rather than left in a temp directory nobody will look in.
    reason = ""
    if not samples:
        blob = (proc.stdout or "") + "\n" + (proc.stderr or "")
        try:
            with open(log_path, "r", encoding="utf-8", errors="replace") as fh:
                blob += "\n" + fh.read()
        except OSError:
            pass
        # Rank by usefulness. JMeter logs INFO lines containing the word
        # "error" ("Thread will continue on error"), so a first-match scan
        # reports noise and hides the exception three lines below it.
        best = ""
        for line in blob.splitlines():
            text = re.sub(r"^\d{4}-\d\d-\d\d \S+ \S+ \S+: ", "", line.strip())
            low = text.lower()
            if "must exist" in low or "no such file" in low:
                best = text
                break                                   # the definitive one
            if not best and ("caused by:" in low or "exception:" in low
                             or low.startswith("error") or " error " in low
                             and "continue on error" not in low):
                best = text
        reason = best[:200]
        if not reason:
            reason = ("JMeter ran but produced no samples - the plan may have no "
                      "enabled thread group, or its data file is missing")

    if not quiet:
        print("== replay %s ==" % path)
        print("  ran %d sampler(s) as a single user" % len(samples))
        for x in unresolved:
            print("  !  %-38s sent %s - a correlation did not resolve"
                  % (x["label"][:38], ", ".join(x["unresolved"])))
        for x in failures:
            detail = x["assertion"] or x["message"]
            print("  X  %-38s %-4s %s" % (x["label"][:38], x["code"], detail[:70]))
            if x["depends_on"] and not x["unresolved"]:
                print("       depends on ${%s} - check those correlations resolved"
                      % ("}, ${".join(x["depends_on"])))
        if not samples:
            print("  X  nothing ran - %s" % reason)
        elif not failures and not unresolved:
            print("  PASS - every request succeeded and every variable resolved")
        else:
            print("  FAIL - %d failing sampler(s), %d with unresolved variables"
                  % (len(failures), len(unresolved)))
        print("  artifacts: %s" % work)
    return {"samples": samples, "failures": failures, "unresolved": unresolved,
            "ran": bool(samples), "reason": reason, "dir": work}


def _deep_load_check(path):
    """Ask the real JMeter to deserialize the tree, with every thread group disabled
    so nothing is actually executed."""
    import shutil
    import subprocess
    import tempfile

    jmeter = shutil.which("jmeter")
    if not jmeter:
        print("  ?  jmeter not on PATH - skipping deep load check")
        return []
    text = open(path, "r", encoding="utf-8", errors="replace").read()
    disabled = re.sub(r'(<(?:\w+\.)*\w*ThreadGroup\b[^>]*?)enabled="true"', r'\1enabled="false"', text)
    disabled = re.sub(r'(<(?:\w+\.)*\w*ThreadGroup\b(?![^>]*enabled=)[^>]*?)>', r'\1 enabled="false">', disabled)
    tmpdir = tempfile.mkdtemp(prefix="jmxgen-verify-")
    tmp = os.path.join(tmpdir, "check.jmx")
    open(tmp, "w", encoding="utf-8").write(disabled)
    try:
        proc = subprocess.run([jmeter, "-n", "-t", tmp, "-l", os.path.join(tmpdir, "o.jtl"),
                               "-j", os.path.join(tmpdir, "j.log")],
                              capture_output=True, text=True, timeout=180)
    except Exception as exc:
        print("  ?  deep check could not run: %s" % exc)
        return []
    out = (proc.stdout or "") + (proc.stderr or "")
    if "Created the tree successfully" in out:
        print("  .  JMeter loaded the tree successfully (deep check)")
        return []
    detail = [l for l in out.splitlines()
              if re.search(r"error|exception|invalid|cannot|could not", l, re.I)][:5]
    return ["JMeter could not load the plan: " + (" | ".join(detail) or out.strip()[-300:])]


# --------------------------------------------------------------------------
# shared helpers for the authoring adapters
# --------------------------------------------------------------------------

STATIC_EXT = re.compile(
    r"\.(js|mjs|css|png|jpe?g|gif|svg|webp|ico|woff2?|ttf|eot|map|mp4|webm|avif)(\?|$)", re.I)

NOISE_PATHS = re.compile(
    r"/cdn-cgi/(rum|speculation|challenge-platform|trace)|/__cf|/beacon(\.|/|$)|"
    r"/collect(\?|$)|/gtm\.js|/gtag/|/analytics/(event|collect)|/rum(\?|$)", re.I)

TRACKER_HINTS = (
    "google-analytics", "googletagmanager", "doubleclick", "gstatic", "googleapis",
    "facebook", "fbcdn", "snapchat", "tiktok", "bing.com", "bat.bing", "clarity.ms",
    "hotjar", "crazyegg", "onetrust", "cookielaw", "segment", "mixpanel", "amplitude",
    "newrelic", "nr-data", "sentry", "optimizely", "adsrvr", "adnxs", "criteo",
    "linkedin", "licdn", "twitter", "t.co", "pinterest", "cloudflareinsights",
    "hubspot", "intercom", "zendesk", "cdn.jsdelivr", "unpkg", "recaptcha",
)


def _is_cors_preflight(req):
    """A preflight is an OPTIONS carrying Access-Control-Request-*, which only a
    browser sends. Recognised by the header rather than the method alone, so a
    genuine OPTIONS endpoint someone means to test is still kept."""
    for h in req.get("headers") or []:
        if str(h.get("name", "")).lower().startswith("access-control-request-"):
            return True
    return False


def _is_third_party(host, keep_hosts):
    if not host:
        return False
    if any(host == k or host.endswith("." + k) for k in keep_hosts):
        return False
    return True


def _looks_like_tracker(host):
    h = (host or "").lower()
    return any(t in h for t in TRACKER_HINTS)


def _reg_domain(host):
    """Rough registrable domain: last two labels (good enough for grouping)."""
    parts = (host or "").split(".")
    return ".".join(parts[-2:]) if len(parts) >= 2 else host


def _var(name):
    return "${%s}" % re.sub(r"\W+", "_", str(name)).upper().strip("_")


def _load_json_or_yaml(path_or_url):
    """Accept a local file or an http(s) URL, JSON or YAML."""
    if re.match(r"^https?://", str(path_or_url)):
        text, _ = _fetch(path_or_url)
    else:
        text = open(path_or_url, "r", encoding="utf-8-sig").read()
    text = text.lstrip()
    if text.startswith("{") or text.startswith("["):
        return json.loads(text)
    if _have_yaml():
        import yaml
        return yaml.safe_load(text)
    return json.loads(text)          # will raise with a clear message


# --------------------------------------------------------------------------
# OpenAPI / Swagger -> spec
# --------------------------------------------------------------------------


def _resolve_ref(doc, ref, seen=None):
    if not ref.startswith("#/"):
        return {}
    node = doc
    for part in ref[2:].split("/"):
        part = part.replace("~1", "/").replace("~0", "~")
        if not isinstance(node, dict) or part not in node:
            return {}
        node = node[part]
    return node


def _example_from_schema(schema, doc, depth=0, name=""):
    """Synthesize a plausible value for a JSON schema."""
    if not isinstance(schema, dict) or depth > 6:
        return "value"
    if "$ref" in schema:
        return _example_from_schema(_resolve_ref(doc, schema["$ref"]), doc, depth + 1, name)
    for key in ("example", "default"):
        if key in schema:
            return schema[key]
    if "examples" in schema and isinstance(schema["examples"], list) and schema["examples"]:
        return schema["examples"][0]
    for comb in ("allOf", "oneOf", "anyOf"):
        if schema.get(comb):
            merged = {}
            for sub in schema[comb]:
                got = _example_from_schema(sub, doc, depth + 1, name)
                if isinstance(got, dict):
                    merged.update(got)
                else:
                    return got
            return merged
    if schema.get("enum"):
        return schema["enum"][0]
    t = schema.get("type")
    if t == "object" or ("properties" in schema and not t):
        out = {}
        props = schema.get("properties") or {}
        required = set(schema.get("required") or [])
        for k, v in props.items():
            if required and k not in required and len(out) >= 8:
                continue
            out[k] = _example_from_schema(v, doc, depth + 1, k)
        return out
    if t == "array":
        return [_example_from_schema(schema.get("items") or {}, doc, depth + 1, name)]
    if t in ("integer", "number"):
        return schema.get("minimum", 1)
    if t == "boolean":
        return True
    fmt = (schema.get("format") or "").lower()
    if fmt == "date-time":
        return "2026-01-01T00:00:00Z"
    if fmt == "date":
        return "2026-01-01"
    if fmt == "uuid":
        return "00000000-0000-0000-0000-000000000000"
    if fmt in ("email", "idn-email"):
        return "user@example.com"
    if fmt == "binary":
        return "BINARY"
    return _var(name) if name else "value"


def openapi_to_spec(src, name=None, auth=None, include=None, exclude=None,
                    server=None, think_time=500):
    doc = _load_json_or_yaml(src)
    swagger2 = "swagger" in doc and "openapi" not in doc

    if server:
        base = server
    elif swagger2:
        host = doc.get("host") or ""
        scheme = (doc.get("schemes") or ["https"])[0]
        base = "%s://%s%s" % (scheme, host, doc.get("basePath", "")) if host else ""
    else:
        servers = doc.get("servers") or []
        base = (servers[0].get("url") if servers else "") or ""
        for var, cfg in ((servers[0].get("variables") or {}) if servers else {}).items():
            base = base.replace("{%s}" % var, str(cfg.get("default", var)))
    u = urllib.parse.urlsplit(base if "//" in base else "https://" + base if base else "")
    protocol, domain = (u.scheme or "https"), (u.hostname or "")
    port = str(u.port) if u.port else ""
    base_path = (u.path or "").rstrip("/")

    variables = {}
    groups = {}
    order = []
    count = 0

    for raw_path, item in (doc.get("paths") or {}).items():
        if not isinstance(item, dict):
            continue
        common = item.get("parameters") or []
        for method, op in item.items():
            if method.lower() not in ("get", "post", "put", "patch", "delete", "head", "options"):
                continue
            if not isinstance(op, dict):
                continue
            label = op.get("operationId") or "%s %s" % (method.upper(), raw_path)
            hay = "%s %s %s" % (method, raw_path, label)
            if include and not re.search(include, hay, re.I):
                continue
            if exclude and re.search(exclude, hay, re.I):
                continue

            params = list(common) + list(op.get("parameters") or [])
            params = [_resolve_ref(doc, p["$ref"]) if "$ref" in p else p for p in params]

            path = raw_path
            query, headers = {}, {}
            for p in params:
                pname, loc = p.get("name"), p.get("in")
                if not pname:
                    continue
                schema = p.get("schema") or ({"type": p.get("type")} if p.get("type") else {})
                val = _example_from_schema(schema, doc, name=pname) if schema else _var(pname)
                if isinstance(val, (dict, list)):
                    val = _var(pname)
                if loc == "path":
                    variables.setdefault(re.sub(r"\W+", "_", pname).upper(), str(val))
                    path = path.replace("{%s}" % pname, _var(pname))
                elif loc == "query" and (p.get("required") or len(query) < 4):
                    query[pname] = str(val)
                elif loc == "header" and pname.lower() not in ("authorization", "content-type"):
                    headers[pname] = str(val)
                elif loc == "body":                       # swagger 2
                    body_schema = p.get("schema") or {}
                    op["_body_example"] = _example_from_schema(body_schema, doc)

            body = op.get("_body_example")
            ctype = None
            rb = op.get("requestBody")
            if rb:
                rb = _resolve_ref(doc, rb["$ref"]) if "$ref" in rb else rb
                content = rb.get("content") or {}
                for mt in ("application/json", "application/x-www-form-urlencoded",
                           "text/plain", "*/*"):
                    if mt in content:
                        ctype = mt
                        media = content[mt]
                        body = media.get("example")
                        if body is None and media.get("examples"):
                            first = list(media["examples"].values())[0]
                            body = (first or {}).get("value")
                        if body is None:
                            body = _example_from_schema(media.get("schema") or {}, doc)
                        break
                else:
                    if content:
                        ctype = list(content)[0]
                        body = _example_from_schema(
                            (content[ctype] or {}).get("schema") or {}, doc)

            step = {"name": label[:110], "method": method.upper(),
                    "path": base_path + path, "think_time": think_time}
            if query:
                step["params"] = query
            if headers:
                step["headers"] = headers
            if body is not None:
                if ctype == "application/x-www-form-urlencoded" and isinstance(body, dict):
                    step.setdefault("params", {}).update({k: str(v) for k, v in body.items()})
                else:
                    step["body"] = body
                    step.setdefault("headers", {}).setdefault(
                        "Content-Type", ctype or "application/json")
            if auth:
                step.setdefault("headers", {})["Authorization"] = auth

            ok = "200"
            for code in (op.get("responses") or {}):
                if str(code).startswith("2"):
                    ok = str(code)
                    break
            step["assert"] = [{"field": "code", "match": "equals", "pattern": ok}]

            tag = (op.get("tags") or ["API"])[0]
            if tag not in groups:
                groups[tag] = []
                order.append(tag)
            groups[tag].append(step)
            count += 1

    if auth and "${" in str(auth):
        for v in re.findall(r"\$\{(\w+)\}", str(auth)):
            variables.setdefault(v, "CHANGE_ME")

    title = (doc.get("info") or {}).get("title") or "API"
    spec = {
        "name": name or "%s load test" % title,
        "platform": "local",
        "comments": "Authored by jmxgen from %s (%d operation(s))" % (src, count),
        "variables": variables,
        "defaults": {"protocol": protocol, "domain": domain, "port": port,
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": {"Accept": "application/json"},
        "cookies": True,
        "thread_groups": [{
            "name": "API Load",
            "threads": 10, "ramp_up": 30, "duration": 300,
            "steps": [{"transaction": t, "steps": groups[t]} for t in order],
        }],
    }
    return spec, count


# --------------------------------------------------------------------------
# Postman collection -> spec
# --------------------------------------------------------------------------


def _pm_vars(text):
    """{{token}} -> ${token}"""
    if not isinstance(text, str):
        return text
    return re.sub(r"\{\{(\w+)\}\}", r"${\1}", text)


def _pm_url(url, known=None):
    if isinstance(url, str):
        raw = url
        query = {}
    else:
        raw = url.get("raw") or ""
        query = {q.get("key"): q.get("value", "") for q in (url.get("query") or [])
                 if q.get("key") and not q.get("disabled")}
    # a collection variable holding the base URL has to be resolved to a real host -
    # JMeter needs a concrete domain, not "${baseUrl}" with a scheme inside it
    for k, v in (known or {}).items():
        if v:
            raw = raw.replace("{{%s}}" % k, str(v))
    raw = _pm_vars(raw).split("#")[0]
    return raw, query


def _pm_tests_to_steps(events):
    """Turn common pm.* test lines into assertions and extractors."""
    asserts, extracts = [], []
    for ev in events or []:
        if ev.get("listen") != "test":
            continue
        script = "\n".join((ev.get("script") or {}).get("exec") or [])
        for code in re.findall(r"to\.have\.status\((\d{3})\)", script):
            asserts.append({"field": "code", "match": "equals", "pattern": code})
        for code in re.findall(r"response\.code\)\.to\.eql\((\d{3})\)", script):
            asserts.append({"field": "code", "match": "equals", "pattern": code})
        for var, path in re.findall(
                r"(?:pm\.)?(?:environment|collectionVariables|globals)\.set\(\s*['\"](\w+)['\"]\s*,\s*"
                r"(?:jsonData|response|pm\.response\.json\(\))((?:\.\w+|\[\d+\])+)", script):
            jsonpath = "$" + path
            extracts.append({"type": "json", "var": var, "query": jsonpath})
        for text in re.findall(r"to\.include\(\s*['\"]([^'\"]+)['\"]", script):
            asserts.append({"field": "body", "match": "contains", "pattern": text})
    return asserts, extracts


def postman_to_spec(path, name=None, think_time=500):
    doc = _load_json_or_yaml(path)
    info = doc.get("info") or {}
    variables = {}
    for v in doc.get("variable") or []:
        if v.get("key"):
            variables[v["key"]] = str(v.get("value", ""))

    coll_auth = doc.get("auth") or {}
    domains = {}
    count = [0]

    def auth_header(auth):
        auth = auth or coll_auth
        kind = (auth or {}).get("type")
        if kind == "bearer":
            items = auth.get("bearer") or []
            tok = next((i.get("value") for i in items if i.get("key") == "token"), "${TOKEN}")
            return {"Authorization": "Bearer %s" % _pm_vars(tok)}
        if kind == "basic":
            items = {i.get("key"): i.get("value") for i in (auth.get("basic") or [])}
            import base64
            u, p = _pm_vars(items.get("username", "")), _pm_vars(items.get("password", ""))
            if "${" in u + p:
                return {"Authorization": "Basic ${__base64Encode(%s:%s)}" % (u, p)}
            return {"Authorization": "Basic " + base64.b64encode(
                ("%s:%s" % (u, p)).encode()).decode()}
        if kind == "apikey":
            items = {i.get("key"): i.get("value") for i in (auth.get("apikey") or [])}
            return {_pm_vars(items.get("key", "X-API-Key")): _pm_vars(items.get("value", "${API_KEY}"))}
        return {}

    def convert(item):
        req = item.get("request")
        if isinstance(req, str):
            req = {"method": "GET", "url": req}
        raw, query = _pm_url(req.get("url") or "", known=variables)
        if not raw:
            return None
        if "//" not in raw:
            raw = "https://" + raw
        u = urllib.parse.urlsplit(raw)
        host = u.hostname or ""
        if host and "$" not in host:
            domains[host] = domains.get(host, 0) + 1
        step = {"name": (item.get("name") or "%s %s" % (req.get("method", "GET"), u.path))[:110],
                "method": (req.get("method") or "GET").upper(),
                "path": (u.path or "/") + (("?" + u.query) if u.query and not query else ""),
                "think_time": think_time}
        step["_host"] = host
        step["_scheme"] = u.scheme or "https"
        if u.port:
            step["_port"] = str(u.port)
        if query:
            step["params"] = {k: _pm_vars(str(v)) for k, v in query.items()}

        headers = {}
        for h in req.get("header") or []:
            if h.get("disabled") or not h.get("key"):
                continue
            headers[h["key"]] = _pm_vars(str(h.get("value", "")))
        headers.update(auth_header(item.get("auth") or req.get("auth")))
        if headers:
            step["headers"] = headers

        body = req.get("body") or {}
        mode = body.get("mode")
        if mode == "raw" and body.get("raw"):
            step["body"] = _pm_vars(body["raw"])
            lang = ((body.get("options") or {}).get("raw") or {}).get("language")
            step.setdefault("headers", {}).setdefault(
                "Content-Type", {"json": "application/json", "xml": "application/xml"}
                .get(lang, "application/json" if body["raw"].lstrip().startswith(("{", "["))
                     else "text/plain"))
        elif mode == "urlencoded":
            step.setdefault("params", {}).update(
                {p["key"]: _pm_vars(str(p.get("value", "")))
                 for p in body.get("urlencoded") or [] if p.get("key") and not p.get("disabled")})
        elif mode == "formdata":
            step["multipart"] = True
            step.setdefault("params", {}).update(
                {p["key"]: _pm_vars(str(p.get("value", "")))
                 for p in body.get("formdata") or [] if p.get("key") and not p.get("disabled")})
        elif mode == "graphql":
            gql = body.get("graphql") or {}
            step["body"] = json.dumps({"query": gql.get("query", ""),
                                       "variables": json.loads(gql.get("variables") or "{}")})
            step.setdefault("headers", {})["Content-Type"] = "application/json"

        asserts, extracts = _pm_tests_to_steps(item.get("event"))
        if asserts:
            step["assert"] = asserts
        if extracts:
            step["extract"] = extracts
        for v in re.findall(r"\$\{(\w+)\}", json.dumps(step)):
            variables.setdefault(v, "CHANGE_ME")
        count[0] += 1
        return step

    def walk(items, depth=0):
        out = []
        for item in items or []:
            if item.get("item") is not None:                 # folder
                kids = walk(item["item"], depth + 1)
                if kids:
                    out.append({"transaction": item.get("name") or "Folder", "steps": kids})
            else:
                step = convert(item)
                if step:
                    out.append(step)
        return out

    steps = walk(doc.get("item"))

    top = max(domains, key=domains.get) if domains else ""

    def strip(nodes):
        for n in nodes:
            if "transaction" in n:
                strip(n["steps"])
                continue
            host = n.pop("_host", "")
            scheme = n.pop("_scheme", "https")
            port = n.pop("_port", "")
            if host and host != top:
                n["domain"], n["protocol"] = host, scheme
                if port:
                    n["port"] = port
    strip(steps)

    return {
        "name": name or (info.get("name") or "Postman collection") + " load test",
        "platform": "local",
        "comments": "Authored by jmxgen from %s (%d request(s))" % (os.path.basename(str(path)), count[0]),
        "variables": variables,
        "defaults": {"protocol": "https", "domain": top, "port": "",
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": {"Accept": "*/*"},
        "cookies": True,
        "thread_groups": [{"name": "Collection Run", "threads": 10, "ramp_up": 30,
                           "duration": 300, "steps": steps}],
    }, count[0]


# --------------------------------------------------------------------------
# cURL -> spec
# --------------------------------------------------------------------------


def curl_to_spec(text, name=None, think_time=500):
    import shlex

    text = re.sub(r"\\\s*\n", " ", text)                 # join line continuations
    blocks = [b.strip() for b in re.split(r"(?m)^\s*(?=curl\b)", text) if b.strip().startswith("curl")]
    steps, domains, variables = [], {}, {}

    for block in blocks:
        try:
            argv = shlex.split(block)
        except ValueError:
            argv = block.split()
        url, method, headers, data, form, user, multipart = None, None, {}, [], {}, None, False
        i = 1
        while i < len(argv):
            a = argv[i]
            nxt = argv[i + 1] if i + 1 < len(argv) else ""
            if a in ("-X", "--request"):
                method = nxt.upper(); i += 2; continue
            if a in ("-H", "--header"):
                if ":" in nxt:
                    k, v = nxt.split(":", 1)
                    headers[k.strip()] = v.strip()
                i += 2; continue
            if a in ("-d", "--data", "--data-raw", "--data-binary", "--data-ascii"):
                data.append(nxt); i += 2; continue
            if a in ("--data-urlencode",):
                data.append(nxt); i += 2; continue
            if a in ("-F", "--form"):
                multipart = True
                if "=" in nxt:
                    k, v = nxt.split("=", 1)
                    form[k] = v
                i += 2; continue
            if a in ("-u", "--user"):
                user = nxt; i += 2; continue
            if a in ("--url",):
                url = nxt; i += 2; continue
            if a.startswith("-"):
                i += 1 + (1 if a in ("-e", "--referer", "-A", "--user-agent", "-b", "--cookie",
                                     "-o", "--output", "--connect-timeout", "-m", "--max-time",
                                     "--retry", "-w", "--write-out", "--cert", "--key") else 0)
                continue
            if not url and re.match(r"^(https?://|[\w.-]+\.\w)", a):
                url = a
            i += 1

        if not url:
            continue
        if "//" not in url:
            url = "https://" + url
        u = urllib.parse.urlsplit(url)
        domains[u.hostname or ""] = domains.get(u.hostname or "", 0) + 1
        if not method:
            method = "POST" if (data or form) else "GET"

        step = {"name": "%s %s" % (method, u.path or "/"),
                "method": method,
                "path": (u.path or "/") + (("?" + u.query) if u.query else ""),
                "think_time": think_time,
                "_host": u.hostname or "", "_scheme": u.scheme or "https",
                "_port": str(u.port) if u.port else ""}
        if headers:
            step["headers"] = headers
        if user:
            import base64
            step.setdefault("headers", {})["Authorization"] = "Basic " + base64.b64encode(
                user.encode()).decode()
        if form:
            step["params"] = form
            step["multipart"] = True
        elif data:
            joined = "&".join(data)
            ctype = (headers.get("Content-Type") or headers.get("content-type") or "").lower()
            if "x-www-form-urlencoded" in ctype or (
                    not ctype and re.fullmatch(r"[\w.\-%]+=[^&]*(&[\w.\-%]+=[^&]*)*", joined)):
                step["params"] = dict(urllib.parse.parse_qsl(joined, keep_blank_values=True))
            else:
                step["body"] = joined
                step.setdefault("headers", {}).setdefault(
                    "Content-Type", ctype or "application/json")
        step["assert"] = [{"field": "code", "match": "equals", "pattern": "200"}]
        steps.append(step)

    if not steps:
        raise ValueError("no curl command found in the input")

    top = max(domains, key=domains.get)
    scheme = steps[0]["_scheme"]
    port = steps[0]["_port"]
    for s in steps:
        h, sc, p = s.pop("_host"), s.pop("_scheme"), s.pop("_port")
        if h != top:
            s["domain"], s["protocol"] = h, sc
            if p:
                s["port"] = p
        for v in re.findall(r"\$\{(\w+)\}", json.dumps(s)):
            variables.setdefault(v, "CHANGE_ME")

    return {
        "name": name or "%s load test" % top,
        "platform": "local",
        "comments": "Authored by jmxgen from %d curl command(s)" % len(steps),
        "variables": variables,
        "defaults": {"protocol": scheme, "domain": top, "port": port,
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": {"Accept": "*/*"},
        "cookies": True,
        "thread_groups": [{"name": "Requests", "threads": 5, "ramp_up": 10, "loops": 1,
                           "steps": [{"transaction": "Flow", "steps": steps}]}],
    }, len(steps)


# --------------------------------------------------------------------------
# .jmx -> spec  (migration: import a plan you already have, then edit it)
# --------------------------------------------------------------------------


def _prop(el, name, ET=None):
    for kind in ("stringProp", "intProp", "longProp", "boolProp", "doubleProp"):
        node_ = el.find("%s[@name='%s']" % (kind, name))
        if node_ is not None:
            return (node_.text or "").strip()
    return ""


def _sampler_to_step(el, sub, ET):
    step = {"name": el.get("testname", "request"),
            "method": _prop(el, "HTTPSampler.method") or "GET",
            "path": _prop(el, "HTTPSampler.path") or "/"}
    for key, prop in (("domain", "HTTPSampler.domain"), ("port", "HTTPSampler.port"),
                      ("protocol", "HTTPSampler.protocol")):
        val = _prop(el, prop)
        if val:
            step[key] = val
    if _prop(el, "HTTPSampler.postBodyRaw") == "true":
        for arg in el.iter("elementProp"):
            body = arg.findtext("stringProp[@name='Argument.value']")
            if body:
                step["body"] = body
                break
    else:
        params = {}
        for arg in el.iter("elementProp"):
            n = arg.findtext("stringProp[@name='Argument.name']")
            v = arg.findtext("stringProp[@name='Argument.value']")
            if n:
                params[n] = v or ""
        if params:
            step["params"] = params

    if sub is None:
        return step
    for child in sub:
        tag = child.tag
        if tag == "HeaderManager":
            hdrs = {}
            for ep in child.iter("elementProp"):
                n = ep.findtext("stringProp[@name='Header.name']")
                v = ep.findtext("stringProp[@name='Header.value']")
                if n:
                    hdrs[n] = v or ""
            if hdrs:
                step["headers"] = hdrs
        elif tag == "JSONPostProcessor":
            step.setdefault("extract", []).append({
                "type": "json",
                "var": _prop(child, "JSONPostProcessor.referenceNames"),
                "query": _prop(child, "JSONPostProcessor.jsonPathExprs")})
        elif tag == "RegexExtractor":
            step.setdefault("extract", []).append({
                "type": "regex",
                "var": _prop(child, "RegexExtractor.refname"),
                "query": _prop(child, "RegexExtractor.regex")})
        elif tag == "BoundaryExtractor":
            step.setdefault("extract", []).append({
                "type": "boundary",
                "var": _prop(child, "BoundaryExtractor.refname"),
                "left": _prop(child, "BoundaryExtractor.lboundary"),
                "right": _prop(child, "BoundaryExtractor.rboundary")})
        elif tag == "ResponseAssertion":
            field = _prop(child, "Assertion.test_field")
            patterns = [c.text or "" for c in child.iter("stringProp")
                        if (c.get("name") or "").isdigit()]
            for pat in patterns:
                step.setdefault("assert", []).append({
                    "field": {"Assertion.response_code": "code",
                              "Assertion.response_data": "body"}.get(field, "body"),
                    "match": "equals" if field.endswith("code") else "contains",
                    "pattern": pat})
        elif tag == "ConstantTimer":
            step["think_time"] = int(_prop(child, "ConstantTimer.delay") or 0)
    return step


def _num_prop(value, fallback=1):
    """Read a numeric plan property, unwrapping ${__P(name,default)} to its default.

    Load values are emitted as JMeter properties so a plan can be re-run at any
    load with -J - which means anything reading them back has to cope with that.
    """
    if value is None:
        return fallback
    text = str(value).strip()
    m = re.match(r"^\$\{__P\(\s*[^,)]*\s*,\s*([^)]*?)\s*\)\}$", text)
    if m:
        text = m.group(1)
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return fallback


def jmx_to_spec(path, name=None):
    """Reverse a plan back into an editable spec - the migration path in."""
    import xml.etree.ElementTree as ET

    text, _blobs, _refs = _sanitize_xml(
        open(path, "r", encoding="utf-8", errors="replace").read())
    root = ET.fromstring(text)
    counts = {"samplers": 0, "raw": 0}

    def walk(ht):
        steps = []
        kids = list(ht)
        i = 0
        while i < len(kids):
            el = kids[i]
            sub = kids[i + 1] if i + 1 < len(kids) and kids[i + 1].tag == "hashTree" else None
            tag = el.tag
            if tag == "HTTPSamplerProxy":
                counts["samplers"] += 1
                steps.append(_sampler_to_step(el, sub, ET))
            elif tag == "TransactionController":
                steps.append({"transaction": el.get("testname", "Transaction"),
                              "steps": walk(sub) if sub is not None else []})
            elif "ParallelSampler" in tag:
                steps.append({"parallel": el.get("testname", "Parallel"),
                              "max_threads": int(_prop(el, "MAX_THREAD_NUMBER") or 6),
                              "steps": walk(sub) if sub is not None else []})
            elif tag == "IfController":
                steps.append({"if": _prop(el, "IfController.condition"),
                              "steps": walk(sub) if sub is not None else []})
            elif tag in ("LoopController", "OnceOnlyController", "WhileController"):
                key = {"LoopController": "loop", "OnceOnlyController": "once",
                       "WhileController": "while"}[tag]
                val = (_prop(el, "LoopController.loops") if tag == "LoopController"
                       else (_prop(el, "WhileController.condition") if tag == "WhileController"
                             else True))
                steps.append({key: val, "steps": walk(sub) if sub is not None else []})
            elif tag in ("ResultCollector", "HeaderManager", "CookieManager",
                         "CacheManager", "ConfigTestElement", "Arguments", "CSVDataSet"):
                pass                       # plan-level config, handled separately
            elif tag.endswith("Sampler") or "Controller" in tag or tag.endswith("Timer"):
                counts["raw"] += 1
                blob = ET.tostring(el, encoding="unicode")
                steps.append({"type": "raw", "xml": blob})
            i += 2 if sub is not None else 1
        return steps

    tgs = []
    for parent in root.iter("hashTree"):
        kids = list(parent)
        for j, el in enumerate(kids):
            if not el.tag.endswith("ThreadGroup"):
                continue
            sub = kids[j + 1] if j + 1 < len(kids) and kids[j + 1].tag == "hashTree" else None
            tg = {"name": el.get("testname", "Thread Group"),
                  "threads": _num_prop(_prop(el, "ThreadGroup.num_threads"), 1),
                  "ramp_up": _num_prop(_prop(el, "ThreadGroup.ramp_time"), 1)}
            dur = _prop(el, "ThreadGroup.duration")
            if _prop(el, "ThreadGroup.scheduler") == "true" and dur:
                tg["duration"] = _num_prop(dur, 0)
            else:
                loops = _prop(el, "LoopController.loops") or "1"
                tg["loops"] = int(loops) if str(loops).lstrip("-").isdigit() else 1
            tg["steps"] = walk(sub) if sub is not None else []
            tgs.append(tg)

    defaults, headers = {}, {}
    for el in root.iter("ConfigTestElement"):
        if el.get("guiclass") == "HttpDefaultsGui":
            for key, prop in (("protocol", "HTTPSampler.protocol"),
                              ("domain", "HTTPSampler.domain"),
                              ("port", "HTTPSampler.port"),
                              ("connect_timeout", "HTTPSampler.connect_timeout"),
                              ("response_timeout", "HTTPSampler.response_timeout")):
                val = _prop(el, prop)
                if val:
                    defaults[key] = val
    for el in root.iter("HeaderManager"):
        for ep in el.iter("elementProp"):
            n = ep.findtext("stringProp[@name='Header.name']")
            v = ep.findtext("stringProp[@name='Header.value']")
            if n and n not in headers:
                headers[n] = v or ""
        break

    csvs = []
    for el in root.iter("CSVDataSet"):
        csvs.append({"file": _prop(el, "filename"),
                     "variables": [v for v in _prop(el, "variableNames").split(",") if v]})

    plan_name = name or (root.find(".//TestPlan").get("testname")
                         if root.find(".//TestPlan") is not None else "Imported plan")
    spec = {"name": plan_name, "platform": "local",
            "comments": "Imported by jmxgen from %s" % os.path.basename(path),
            "variables": {},
            "defaults": defaults or {"protocol": "https", "domain": ""},
            "headers": headers or {"Accept": "*/*"},
            "cookies": True,
            "thread_groups": tgs or [{"name": "Thread Group", "threads": 1,
                                      "loops": 1, "steps": []}]}
    if csvs:
        spec["csv"] = csvs
    return spec, counts


# --------------------------------------------------------------------------
# optimize - repair / slim an existing .jmx (works on unparseable files too)
# --------------------------------------------------------------------------

INVALID_REF = re.compile(r"&#x(?:0|1|2|3|4|5|6|7|8|b|c|e|f|1[0-9a-f])[0-9a-f]?;", re.I)
BINARY_MARKER = "REMOVED_BY_JMXGEN_BINARY_BODY"


def _sanitize_xml(text):
    """Strip XML-1.0-illegal character references and raw control bytes.

    PerfPilot-style generators paste raw (often gzipped) request bodies into
    Argument.value; those bytes are not legal XML and no standard tool can read
    the file. Whole binary blobs are replaced, stray control chars dropped."""
    n_blobs = [0]

    def fix_value(m):
        inner = m.group(1)
        if INVALID_REF.search(inner) or "\x00" in inner:
            n_blobs[0] += 1
            return '<stringProp name="Argument.value">%s</stringProp>' % BINARY_MARKER
        return m.group(0)

    n_refs = len(INVALID_REF.findall(text))
    text = re.sub(r'<stringProp name="Argument\.value">(.*?)</stringProp>',
                  fix_value, text, flags=re.S)
    text = INVALID_REF.sub("", text)
    text = "".join(c for c in text if c in "\t\n\r" or ord(c) >= 0x20)
    return text, n_blobs[0], n_refs


SAMPLER_TAGS = ("HTTPSamplerProxy", "JSR223Sampler", "DebugSampler", "TestAction")
LISTENER_GUIS = ("ViewResultsFullVisualizer", "ViewResultsTree", "StatVisualizer",
                 "StatGraphVisualizer", "SummaryReport", "TableVisualizer",
                 "GraphVisualizer", "AssertionVisualizer", "RespTimeGraphVisualizer")


def _sampler_url(el):
    def prop(n):
        return (el.findtext("stringProp[@name='%s']" % n) or "").strip()
    return prop("HTTPSampler.domain"), prop("HTTPSampler.path")


def _pairs(ht):
    """Yield (index, element, hashTree) pairs inside a hashTree."""
    kids = list(ht)
    i = 0
    while i < len(kids):
        el = kids[i]
        sub = kids[i + 1] if i + 1 < len(kids) and kids[i + 1].tag == "hashTree" else None
        yield el, sub
        i += 2 if sub is not None else 1


def _remove_pair(parent, el):
    kids = list(parent)
    try:
        i = kids.index(el)
    except ValueError:
        return
    if i + 1 < len(kids) and kids[i + 1].tag == "hashTree":
        parent.remove(kids[i + 1])
    parent.remove(el)


def _set_proxy_on_defaults(root, proxy, ET):
    """Put the proxy on HTTP Request Defaults, creating it if the plan has none."""
    props = {"HTTPSampler.proxyHost": proxy.get("host", ""),
             "HTTPSampler.proxyPort": str(proxy.get("port", "")),
             "HTTPSampler.proxyScheme": proxy.get("scheme", "http")}
    if proxy.get("user"):
        props["HTTPSampler.proxyUser"] = proxy["user"]
        props["HTTPSampler.proxyPass"] = proxy.get("password", "")

    targets = [e for e in root.iter("ConfigTestElement")
               if e.get("guiclass") == "HttpDefaultsGui"]
    if not targets:
        el = ET.Element("ConfigTestElement", {
            "guiclass": "HttpDefaultsGui", "testclass": "ConfigTestElement",
            "testname": "HTTP Request Defaults (proxy)", "enabled": "true"})
        targets = [el]
        plan_ht = root.find("hashTree/hashTree")
        if plan_ht is not None:
            plan_ht.insert(0, ET.Element("hashTree"))
            plan_ht.insert(0, el)
    for el in targets:
        for name, val in props.items():
            node_ = el.find("stringProp[@name='%s']" % name)
            if node_ is None:
                node_ = ET.SubElement(el, "stringProp")
                node_.set("name", name)
            node_.text = val
    return len(targets)


def optimize(path, out, keep_domains=None, drop_third_party=False, drop_static=False,
             drop_disabled=True, strip_listeners=True, cap_parallel=None,
             dedupe_headers=False, add_timeouts=True, keep_trackers=False,
             max_samplers=None, dry_run=False, proxy=None):
    import xml.etree.ElementTree as ET

    raw = open(path, "r", encoding="utf-8", errors="replace").read()
    before_bytes = len(raw.encode("utf-8"))
    text, n_blobs, n_refs = _sanitize_xml(raw)
    root = ET.fromstring(text)

    stats = {"binary_bodies_removed": n_blobs, "illegal_char_refs_removed": n_refs}

    all_samplers = [e for e in root.iter() if e.tag in SAMPLER_TAGS]
    stats["samplers_before"] = len(all_samplers)

    # work out what the system under test actually is
    hosts = {}
    for s in root.iter("HTTPSamplerProxy"):
        d = (s.findtext("stringProp[@name='HTTPSampler.domain']") or "").strip()
        if d and "${" not in d:
            hosts[d] = hosts.get(d, 0) + 1
    if keep_domains:
        keep = [d.strip().lower() for d in keep_domains if d.strip()]
    elif drop_third_party and hosts:
        top = max(hosts, key=hosts.get)
        keep = [_reg_domain(top)]
    else:
        keep = []
    stats["kept_domains"] = keep
    stats["distinct_domains_before"] = len(hosts)

    dropped = {"third_party": 0, "static": 0, "disabled": 0, "listeners": 0, "over_cap": 0}
    kept_count = [0]

    def prune(ht):
        for el, sub in list(_pairs(ht)):
            tag, guic = el.tag, el.get("guiclass", "")
            if drop_disabled and el.get("enabled") == "false":
                _remove_pair(ht, el)
                dropped["disabled"] += 1
                continue
            if strip_listeners and tag == "ResultCollector" and guic in LISTENER_GUIS:
                _remove_pair(ht, el)
                dropped["listeners"] += 1
                continue
            if tag == "HTTPSamplerProxy":
                domain, spath = _sampler_url(el)
                host = domain.lower()
                if keep and host and "${" not in host:
                    third = not any(host == k or host.endswith("." + k) for k in keep)
                    if third and (not keep_trackers or _looks_like_tracker(host)):
                        _remove_pair(ht, el)
                        dropped["third_party"] += 1
                        continue
                if drop_static and STATIC_EXT.search(spath or ""):
                    _remove_pair(ht, el)
                    dropped["static"] += 1
                    continue
                if max_samplers and kept_count[0] >= max_samplers:
                    _remove_pair(ht, el)
                    dropped["over_cap"] += 1
                    continue
                kept_count[0] += 1
                if add_timeouts:
                    for prop, val in (("HTTPSampler.connect_timeout", "5000"),
                                      ("HTTPSampler.response_timeout", "30000")):
                        node_ = el.find("stringProp[@name='%s']" % prop)
                        if node_ is None:
                            node_ = ET.SubElement(el, "stringProp")
                            node_.set("name", prop)
                        if not (node_.text or "").strip():
                            node_.text = val
            if cap_parallel and "ParallelSampler" in tag:
                for name, val in (("MAX_THREAD_NUMBER", str(cap_parallel)),
                                  ("ParallelSampler.maxThreadNumber", str(cap_parallel))):
                    n_ = el.find("intProp[@name='%s']" % name)
                    if n_ is None:
                        n_ = ET.SubElement(el, "intProp")
                        n_.set("name", name)
                    n_.text = val
                lim = el.find("boolProp[@name='LIMIT_MAX_THREAD_NUMBER']")
                if lim is None:
                    lim = ET.SubElement(el, "boolProp")
                    lim.set("name", "LIMIT_MAX_THREAD_NUMBER")
                lim.text = "true"
            if sub is not None:
                prune(sub)

    top_ht = root.find("hashTree")
    if top_ht is not None:
        prune(top_ht)

    # drop controllers left with no samplers under them
    def prune_empty(ht):
        changed = False
        for el, sub in list(_pairs(ht)):
            if sub is None:
                continue
            prune_empty(sub)
            is_ctrl = ("Controller" in el.tag or "ParallelSampler" in el.tag)
            if is_ctrl and not [e for e in sub.iter() if e.tag in SAMPLER_TAGS]:
                _remove_pair(ht, el)
                changed = True
        return changed
    if top_ht is not None:
        prune_empty(top_ht)

    # hoist header pairs shared by every remaining sampler into one plan-level manager
    if dedupe_headers and top_ht is not None:
        managers = []
        for parent in root.iter("hashTree"):
            for el, _sub in _pairs(parent):
                if el.tag == "HeaderManager":
                    managers.append((parent, el))
        sets = []
        for _p, hm in managers:
            pairs = set()
            for ep in hm.iter("elementProp"):
                n = ep.findtext("stringProp[@name='Header.name']")
                v = ep.findtext("stringProp[@name='Header.value']")
                if n is not None:
                    pairs.add((n, v or ""))
            sets.append(pairs)
        common = set.intersection(*sets) if len(sets) > 1 else set()
        # keep only headers that are genuinely global and static
        common = {(n, v) for (n, v) in common if "${" not in (v or "")}
        if common:
            for (parent, hm), pairs in zip(managers, sets):
                coll = hm.find("collectionProp")
                if coll is None:
                    continue
                for ep in list(coll):
                    n = ep.findtext("stringProp[@name='Header.name']")
                    v = ep.findtext("stringProp[@name='Header.value']") or ""
                    if (n, v) in common:
                        coll.remove(ep)
                if len(list(coll)) == 0:
                    _remove_pair(parent, hm)
            shared = ET.Element("HeaderManager", {
                "guiclass": "HeaderPanel", "testclass": "HeaderManager",
                "testname": "HTTP Header Manager (shared)", "enabled": "true"})
            coll = ET.SubElement(shared, "collectionProp", {"name": "HeaderManager.headers"})
            for n, v in sorted(common):
                ep = ET.SubElement(coll, "elementProp", {"name": "", "elementType": "Header"})
                a = ET.SubElement(ep, "stringProp", {"name": "Header.name"}); a.text = n
                b = ET.SubElement(ep, "stringProp", {"name": "Header.value"}); b.text = v
            plan_ht = top_ht.find("hashTree")
            if plan_ht is not None:
                plan_ht.insert(0, ET.Element("hashTree"))
                plan_ht.insert(0, shared)
            stats["headers_hoisted"] = len(common)
            stats["header_managers_removed"] = sum(1 for s in sets if s and s <= common)

    if proxy:
        stats["proxy_set_on"] = _set_proxy_on_defaults(root, proxy, ET)

    stats["samplers_after"] = len([e for e in root.iter() if e.tag in SAMPLER_TAGS])
    stats["dropped"] = dropped

    try:
        ET.indent(root, space="  ")
    except AttributeError:
        pass
    body = ET.tostring(root, encoding="unicode")
    result = '<?xml version="1.0" encoding="UTF-8"?>\n' + body + "\n"
    stats["bytes_before"] = before_bytes
    stats["bytes_after"] = len(result.encode("utf-8"))

    if not dry_run:
        open(out, "w", encoding="utf-8").write(result)
    return stats


def print_optimize_report(path, out, stats):
    mb = lambda b: b / 1048576.0
    print("== optimize %s -> %s ==" % (path, out))
    if stats["binary_bodies_removed"] or stats["illegal_char_refs_removed"]:
        print("  fixed  %d binary request body(ies), %d illegal XML char refs "
              "(file is now standard XML)"
              % (stats["binary_bodies_removed"], stats["illegal_char_refs_removed"]))
    d = stats["dropped"]
    for key, label in (("third_party", "third-party/tracker requests"),
                       ("static", "static asset requests"),
                       ("disabled", "disabled elements"),
                       ("listeners", "heavy listeners"),
                       ("over_cap", "requests over --max-samplers")):
        if d.get(key):
            print("  dropped %d %s" % (d[key], label))
    if stats.get("headers_hoisted"):
        print("  hoisted %d shared header(s) into one plan-level Header Manager "
              "(%d per-sampler managers emptied)"
              % (stats["headers_hoisted"], stats.get("header_managers_removed", 0)))
    if stats.get("proxy_set_on"):
        print("  proxy   set on %d HTTP Request Defaults element(s)" % stats["proxy_set_on"])
    if stats.get("kept_domains"):
        print("  kept    domains: %s (of %d seen)"
              % (", ".join(stats["kept_domains"]), stats["distinct_domains_before"]))
    print("  result  %d -> %d samplers, %.1f MB -> %.1f MB"
          % (stats["samplers_before"], stats["samplers_after"],
             mb(stats["bytes_before"]), mb(stats["bytes_after"])))


# --------------------------------------------------------------------------
# probe - give it page URLs, it loads each page and authors the plan
# --------------------------------------------------------------------------

# HTML in the wild is not always quoted - accept href=/a.css as well as href="/a.css"
_VAL = r"""(?:"([^"]+)"|'([^']+)'|([^\s"'>]+))"""


def _asset_pattern(tag, attr):
    return r"<%s\b[^>]*?\b%s\s*=\s*%s" % (tag, attr, _VAL)


ASSET_ATTRS = (
    (_asset_pattern("script", "src"), "script"),
    (_asset_pattern("link", "href"), "link"),
    (_asset_pattern("img", "src"), "img"),
    (_asset_pattern("iframe", "src"), "iframe"),
    (_asset_pattern("source", "src"), "source"),
    (_asset_pattern("video", "poster"), "poster"),
    (r"\burl\(\s*['\"]?([^'\")]+)['\"]?\s*\)", "css"),
    (r"<[^>]*?\bdata-src\s*=\s*" + _VAL, "lazy"),
)


def _page_assets(html, base_url, limit=400):
    """Every sub-resource the page declares, in document order, de-duplicated."""
    seen, out = set(), []
    for pattern, kind in ASSET_ATTRS:
        for m in re.finditer(pattern, html, re.I):
            raw = next((g for g in m.groups() if g), "").strip()
            if not raw or raw.startswith(("data:", "javascript:", "mailto:", "#", "about:")):
                continue
            url = urllib.parse.urljoin(base_url, raw)
            if not url.startswith("http"):
                continue
            url = url.split("#")[0]
            if url in seen:
                continue
            seen.add(url)
            out.append((url, kind))
            if len(out) >= limit:
                return out
    # srcset lists: "a.jpg 1x, b.jpg 2x"
    for m in re.finditer(r"\bsrcset\s*=\s*['\"]([^'\"]+)['\"]", html, re.I):
        for part in m.group(1).split(","):
            cand = part.strip().split(" ")[0]
            if not cand:
                continue
            url = urllib.parse.urljoin(base_url, cand).split("#")[0]
            if url.startswith("http") and url not in seen and len(out) < limit:
                seen.add(url)
                out.append((url, "srcset"))
    return out


# high-confidence: these strings only ever appear on a challenge page
BOT_WALL_STRONG = re.compile(
    r"awswaf|gokuProps|cf_chl_|cf-browser-verification|_pxCaptcha|px-captcha|"
    r"distil_r_captcha|Incapsula|perimeterx|__cf_bm_challenge", re.I)

# ambiguous: a real page may legitimately contain these words, so they only count
# on a page too small to be the real thing
BOT_WALL_WEAK = re.compile(
    r"captcha|Just a moment|Enter the characters you see below|Access Denied|"
    r"Checking your browser|Attention Required|Request unsuccessful|Bot Manager", re.I)


def _looks_blocked(html, title):
    """A WAF / bot-challenge page instead of the real one.

    Authoring a plan from one produces a single meaningless sampler, so it must be
    loud - but a 200 KB page that merely mentions "captcha" is a real page."""
    head = html[:20000]
    if BOT_WALL_STRONG.search(head):
        return True
    if len(html) < 15000 and BOT_WALL_WEAK.search(head):
        return True
    return len(html) < 4000 and not title


def _read_url_list(source):
    """URLs from: a literal URL, a .txt/.csv/.xlsx file, or '-' for stdin."""
    if source == "-":
        return [l.strip() for l in sys.stdin if l.strip() and not l.startswith("#")]
    if re.match(r"^https?://", source):
        return [source]
    low = source.lower()
    if low.endswith((".xlsx", ".xlsm")):
        _cfg, rows = _rows_from_xlsx(source)
        urls = []
        for r in rows:
            for key in ("url", "page", "page_url", "path", "link"):
                if r.get(key):
                    urls.append(str(r[key]).strip())
                    break
        return urls
    if low.endswith(".csv"):
        _cfg, rows = _rows_from_csv(source)
        return [str(r.get("url") or r.get("page") or "").strip()
                for r in rows if (r.get("url") or r.get("page"))]
    out = []
    for line in open(source, "r", encoding="utf-8-sig"):
        line = line.strip()
        if line and not line.startswith("#"):
            out.append(line.split(",")[0].strip() if "," in line and " " not in line else line)
    return out


def probe_to_spec(sources, name=None, assets=True, same_domain_assets=True,
                  keep_static=True, keep_trackers=False, parallel=None,
                  think_time=1000, timeout=20, forms=True, max_assets=400):
    """Load each page URL and author a plan: main request + its sub-resources.

    No browser and no HAR: each page is fetched once and its declared
    sub-resources are read out of the markup."""
    urls = []
    for src in sources:
        urls += _read_url_list(src)
    urls = [u if "//" in u else "https://" + u for u in urls if u]
    if not urls:
        raise ValueError("no URLs found in %s" % ", ".join(map(str, sources)))

    hosts = {}
    for u in urls:
        h = urllib.parse.urlsplit(u).hostname
        if h:
            hosts[h] = hosts.get(h, 0) + 1
    top_host = max(hosts, key=hosts.get)
    keep_root = _reg_domain(top_host)
    scheme = urllib.parse.urlsplit(urls[0]).scheme or "https"
    port = str(urllib.parse.urlsplit(urls[0]).port or "")

    pages, variables = [], {}
    stats = {"pages": 0, "failed": [], "blocked": [], "assets": 0,
             "skipped_third_party": 0, "skipped_static": 0}

    for i, url in enumerate(urls, 1):
        try:
            html, ctype = _fetch(url, timeout)
        except Exception as exc:
            stats["failed"].append((url, str(exc)))
            continue
        u = urllib.parse.urlsplit(url)

        title = ""
        m = re.search(r"<title[^>]*>(.*?)</title>", html, re.S | re.I)
        if m:
            title = re.sub(r"\s+", " ", m.group(1)).strip()[:60]
        label = "Page :%d [ %s ]" % (i, title or (u.path or "/"))

        if _looks_blocked(html, title):
            stats["blocked"].append(url)
            continue

        main = {"name": "Main URL: %s" % ((u.path or "/") + (("?" + u.query) if u.query else ""))[:90],
                "method": "GET",
                "path": (u.path or "/") + (("?" + u.query) if u.query else ""),
                "assert": [{"field": "code", "match": "equals", "pattern": "200"}]}
        if (u.hostname or "") != top_host:
            main["domain"], main["protocol"] = u.hostname, u.scheme
        sub_steps, form_steps = [], []

        if forms and "html" in (ctype or "").lower():
            p = _PageParser()
            try:
                p.feed(html)
            except Exception:
                pass
            for fi, form in enumerate(p.forms, 1):
                action = urllib.parse.urljoin(url, form["action"] or url)
                ap = urllib.parse.urlsplit(action)
                if ap.hostname and ap.hostname != u.hostname:
                    continue
                params = {}
                for inp in form["inputs"]:
                    if inp["type"] == "hidden":
                        var = re.sub(r"\W+", "_", inp["name"]).upper()
                        # several forms on a page often share the same hidden field -
                        # one extractor for it is enough
                        if not any(e.get("var") == var for e in main.get("extract", [])):
                            main.setdefault("extract", []).append({
                                "type": "boundary", "var": var,
                                "left": 'name="%s" value="' % inp["name"], "right": '"',
                                "default": inp["value"] or "NOT_FOUND"})
                        params[inp["name"]] = "${%s}" % var
                    else:
                        val = _placeholder(inp)
                        params[inp["name"]] = val
                        mv = re.fullmatch(r"\$\{(\w+)\}", val)
                        if mv:
                            variables.setdefault(mv.group(1), "CHANGE_ME")
                if params:
                    form_steps.append({
                        "name": "Form %d: %s %s" % (fi, form["method"], ap.path or "/"),
                        "method": form["method"],
                        "path": (ap.path or "/") + (("?" + ap.query) if ap.query else ""),
                        "params": params,
                        "assert": [{"field": "code", "match": "equals", "pattern": "200"}]})

        if assets and "html" in (ctype or "").lower():
            for a_url, kind in _page_assets(html, url, max_assets):
                au = urllib.parse.urlsplit(a_url)
                ahost = au.hostname or ""
                third = not (ahost == keep_root or ahost.endswith("." + keep_root))
                if _looks_like_tracker(ahost) and not keep_trackers:
                    stats["skipped_third_party"] += 1
                    continue
                if third and same_domain_assets:
                    stats["skipped_third_party"] += 1
                    continue
                if not keep_static and STATIC_EXT.search(au.path or ""):
                    stats["skipped_static"] += 1
                    continue
                sub_path = (au.path or "/") + (("?" + au.query) if au.query else "")
                sub = {"name": ("Sub URL: %s%s" % ("" if ahost == top_host else ahost,
                                                   sub_path))[:90],
                       "method": "GET",
                       "path": sub_path}
                if ahost != top_host:
                    sub["domain"], sub["protocol"] = ahost, au.scheme
                    if au.port:
                        sub["port"] = str(au.port)
                sub_steps.append(sub)
                stats["assets"] += 1

        # order matters: the HTML must complete before anything that uses a token
        # extracted from it, so only the sub-resources are ever fanned out
        page_steps = [main]
        if sub_steps:
            page_steps.append({"parallel": "Assets - page %d" % i,
                               "max_threads": parallel, "steps": sub_steps}
                              if parallel else
                              {"transaction": "Assets - page %d" % i, "steps": sub_steps})
        page_steps += form_steps
        if think_time:
            page_steps.append({"pause": think_time, "name": "Pause after page %d" % i})
        pages.append({"transaction": label, "steps": page_steps})
        stats["pages"] += 1

    if not pages:
        detail = ""
        if stats["failed"]:
            detail += "\n" + "\n".join("  %s -> %s" % (u, e) for u, e in stats["failed"][:5])
        if stats["blocked"]:
            detail = ("\n%d of them answered with a WAF / bot-challenge page rather than "
                      "the real one. A plain HTTP client cannot get past that - record a "
                      "HAR from your own browser and use `from-har`, or drive the API "
                      "directly with `from-openapi` / `from-postman`." % len(stats["blocked"]))
        raise ValueError("no page could be authored from the %d URL(s)%s" % (len(urls), detail))

    spec = {
        "name": name or ("%s page load test" % top_host),
        "platform": "local",
        "comments": "Authored by jmxgen: probed %d page URL(s), %d sub-resource(s)"
                    % (stats["pages"], stats["assets"]),
        "variables": variables,
        "defaults": {"protocol": scheme, "domain": top_host, "port": port,
                     "connect_timeout": 5000, "response_timeout": 30000},
        "headers": {"User-Agent": "Mozilla/5.0 (JMeter)",
                    "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"},
        "cookies": True,
        "cache": True,
        "thread_groups": [{
            "name": "Page Journey", "threads": 10, "ramp_up": 30, "duration": 300,
            "steps": pages,
        }],
    }
    return spec, stats


# --------------------------------------------------------------------------
# record - launch a real browser, you click, it captures, you get a .jmx
# --------------------------------------------------------------------------

PLAYWRIGHT_HELP = """this needs Playwright:

    pip install playwright
    python3 -m playwright install chromium

(one-off, ~90 MB for the browser). Use --browser chrome to drive your installed
Google Chrome instead of downloading Chromium."""


def _wait_for_finish(browser, prompt):
    """Return when the user presses Enter or closes the browser window."""
    import time
    print(prompt, flush=True)
    try:
        import select
        while True:
            if not browser.is_connected():
                print("  browser closed - wrapping up", flush=True)
                return
            ready, _, _ = select.select([sys.stdin], [], [], 0.5)
            if ready:
                sys.stdin.readline()
                return
    except (ImportError, OSError):          # no select on this stdin (Windows, pipes)
        try:
            input()
        except EOFError:
            while browser.is_connected():
                time.sleep(0.5)


def record_session(urls, har_path, browser_name="chromium", profile_dir=None,
                   manual=False, settle_ms=1500, timeout=60000, viewport=None,
                   ignore_https_errors=True, proxy=None):
    """Drive a real browser over a list of URLs, capturing everything into a HAR.

    Headless and unattended by default - hand it a URL list and it walks them.
    With manual=True the window is visible and you drive it yourself."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        raise ValueError(PLAYWRIGHT_HELP)

    ctx_args = {
        "record_har_path": har_path,
        "record_har_content": "embed",   # response bodies - correlation needs them
        "record_har_mode": "full",
        "ignore_https_errors": ignore_https_errors,
    }
    if manual:
        ctx_args["no_viewport"] = viewport is None
        ctx_args["viewport"] = viewport
    launch_args = {"headless": not manual}
    if proxy and proxy.get("host"):
        server = "%s://%s:%s" % (proxy.get("scheme", "http"), proxy["host"],
                                 proxy.get("port", 8080))
        pcfg = {"server": server}
        if proxy.get("user"):
            pcfg["username"] = proxy["user"]
            pcfg["password"] = proxy.get("password", "")
        launch_args["proxy"] = pcfg
    if manual:
        launch_args["args"] = ["--start-maximized"]
    if browser_name == "chrome":
        launch_args["channel"] = "chrome"
    elif browser_name == "edge":
        launch_args["channel"] = "msedge"

    warnings, visited = [], 0

    with sync_playwright() as pw:
        engine = pw.firefox if browser_name == "firefox" else \
            (pw.webkit if browser_name == "webkit" else pw.chromium)
        if profile_dir:
            # persistent profile: log in once, reuse the session next time
            context = engine.launch_persistent_context(profile_dir, **launch_args, **ctx_args)
            browser = context.browser or context
        else:
            browser = engine.launch(**launch_args)
            context = browser.new_context(**ctx_args)

        def visit(url, page):
            try:
                page.goto(url, wait_until="load", timeout=timeout)
            except Exception as exc:
                warnings.append("%s -> %s" % (url, str(exc).splitlines()[0]))
                return False
            try:                       # let lazy loads and XHR settle
                page.wait_for_load_state("networkidle", timeout=min(timeout, 15000))
            except Exception:
                pass
            if settle_ms:
                page.wait_for_timeout(settle_ms)
            try:
                body = page.content()[:20000]
                if BOT_WALL_STRONG.search(body):
                    warnings.append("%s -> WAF / bot-challenge page (try --manual)" % url)
            except Exception:
                pass
            return True

        if manual:
            page = context.pages[0] if context.pages else context.new_page()
            if visit(urls[0], page):
                visited = 1
            for extra in urls[1:]:
                warnings.append("--manual drives one starting URL; %s not opened "
                                "automatically" % extra)
            _wait_for_finish(
                browser if hasattr(browser, "is_connected") else context,
                "\n  >> Browser is open. Do your journey - log in, click, submit.\n"
                "     Everything the browser requests is being captured.\n"
                "     Press ENTER here (or close the browser) when you are done.\n")
        else:
            for i, url in enumerate(urls, 1):
                # a fresh page per URL gives each one its own entry in the HAR's
                # page list, which becomes one transaction in the plan
                page = context.new_page()
                print("  [%d/%d] %s" % (i, len(urls), url[:90]), flush=True)
                if visit(url, page):
                    visited += 1
                try:
                    page.close()
                except Exception:
                    pass

        try:
            context.close()              # flushes the HAR
        except Exception:
            pass
        try:
            if hasattr(browser, "close"):
                browser.close()
        except Exception:
            pass

    if not os.path.exists(har_path):
        raise ValueError("no HAR was written - nothing was captured")
    if not visited:
        raise ValueError("none of the %d URL(s) could be loaded:\n  %s"
                         % (len(urls), "\n  ".join(warnings[:5])))
    return har_path, warnings


def capture_via_proxy(out_har, port=8080, duration=None, extra_args=None):
    """Record anything that can be pointed at a proxy - mobile, desktop, Postman,
    a backend service. Chrome is not required."""
    import shutil
    import subprocess

    mitm = shutil.which("mitmdump") or os.path.expanduser(
        "~/Library/Python/%d.%d/bin/mitmdump" % sys.version_info[:2])
    if not (shutil.which("mitmdump") or os.path.exists(mitm)):
        raise ValueError(
            "mitmdump not found. Install it with:\n"
            "    pip install mitmproxy\n"
            "(it may land in ~/Library/Python/*/bin - add that to PATH)")
    addon = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mitm_har.py")
    if not os.path.exists(addon):
        raise ValueError("mitm_har.py must sit next to jmxgen.py")

    cmd = [mitm, "-s", addon, "--set", "jmxgen_har=%s" % out_har,
           "-p", str(port), "-q"] + list(extra_args or [])
    print("proxy listening on port %d" % port)
    print("  1. point the device / app / Postman at  <this-machine-ip>:%d" % port)
    print("  2. install mitmproxy's CA on it from    http://mitm.it")
    print("  3. do the journey, then press ctrl-c here%s"
          % ("" if duration is None else " (or wait %ds)" % duration))
    proc = subprocess.Popen(cmd)
    try:
        proc.wait(timeout=duration)
    except KeyboardInterrupt:
        proc.terminate()
        proc.wait(timeout=20)
    except Exception:
        proc.terminate()                 # duration elapsed - stop and flush the HAR
        proc.wait(timeout=20)
    if not os.path.exists(out_har):
        raise ValueError("nothing was captured - did traffic actually go through the proxy?")
    return out_har


def start_console(port=8770, open_browser=True, idle_timeout=60):
    """Launch the web console - the entry point for anyone who is not scripting."""
    import importlib.util
    import threading
    import time
    import webbrowser

    print("jmxgen console  starting on port %d ..." % port, flush=True)
    here = os.path.dirname(os.path.abspath(__file__))
    server_py = os.path.join(here, "jmxgen_server.py")
    if not os.path.exists(server_py):
        sys.exit("jmxgen_server.py is not next to jmxgen.py")
    if os.environ.get("JMXGEN_DEBUG"):
        print("[dbg] loading %s" % server_py, flush=True)
    spec = importlib.util.spec_from_file_location("jmxgen_server", server_py)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    if os.environ.get("JMXGEN_DEBUG"):
        print("[dbg] loaded; starting server", flush=True)

    url = "http://localhost:%d" % port
    if open_browser:
        # a frozen binary needs a few seconds to unpack, so poll until the port
        # actually answers - opening on a timer lands on a connection error
        def _open_when_ready():
            import socket
            for _ in range(600):                       # up to 60s, then give up
                try:
                    with socket.create_connection(("127.0.0.1", port), 0.2):
                        break
                except OSError:
                    time.sleep(0.1)
            else:
                return
            webbrowser.open(url)
        threading.Thread(target=_open_when_ready, daemon=True).start()
    os.environ["PORT"] = str(port)
    os.environ["IDLE_TIMEOUT"] = str(idle_timeout)
    return mod.main() or 0


def doctor():
    """What is installed, what it unlocks, and how to get the rest."""
    import shutil
    import subprocess

    rows = []

    def check(name, ok, detail, unlocks, fix=""):
        rows.append((name, ok, detail, unlocks, fix))

    check("python", True, "%d.%d.%d" % sys.version_info[:3], "everything")

    try:
        import yaml  # noqa: F401
        check("PyYAML", True, "installed", "YAML specs and rule files")
    except ImportError:
        check("PyYAML", False, "missing", "YAML specs (JSON still works)",
              "pip install pyyaml")

    try:
        import openpyxl  # noqa: F401
        check("openpyxl", True, "installed", "Excel input (.xlsx)",)
    except ImportError:
        check("openpyxl", False, "missing", "Excel input (.csv still works)",
              "pip install openpyxl")

    try:
        import requests  # noqa: F401
        check("requests", True, "installed", "running on HyperExecute")
    except ImportError:
        check("requests", False, "missing", "running on HyperExecute",
              "pip install requests")

    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            path = pw.chromium.executable_path
        have = bool(path and os.path.exists(path))
        check("Playwright", have, "browser ready" if have else "browser not downloaded",
              "`record` - headless / manual browser capture",
              "" if have else "python3 -m playwright install chromium")
    except Exception:
        check("Playwright", False, "missing",
              "`record` - headless / manual browser capture",
              "pip install playwright && python3 -m playwright install chromium")

    mitm = shutil.which("mitmdump") or os.path.expanduser(
        "~/Library/Python/%d.%d/bin/mitmdump" % sys.version_info[:2])
    have_mitm = bool(shutil.which("mitmdump") or os.path.exists(mitm))
    check("mitmproxy", have_mitm, "installed" if have_mitm else "missing",
          "`capture` - mobile / desktop / Postman / backend traffic",
          "" if have_mitm else "pip install mitmproxy")

    java = shutil.which("java")
    jver = ""
    if java:
        try:
            out = subprocess.run([java, "-version"], capture_output=True, text=True,
                                 timeout=15)
            jver = (out.stderr or out.stdout).splitlines()[0][:40]
        except Exception:
            jver = "present"
    check("Java", bool(java), jver or "missing", "JMeter itself",
          "" if java else "install a JDK 11 or 17")

    jmeter = shutil.which("jmeter")
    check("JMeter", bool(jmeter), jmeter or "missing",
          "`verify --deep` and `replay` - the validation gates",
          "" if jmeter else "brew install jmeter (or unpack the Apache tarball)")

    plugins = {}
    if jmeter:
        base = os.path.realpath(jmeter)
        for _ in range(4):
            base = os.path.dirname(base)
            libext = os.path.join(base, "libexec", "lib", "ext")
            if os.path.isdir(libext):
                break
            libext = os.path.join(base, "lib", "ext")
            if os.path.isdir(libext):
                break
        else:
            libext = ""
        if libext and os.path.isdir(libext):
            jars = " ".join(os.listdir(libext)).lower()
            plugins["bzm Parallel Controller"] = "parallel" in jars or "bzm" in jars
            plugins["WebDriver Set"] = "webdriver" in jars
    for name, present in plugins.items():
        check("plugin: " + name, present, "installed" if present else "not installed",
              "plans that use it",
              "" if present else "install via JMeter's Plugins Manager")

    print("== jmxgen doctor ==")
    width = max(len(r[0]) for r in rows)
    for name, ok, detail, unlocks, fix in rows:
        print("  %s %-*s %-22s %s"
              % ("." if ok else "X", width, name, detail[:22], unlocks))
        if not ok and fix:
            print("     %s%s" % (" " * width, fix))
    missing = [r for r in rows if not r[1]]
    print("  %s" % ("everything needed is installed"
                    if not missing else
                    "%d optional component(s) missing - each only limits the feature listed"
                    % len(missing)))
    return 0


def _report_correlations(found, report_path=None):
    """Print what was correlated and why - a silent rewrite is not reviewable."""
    if not found:
        print("  no dynamic values correlated "
              "(nothing a later request sent came from an earlier response)")
        return
    print("  correlated %d value(s):" % len(found))
    for c in found:
        print("    %-24s %-18s %-9s from %-7s  %s -> %s"
              % ("${%s}" % c["var"], c["rule"], c["confidence"], c["found_in"],
                 c["source_step"][:28], c["used_in"][:28]))
    weak = [c for c in found if c["confidence"] == "low"]
    if weak:
        print("  %d of those matched no rule - review them before a long run"
              % len(weak))
    if report_path:
        with open(report_path, "w", encoding="utf-8") as fh:
            json.dump(found, fh, indent=2)
        print("  correlation report: %s" % report_path)


def main():
    ap = argparse.ArgumentParser(prog="jmxgen", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    if len(sys.argv) == 1:
        return start_console(int(os.environ.get("PORT", 8770)))

    sub = ap.add_subparsers(dest="cmd", required=True)

    prox = argparse.ArgumentParser(add_help=False)
    prox.add_argument("--proxy", metavar="HOST:PORT",
                      help="route the generated plan through a forward proxy "
                           "(also used while recording)")
    prox.add_argument("--proxy-auth", metavar="USER:PASS", help="proxy credentials")
    prox.add_argument("--login", metavar="[METHOD ]PATH",
                      help="authenticate once in a setUp thread group, e.g. "
                           "\"POST /oauth/token\"; the token is published as a "
                           "property and sent on every request")
    prox.add_argument("--login-body", metavar="JSON",
                      help="request body for --login (JSON or form string)")
    prox.add_argument("--login-token", metavar="JSONPATH", default="$.access_token",
                      help="where the token is in the login response (default $.access_token)")
    prox.add_argument("--login-header", metavar="NAME: VALUE",
                      help="how to send it (default 'Authorization: Bearer ${__P(AUTH_TOKEN)}')")
    prox.add_argument("--parameterize", metavar="LITERAL=COLUMN", action="append",
                      help="replace a recorded literal with a ${CSV column}, e.g. "
                           "--parameterize bob@example.com=email (repeatable)")
    prox.add_argument("--suggest-parameters", action="store_true",
                      help="list values that look like per-user data, then exit")
    prox.add_argument("--replay", action="store_true",
                      help="after generating, run the plan once as a single user and "
                           "report anything that fails or fails to resolve")
    prox.add_argument("--csv", metavar="FILE:COL1,COL2", action="append",
                      help="attach a CSV Data Set, e.g. users.csv:username,password "
                           "(repeatable)")

    p = sub.add_parser("init", help="write a starter spec", parents=[prox])
    p.add_argument("out", nargs="?", default="testplan.yaml")
    p.add_argument("--kind", choices=["http", "webdriver"], default="http")

    p = sub.add_parser("build", help="spec -> .jmx", parents=[prox])
    p.add_argument("spec")
    p.add_argument("-o", "--out", help="output .jmx (default: <spec name>.jmx)")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("from-har", help="browser recording (HAR) -> .jmx", parents=[prox])
    p.add_argument("har")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--include", help="regex: keep only URLs matching")
    p.add_argument("--exclude", help="regex: drop URLs matching")
    p.add_argument("--keep-static", action="store_true", help="keep images/css/js requests")
    p.add_argument("--mode", choices=["auto", "api", "web"], default="auto",
                   help="api = service calls only (XHR/fetch/REST); "
                        "web = everything the browser fetched; auto = drop static/trackers")
    p.add_argument("--methods", metavar="LIST",
                   help="keep only these HTTP methods, e.g. POST or POST,PUT,DELETE")
    p.add_argument("--real-think-time", action="store_true",
                   help="use the pauses the real user left, instead of a fixed value")
    p.add_argument("--keep-third-party", action="store_true",
                   help="keep other domains (trackers are still dropped)")
    p.add_argument("--no-pages", action="store_true",
                   help="one flat flow instead of a transaction per page")
    p.add_argument("--no-correlate", action="store_true",
                   help="skip automatic token correlation")
    p.add_argument("--rules", metavar="FILE",
                   help="extra correlation rules (YAML/JSON) layered over the built-ins")
    p.add_argument("--correlation-report", metavar="FILE",
                   help="write what was correlated, and why, as JSON")
    p.add_argument("--think-time", type=int, help="add this think time (ms) between requests")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("record",
                       help="launch a browser, you click through it, get a .jmx", parents=[prox])
    p.add_argument("sources", nargs="+", metavar="URL|FILE",
                   help="URLs, or a .txt/.csv/.xlsx list of them ('-' = stdin)")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--manual", action="store_true",
                   help="open a visible browser and let you drive it; without this it "
                        "walks the URL list headlessly")
    p.add_argument("--settle", type=int, default=1500, metavar="MS",
                   help="wait after each page loads so lazy/XHR traffic is captured")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--har", help="keep the captured HAR at this path")
    p.add_argument("--browser", default="chromium",
                   choices=["chromium", "chrome", "edge", "firefox", "webkit"])
    p.add_argument("--profile", metavar="DIR",
                   help="persistent browser profile - log in once, reuse it next time")
    p.add_argument("--keep-static", action="store_true")
    p.add_argument("--mode", choices=["auto", "api", "web"], default="auto",
                   help="api = service calls only; web = everything the browser "
                        "fetched; auto = drop static/trackers")
    p.add_argument("--methods", metavar="LIST",
                   help="keep only these HTTP methods, e.g. POST or POST,PUT,DELETE")
    p.add_argument("--real-think-time", action="store_true",
                   help="use the pauses the real user left, instead of a fixed value")
    p.add_argument("--keep-third-party", action="store_true")
    p.add_argument("--no-pages", action="store_true",
                   help="one flat flow instead of a transaction per page")
    p.add_argument("--no-correlate", action="store_true")
    p.add_argument("--rules", metavar="FILE",
                   help="extra correlation rules (YAML/JSON) layered over the built-ins")
    p.add_argument("--correlation-report", metavar="FILE",
                   help="write what was correlated, and why, as JSON")
    p.add_argument("--include", help="regex: keep only URLs matching")
    p.add_argument("--exclude", help="regex: drop URLs matching")
    p.add_argument("--think-time", type=int, help="think time between requests, ms")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("capture", parents=[prox],
                       help="record through a proxy (mobile, desktop, Postman, backend)")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--har", default="capture.har", help="where to keep the capture")
    p.add_argument("--port", type=int, default=8080)
    p.add_argument("--seconds", type=int, help="stop automatically after N seconds")
    p.add_argument("--mode", choices=["auto", "api", "web"], default="api")
    p.add_argument("--methods", metavar="LIST")
    p.add_argument("--include")
    p.add_argument("--exclude")
    p.add_argument("--keep-static", action="store_true")
    p.add_argument("--keep-third-party", action="store_true")
    p.add_argument("--no-pages", action="store_true")
    p.add_argument("--no-correlate", action="store_true")
    p.add_argument("--real-think-time", action="store_true")
    p.add_argument("--rules", metavar="FILE")
    p.add_argument("--correlation-report", metavar="FILE")
    p.add_argument("--think-time", type=int)
    p.add_argument("--spec")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true")

    p = sub.add_parser("probe", help="page URL(s) -> .jmx, no HAR and no recording", parents=[prox])
    p.add_argument("sources", nargs="+",
                   metavar="URL|FILE", help="URLs, or a .txt/.csv/.xlsx list of them ('-' = stdin)")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--no-assets", action="store_true", help="main page requests only")
    p.add_argument("--all-domains", action="store_true",
                   help="keep sub-resources from other domains (CDNs); trackers still dropped")
    p.add_argument("--keep-trackers", action="store_true")
    p.add_argument("--no-static", action="store_true", help="drop js/css/image/font sub-resources")
    p.add_argument("--parallel", nargs="?", type=int, const=6, metavar="N",
                   help="fan each page's sub-resources out through a Parallel Controller "
                        "capped at N (default 6) instead of running them serially")
    p.add_argument("--think-time", type=int, default=1000, help="pause between pages, ms")
    p.add_argument("--no-forms", action="store_true", help="skip form submissions")
    p.add_argument("--max-assets", type=int, default=400, help="cap sub-resources per page")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("from-openapi", help="OpenAPI/Swagger file or URL -> .jmx", parents=[prox])
    p.add_argument("spec_src", metavar="openapi")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--auth", help='Authorization header value, e.g. "Bearer ${TOKEN}"')
    p.add_argument("--server", help="override the base URL")
    p.add_argument("--include", help="regex: only operations matching")
    p.add_argument("--exclude", help="regex: skip operations matching")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("from-postman", help="Postman collection -> .jmx", parents=[prox])
    p.add_argument("collection")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("from-curl", help="curl command(s) -> .jmx", parents=[prox])
    p.add_argument("file", nargs="?", help="file with curl commands ('-' or omitted = stdin)")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("import-jmx", parents=[prox],
                       help="turn an existing .jmx back into an editable spec")
    p.add_argument("jmx")
    p.add_argument("-o", "--out", help="output .jmx (round-tripped)")
    p.add_argument("--spec", help="write the editable spec here")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true")

    p = sub.add_parser("to-playwright", parents=[prox],
                       help="spec, .jmx or HAR -> a runnable Playwright browser test")
    p.add_argument("source", help="a spec .yaml, a .jmx, or a recorded .har")
    p.add_argument("-o", "--out", help="output .py (default: alongside the source)")

    p = sub.add_parser("to-taurus", parents=[prox],
                       help="spec or .jmx -> Taurus YAML (runs under bzt / BlazeMeter)")
    p.add_argument("source", help="a spec .yaml or an existing .jmx")
    p.add_argument("-o", "--out", help="output .yml (default: alongside the source)")

    p = sub.add_parser("optimize", help="repair / slim an existing .jmx", parents=[prox])
    p.add_argument("jmx")
    p.add_argument("-o", "--out", required=True)
    p.add_argument("--keep-domains", help="comma-separated domains that are the system under test")
    p.add_argument("--drop-third-party", action="store_true",
                   help="keep only the busiest domain's requests")
    p.add_argument("--drop-static", action="store_true", help="drop js/css/image/font requests")
    p.add_argument("--keep-disabled", action="store_true", help="do not delete disabled elements")
    p.add_argument("--keep-listeners", action="store_true", help="do not remove heavy listeners")
    p.add_argument("--keep-trackers", action="store_true")
    p.add_argument("--cap-parallel", type=int, metavar="N",
                   help="enable and set the Parallel Controller thread cap")
    p.add_argument("--dedupe-headers", action="store_true",
                   help="hoist headers common to every sampler into one manager")
    p.add_argument("--no-timeouts", action="store_true", help="do not fill in missing timeouts")
    p.add_argument("--max-samplers", type=int)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the result")

    p = sub.add_parser("from-url", help="crawl a live URL -> ready-to-run .jmx", parents=[prox])
    p.add_argument("url")
    p.add_argument("-o", "--out", help="output .jmx (default: <domain>.jmx)")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--depth", type=int, default=1, help="link-follow depth (default 1)")
    p.add_argument("--max-pages", type=int, default=8)
    p.add_argument("--no-links", action="store_true", help="only the given page, no crawling")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("from-excel", help="Excel/CSV sheet -> ready-to-run .jmx", parents=[prox])
    p.add_argument("sheet")
    p.add_argument("-o", "--out", help="output .jmx")
    p.add_argument("--spec", help="also write the editable spec here")
    p.add_argument("--name")
    p.add_argument("--deep", action="store_true", help="also make JMeter load the tree")

    p = sub.add_parser("template", help="write an Excel/CSV template to fill in")
    p.add_argument("out", nargs="?", default="testplan.xlsx")

    p = sub.add_parser("replay",
                       help="run the plan once as a single user and diagnose failures")
    p.add_argument("jmx")
    p.add_argument("--keep", metavar="DIR", help="keep replay artifacts here")
    p.add_argument("--timeout", type=int, default=300)

    p = sub.add_parser("console", help="open the web console (this is the default)")
    p.add_argument("--port", type=int, default=8770)
    p.add_argument("--no-open", action="store_true", help="do not open a browser")
    p.add_argument("--idle-timeout", type=int, default=60, metavar="MIN",
                   help="stop and release the port after this many minutes "
                        "with no requests (0 = stay up forever, default 60)")

    p = sub.add_parser("ship", help="author and run on HyperExecute in one step")
    p.add_argument("rest", nargs=argparse.REMAINDER,
                   help="arguments passed straight to the pipeline")

    sub.add_parser("doctor", help="what is installed and what it unlocks")

    p = sub.add_parser("verify", help="check a .jmx is valid and loadable")
    p.add_argument("jmx", nargs="+")
    p.add_argument("--deep", action="store_true",
                   help="also make JMeter load the tree (thread groups disabled)")

    p = sub.add_parser("validate", help="lint an existing .jmx")
    p.add_argument("jmx", nargs="+")

    a = ap.parse_args()

    proxy_cfg = None
    if getattr(a, "proxy", None):
        _host, _, _port = a.proxy.rpartition(":")
        proxy_cfg = {"host": _host.split("//")[-1] or a.proxy, "port": _port or "8080"}
        if getattr(a, "proxy_auth", None) and ":" in a.proxy_auth:
            proxy_cfg["user"], proxy_cfg["password"] = a.proxy_auth.split(":", 1)
        # set before any authoring fetch happens, not just on the generated plan
        _cred = ("%s:%s@" % (proxy_cfg["user"], proxy_cfg.get("password", ""))
                 if proxy_cfg.get("user") else "")
        FETCH_PROXY["http"] = FETCH_PROXY["https"] = \
            "http://%s%s:%s" % (_cred, proxy_cfg["host"], proxy_cfg["port"])

    login_cfg = None
    if getattr(a, "login", None):
        parts = a.login.split(None, 1)
        method, path = (parts[0].upper(), parts[1]) if len(parts) == 2 else ("POST", parts[0])
        body = getattr(a, "login_body", None)
        if body:
            try:
                body = json.loads(body)
            except ValueError:
                pass                      # a form string is fine as-is
        login_cfg = {
            "var": "AUTH_TOKEN",
            "login": {"method": method, "path": path,
                      "extract": {"type": "json", "query": a.login_token}},
        }
        if body is not None:
            login_cfg["login"]["body"] = body
            if isinstance(body, dict):
                login_cfg["login"]["headers"] = {"Content-Type": "application/json"}
        if getattr(a, "login_header", None):
            login_cfg["header"] = a.login_header

    csv_cfg = []
    for item in (getattr(a, "csv", None) or []):
        file_, _, cols = item.partition(":")
        csv_cfg.append({"file": file_,
                        "variables": [c.strip() for c in cols.split(",") if c.strip()]})

    def _apply_proxy(spec):
        if proxy_cfg:
            spec["proxy"] = proxy_cfg
        if login_cfg:
            spec["auth"] = login_cfg
        if csv_cfg:
            spec["csv"] = (spec.get("csv") or []) + csv_cfg
        return spec


    if a.cmd == "init":
        spec = WD_SPEC if a.kind == "webdriver" else HTTP_SPEC
        out = a.out
        if out.endswith((".yaml", ".yml")) and not _have_yaml():
            out = out.rsplit(".", 1)[0] + ".json"
            print("PyYAML not installed -> writing JSON spec instead", file=sys.stderr)
        open(out, "w", encoding="utf-8").write(dump_spec(spec, out))
        print("wrote %s   ->  next: python3 jmxgen.py build %s" % (out, out))
        return 0

    if a.cmd == "build":
        spec = load_spec(a.spec)
        _apply_proxy(spec)
        jmx = build_plan(spec)
        out = a.out or (re.sub(r"[^\w.-]+", "_", spec.get("name", "testplan")) + ".jmx")
        open(out, "w", encoding="utf-8").write(jmx)
        print("wrote %s (%.1f KB)" % (out, len(jmx) / 1024.0))
        props = write_system_properties(spec, out)
        if props:
            print("wrote %s - the client certificate lives here, not in the .jmx"
                  % props)
            print("  run:    jmeter -n -t %s -S %s ..." % (out, os.path.basename(props)))
            print("  on HX:  upload it with the plan and add -S %s to the args"
                  % os.path.basename(props))
        errors, _ = verify(out, deep=a.deep)
        validate(out)
        return 1 if errors else 0

    param_bindings = {}
    for item in (getattr(a, "parameterize", None) or []):
        lit, _, col = item.partition("=")
        if lit and col:
            param_bindings[lit] = col

    def _emit(spec, out_arg, spec_arg, deep, note=""):
        _apply_proxy(spec)
        if getattr(a, "suggest_parameters", False):
            found = suggest_parameters(spec)
            if not found:
                print("no obvious per-user values found")
            for value, kind in found.items():
                print("  %-40s looks like a %s  ->  --parameterize '%s=%s'"
                      % (value[:40], kind, value, kind))
            return 0
        if param_bindings:
            spec, n = parameterize(spec, param_bindings)
            print("parameterized %d occurrence(s) across %d value(s)"
                  % (n, len(param_bindings)))
        if spec_arg:
            open(spec_arg, "w", encoding="utf-8").write(dump_spec(spec, spec_arg))
            print("wrote %s" % spec_arg)
        out = out_arg or (re.sub(r"[^\w.-]+", "_", spec["name"]) + ".jmx")
        open(out, "w", encoding="utf-8").write(build_plan(spec))
        print("wrote %s%s" % (out, (" - " + note) if note else ""))
        props = write_system_properties(spec, out)
        if props:
            print("wrote %s - the client certificate lives here, not in the .jmx"
                  % props)
            print("  run:    jmeter -n -t %s -S %s ..." % (out, os.path.basename(props)))
            print("  on HX:  upload it with the plan and add -S %s to the args"
                  % os.path.basename(props))
        errors, _ = verify(out, deep=deep)
        validate(out)
        if getattr(a, "replay", False) and not errors:
            try:
                result = replay(out)
                if result["failures"] or result["unresolved"]:
                    return 1
            except ValueError as exc:
                print("  ?  replay skipped: %s" % exc, file=sys.stderr)
        return 1 if errors else 0

    if a.cmd == "from-har":
        spec, info = har_to_spec(a.har, a.include, a.exclude, a.keep_static, a.name,
                                 pages=not a.no_pages,
                                 drop_third_party=not a.keep_third_party,
                                 correlate=not a.no_correlate,
                                 think_time=a.think_time, mode=a.mode,
                                 methods=a.methods,
                                 real_think_time=a.real_think_time,
                                 rules=load_rules(getattr(a, "rules", None)))
        sk = info["skipped"]
        print("kept %d of %d recorded requests across %d page(s) "
              "(dropped %d static, %d third-party, %d filtered%s)"
              % (info["kept"], info["total"], info["pages"],
                 sk["static"], sk["third_party"], sk["filtered"],
                 ", %d CORS preflight" % sk["preflight"] if sk.get("preflight") else ""))
        _report_correlations(info["correlated"], getattr(a, "correlation_report", None))
        if not info["kept"]:
            sys.exit("nothing left after filtering - loosen --include/--exclude "
                     "or pass --keep-third-party/--keep-static")
        return _emit(spec, a.out, a.spec, a.deep)

    if a.cmd == "record":
        import tempfile
        urls = []
        for src in a.sources:
            urls += _read_url_list(src)
        urls = [u if "//" in u else "https://" + u for u in urls if u]
        if not urls:
            sys.exit("no URLs found in %s" % ", ".join(a.sources))
        har_path = a.har or os.path.join(tempfile.mkdtemp(prefix="jmxgen-rec-"),
                                         "session.har")
        print("recording %d URL(s) with %s%s"
              % (len(urls), a.browser, " (manual)" if a.manual else " (headless)"))
        rec_proxy = _apply_proxy({}).get("proxy")
        _har, warns = record_session(urls, har_path, browser_name=a.browser,
                                     profile_dir=a.profile, manual=a.manual,
                                     settle_ms=a.settle, proxy=rec_proxy)
        for w in warns:
            print("  ?  %s" % w, file=sys.stderr)
        size = os.path.getsize(har_path) / 1048576.0
        print("captured %s (%.1f MB)%s"
              % (har_path, size, "" if a.har else " [temporary]"))
        spec, info = har_to_spec(har_path, a.include, a.exclude, a.keep_static, a.name,
                                 pages=not a.no_pages,
                                 drop_third_party=not a.keep_third_party,
                                 correlate=not a.no_correlate,
                                 think_time=a.think_time, mode=a.mode,
                                 methods=a.methods,
                                 real_think_time=a.real_think_time,
                                 rules=load_rules(getattr(a, "rules", None)))
        sk = info["skipped"]
        print("kept %d of %d captured requests across %d page(s) "
              "(dropped %d static, %d third-party, %d filtered%s)"
              % (info["kept"], info["total"], info["pages"],
                 sk["static"], sk["third_party"], sk["filtered"],
                 ", %d CORS preflight" % sk["preflight"] if sk.get("preflight") else ""))
        _report_correlations(info["correlated"], getattr(a, "correlation_report", None))
        if not info["kept"]:
            sys.exit("nothing was captured that looks like application traffic")
        rc = _emit(spec, a.out, a.spec, a.deep)
        if not a.har:                  # embedded response bodies make these large
            import shutil
            shutil.rmtree(os.path.dirname(har_path), ignore_errors=True)
        return rc

    if a.cmd == "capture":
        capture_via_proxy(a.har, port=a.port, duration=a.seconds)
        spec, info = har_to_spec(a.har, a.include, a.exclude, a.keep_static, a.name,
                                 pages=not a.no_pages,
                                 drop_third_party=not a.keep_third_party,
                                 correlate=not a.no_correlate,
                                 think_time=a.think_time, mode=a.mode,
                                 methods=a.methods,
                                 real_think_time=a.real_think_time,
                                 rules=load_rules(a.rules))
        sk = info["skipped"]
        print("kept %d of %d captured requests (dropped %d static, %d third-party, %d filtered%s)"
              % (info["kept"], info["total"], sk["static"], sk["third_party"], sk["filtered"],
                 ", %d CORS preflight" % sk["preflight"] if sk.get("preflight") else ""))
        _report_correlations(info["correlated"], a.correlation_report)
        if not info["kept"]:
            sys.exit("nothing captured looked like application traffic")
        return _emit(spec, a.out, a.spec, a.deep)

    if a.cmd == "probe":
        spec, st = probe_to_spec(a.sources, name=a.name, assets=not a.no_assets,
                                 same_domain_assets=not a.all_domains,
                                 keep_static=not a.no_static, keep_trackers=a.keep_trackers,
                                 parallel=a.parallel, think_time=a.think_time,
                                 forms=not a.no_forms, max_assets=a.max_assets)
        print("probed %d page(s), found %d sub-resource(s) "
              "(dropped %d third-party/tracker, %d static)"
              % (st["pages"], st["assets"], st["skipped_third_party"], st["skipped_static"]))
        for url, err in st["failed"]:
            print("  ?  could not fetch %s -> %s" % (url, err), file=sys.stderr)
        for url in st["blocked"]:
            print("  !  WAF / bot-challenge page returned, page skipped: %s" % url[:100],
                  file=sys.stderr)
        if st["blocked"]:
            print("  !  a plain HTTP client cannot pass those challenges - record a HAR "
                  "from your browser (from-har), or use from-openapi / from-postman",
                  file=sys.stderr)
        return _emit(spec, a.out, a.spec, a.deep)

    if a.cmd == "from-openapi":
        spec, n = openapi_to_spec(a.spec_src, name=a.name, auth=a.auth,
                                  include=a.include, exclude=a.exclude, server=a.server)
        if not n:
            sys.exit("no operations found in %s" % a.spec_src)
        return _emit(spec, a.out, a.spec, a.deep, "%d operation(s)" % n)

    if a.cmd == "from-postman":
        spec, n = postman_to_spec(a.collection, name=a.name)
        if not n:
            sys.exit("no requests found in %s" % a.collection)
        return _emit(spec, a.out, a.spec, a.deep, "%d request(s)" % n)

    if a.cmd == "from-curl":
        text = (sys.stdin.read() if a.file in (None, "-")
                else open(a.file, "r", encoding="utf-8").read())
        spec, n = curl_to_spec(text, name=a.name)
        return _emit(spec, a.out, a.spec, a.deep, "%d request(s)" % n)

    if a.cmd == "import-jmx":
        spec, counts = jmx_to_spec(a.jmx, name=a.name)
        print("imported %d HTTP sampler(s); %d element(s) kept verbatim"
              % (counts["samplers"], counts["raw"]))
        spec_path = a.spec or (os.path.splitext(a.jmx)[0] + ".spec.yaml")
        open(spec_path, "w", encoding="utf-8").write(dump_spec(spec, spec_path))
        print("wrote %s" % spec_path)
        if a.out:
            return _emit(spec, a.out, None, a.deep)
        return 0

    if a.cmd == "to-playwright":
        low = a.source.lower()
        if low.endswith(".jmx"):
            spec, _ = jmx_to_spec(a.source)
        elif low.endswith((".har", ".json")):
            spec, _ = har_to_spec(a.source)
        elif low.endswith((".yaml", ".yml")):
            spec = load_spec(a.source)
        else:
            print("cannot read %s - expected a spec (.yaml), a plan (.jmx) "
                  "or a recording (.har)" % a.source)
            return 1
        if not spec_has_browser_steps(spec):
            print("no browser steps in %s" % a.source)
            print("record with the extension and leave 'record browser steps' ticked")
            return 1
        out = a.out or (os.path.splitext(a.source)[0] + "_browser_test.py")
        open(out, "w", encoding="utf-8").write(
            spec_to_playwright(spec, os.path.basename(out)))
        print("wrote %s" % out)
        print("run it with:  python3 %s" % out)
        return 0

    if a.cmd == "to-taurus":
        if a.source.lower().endswith(".jmx"):
            spec, _ = jmx_to_spec(a.source)
        else:
            spec = load_spec(a.source)
        out = a.out or (os.path.splitext(a.source)[0] + ".taurus.yml")
        open(out, "w", encoding="utf-8").write(dump_taurus(spec))
        n = sum(1 for tg in spec.get("thread_groups") or []
                for _ in _walk_steps(tg.get("steps") or []))
        print("wrote %s  (%d request(s))" % (out, n))
        print("run it with:  bzt %s" % out)
        return 0

    if a.cmd == "optimize":
        stats = optimize(a.jmx, a.out,
                         keep_domains=(a.keep_domains.split(",") if a.keep_domains else None),
                         drop_third_party=a.drop_third_party or bool(a.keep_domains),
                         drop_static=a.drop_static,
                         drop_disabled=not a.keep_disabled,
                         strip_listeners=not a.keep_listeners,
                         cap_parallel=a.cap_parallel,
                         dedupe_headers=a.dedupe_headers,
                         add_timeouts=not a.no_timeouts,
                         keep_trackers=a.keep_trackers,
                         max_samplers=a.max_samplers,
                         dry_run=a.dry_run,
                         proxy=_apply_proxy({}).get("proxy"))
        print_optimize_report(a.jmx, a.out, stats)
        if a.dry_run:
            print("  (dry run - nothing written)")
            return 0
        errors, _ = verify(a.out, deep=a.deep)
        validate(a.out)
        return 1 if errors else 0

    if a.cmd == "from-url":
        spec, errs = url_to_spec(a.url, depth=a.depth, max_pages=a.max_pages,
                                 name=a.name, follow_links=not a.no_links)
        for e in errs:
            print("  ?  could not fetch %s" % e, file=sys.stderr)
        n_steps = sum(len(t["steps"]) for t in spec["thread_groups"][0]["steps"])
        if not n_steps:
            sys.exit("nothing could be authored from %s (no page fetched)" % a.url)
        if a.spec:
            open(a.spec, "w", encoding="utf-8").write(dump_spec(spec, a.spec))
            print("wrote %s" % a.spec)
        out = a.out or (re.sub(r"[^\w.-]+", "_", spec["name"]) + ".jmx")
        open(out, "w", encoding="utf-8").write(build_plan(spec))
        print("wrote %s (%d request(s) across %d page(s))"
              % (out, n_steps, len(spec["thread_groups"][0]["steps"])))
        errors, _ = verify(out, deep=a.deep)
        validate(out)
        return 1 if errors else 0

    if a.cmd == "from-excel":
        spec = sheet_to_spec(a.sheet, name=a.name)
        if a.spec:
            open(a.spec, "w", encoding="utf-8").write(dump_spec(spec, a.spec))
            print("wrote %s" % a.spec)
        out = a.out or (re.sub(r"[^\w.-]+", "_", spec["name"]) + ".jmx")
        open(out, "w", encoding="utf-8").write(build_plan(spec))
        n = sum(len(tg["steps"]) for tg in spec["thread_groups"])
        print("wrote %s (%d thread group(s), %d top-level step(s))"
              % (out, len(spec["thread_groups"]), n))
        errors, _ = verify(out, deep=a.deep)
        validate(out)
        return 1 if errors else 0

    if a.cmd == "template":
        out = write_template(a.out)
        print("wrote %s   ->  fill it in, then: python3 jmxgen.py from-excel %s" % (out, out))
        return 0

    if a.cmd == "console":
        return start_console(a.port, open_browser=not a.no_open,
                             idle_timeout=a.idle_timeout)

    if a.cmd == "ship":
        import subprocess
        script = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "jmx_pipeline.py")
        if not os.path.exists(script):
            sys.exit("jmx_pipeline.py is not next to jmxgen.py")
        return subprocess.run([sys.executable, script] + list(a.rest)).returncode

    if a.cmd == "doctor":
        return doctor()

    if a.cmd == "replay":
        result = replay(a.jmx, out_dir=a.keep, timeout=a.timeout)
        return 1 if (result["failures"] or result["unresolved"]) else 0

    if a.cmd == "verify":
        rc = 0
        for f in a.jmx:
            errors, _ = verify(f, deep=a.deep)
            rc |= (1 if errors else 0)
        return rc

    if a.cmd == "validate":
        rc = 0
        for f in a.jmx:
            rc |= validate(f)
        return rc


if __name__ == "__main__":
    try:
        sys.exit(main())
    except ValueError as exc:          # authoring problems are user-facing, not crashes
        sys.exit("error: %s" % exc)
    except KeyboardInterrupt:
        sys.exit(130)
